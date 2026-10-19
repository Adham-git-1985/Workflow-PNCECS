import tempfile
import unittest
from datetime import datetime
from io import BytesIO
from unittest.mock import patch

from flask import Flask
from openpyxl import load_workbook

from extensions import db
from models import WorkflowReportEmailDelivery
import services.workflow_report_email as workflow_email


class WorkflowReportEmailTests(unittest.TestCase):
    def test_default_schedule_is_ten_am_and_daily(self):
        with patch.object(workflow_email, "_setting", side_effect=lambda _key, default=None: default):
            config = workflow_email.get_report_email_config()

        self.assertEqual(config["send_times_list"], ["10:00"])
        self.assertEqual(config["send_time"], "10:00")
        self.assertEqual(config["frequency"], "DAILY")
        self.assertEqual(config["interval_days"], 1)

    def test_cycle_runs_only_the_configured_current_minute(self):
        config = {
            "enabled": True,
            "send_times_list": ["10:00"],
            "send_time": "10:00",
            "frequency": "DAILY",
            "skip_weekly_holidays": False,
            "skip_official_holidays": False,
            "excluded_dates_list": [],
        }
        with (
            patch.object(workflow_email, "get_report_email_config", return_value=config),
            patch.object(workflow_email, "_schedule_is_due", return_value=True),
            patch.object(workflow_email, "_send_day_is_excluded", return_value=False),
            patch.object(workflow_email, "_record_scheduler_state"),
            patch.object(workflow_email, "_send_slot", return_value={"status": "sent"}) as send_slot,
        ):
            result = workflow_email.run_workflow_report_email_cycle(datetime(2026, 10, 19, 10, 0))

        self.assertEqual(result["status"], "sent")
        send_slot.assert_called_once()

    def test_cycle_does_not_catch_up_a_missed_slot(self):
        config = {
            "enabled": True,
            "send_times_list": ["10:00"],
            "send_time": "10:00",
            "frequency": "DAILY",
            "skip_weekly_holidays": False,
            "skip_official_holidays": False,
            "excluded_dates_list": [],
        }
        with (
            patch.object(workflow_email, "get_report_email_config", return_value=config),
            patch.object(workflow_email, "_schedule_is_due", return_value=True),
            patch.object(workflow_email, "_send_day_is_excluded", return_value=False),
            patch.object(workflow_email, "_record_scheduler_state"),
            patch.object(workflow_email, "_send_slot") as send_slot,
        ):
            result = workflow_email.run_workflow_report_email_cycle(datetime(2026, 10, 19, 10, 1))

        self.assertEqual(result["status"], "not_due")
        send_slot.assert_not_called()

    def test_email_content_and_excel_attachment_include_summary_and_details(self):
        data = {
            "include_details": True,
            "scheduled_time": "10:00",
            "summary": [{
                "employee": "Employee One",
                "job_title": "Manager",
                "department": "Operations",
                "total_paths": 2,
                "active_paths": 1,
                "statuses": "In progress: 2",
                "last_assigned_at": "2026-10-19 09:15",
                "last_opened_at": "2026-10-19 09:30",
            }],
            "details": [{
                "employee": "Employee One",
                "request_id": 7,
                "path": "Request path",
                "status": "In progress",
                "requester": "Requester",
                "steps": "1, 2",
                "active": "Yes",
                "created_at": "2026-10-19 08:00",
                "assignment_at": "2026-10-19 09:15",
                "opened_at": "2026-10-19 09:30",
            }],
        }

        subject, text_body, html_body = workflow_email.build_email_content(data)
        workbook = load_workbook(BytesIO(workflow_email._xlsx_bytes(data)), read_only=True)
        try:
            self.assertIn("10:00", subject)
            self.assertIn("Employee One", text_body)
            self.assertIn("Employee One", html_body)
            self.assertGreaterEqual(len(workbook.sheetnames), 2)
        finally:
            workbook.close()

    def test_on_demand_send_records_a_delivery_history_row(self):
        temp_dir = tempfile.TemporaryDirectory()
        app = Flask(__name__, instance_path=temp_dir.name)
        app.config.update(
            TESTING=True,
            SECRET_KEY="workflow-report-email-test",
            SQLALCHEMY_DATABASE_URI="sqlite:///:memory:",
            SQLALCHEMY_TRACK_MODIFICATIONS=False,
        )
        db.init_app(app)
        context = app.app_context()
        context.push()
        WorkflowReportEmailDelivery.__table__.create(db.engine)
        try:
            config = {"include_details": False, "recipient_emails_list": ["recipient@example.test"], "recipient_user_ids_list": []}
            report_data = {
                "include_details": False,
                "summary": [],
                "details": [],
                "total_assignments": 0,
            }
            with (
                patch.object(workflow_email, "get_report_email_config", return_value=config),
                patch.object(workflow_email, "_recipients", return_value=["recipient@example.test"]),
                patch.object(workflow_email, "email_delivery_enabled", return_value=True),
                patch.object(workflow_email, "_mail_config", return_value={"ready": True}),
                patch.object(workflow_email, "_report_viewer", return_value=object()),
                patch.object(workflow_email, "build_workflow_report_data", return_value=report_data),
                patch.object(workflow_email, "_send_email"),
            ):
                result = workflow_email.send_workflow_report_email_now()

            row = WorkflowReportEmailDelivery.query.one()
            self.assertEqual(result["status"], "sent")
            self.assertEqual(row.trigger, "ON_DEMAND")
            self.assertEqual(row.status, "SENT")
            self.assertEqual(row.recipient_count, 1)
        finally:
            db.session.remove()
            context.pop()
            temp_dir.cleanup()


if __name__ == "__main__":
    unittest.main()
