import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
import zipfile

from werkzeug.datastructures import FileStorage

from services.correspondence_intake import (
    CorrespondenceIntakeError,
    OcrConfig,
    analyze_correspondence_attachment,
    analyze_workflow_attachment,
    extract_attachment_text,
    read_limited_upload,
)


class CorrespondenceIntakeTests(unittest.TestCase):
    def test_workflow_attachment_produces_request_and_template_suggestions(self):
        source = """
        الموضوع: طلب شراء أجهزة حاسوب
        يرجى تفريغ الطلب على مسار المشتريات ومتابعة إجراءات Procurement.
        """.encode("utf-8")

        result = analyze_workflow_attachment(
            source,
            "طلب_شراء.txt",
            request_type_choices=[
                {
                    "value": "7",
                    "label": "طلب شراء",
                    "match_text": "طلب شراء Procurement PRC",
                },
                {"value": "8", "label": "طلب إداري"},
            ],
            workflow_choices=[
                {"value": "12", "label": "مسار المشتريات"},
                {"value": "13", "label": "مسار إداري"},
            ],
        )

        suggestions = result["suggestions"]
        self.assertEqual(suggestions["title"]["value"], "طلب شراء أجهزة حاسوب")
        self.assertIn("إجراءات Procurement", suggestions["description"]["value"])
        self.assertEqual(suggestions["request_type"]["select_value"], "7")
        self.assertEqual(suggestions["request_type"]["value"], "طلب شراء")
        self.assertEqual(suggestions["workflow_template"]["select_value"], "12")
        self.assertEqual(result["privacy"], "LOCAL_ONLY")

    def test_arabic_text_produces_reviewable_correspondence_suggestions(self):
        source = """
        الجهة المرسلة: وزارة التربية والتعليم
        الموضوع: طلب صرف دفعة مالية عاجل
        رقم المرجع: FIN/2026/44
        الموعد النهائي: 2026-08-30
        سري
        يرجى تحويل الطلب إلى دائرة الشؤون المالية عبر مسار المالية.
        """.encode("utf-8")

        result = analyze_correspondence_attachment(
            source,
            "طلب_مالي.txt",
            sender_choices=[
                {"value": "وزارة التربية والتعليم", "label": "وزارة التربية والتعليم"}
            ],
            category_choices=[
                {"value": "GENERAL", "label": "عام"},
                {"value": "FIN", "label": "مالية"},
            ],
            competence_choices=[
                {
                    "value": "DEPARTMENT:7",
                    "label": "دائرة الشؤون المالية (دائرة)",
                }
            ],
            workflow_choices=[
                {"value": "12", "label": "مسار المالية"},
                {"value": "13", "label": "مسار إداري"},
            ],
            received_date="2026-08-17",
        )

        suggestions = result["suggestions"]
        self.assertEqual(suggestions["received_date"]["value"], "2026-08-17")
        self.assertEqual(suggestions["subject"]["value"], "طلب صرف دفعة مالية عاجل")
        self.assertEqual(suggestions["sender"]["select_value"], "وزارة التربية والتعليم")
        self.assertTrue(suggestions["sender"]["is_known"])
        self.assertEqual(suggestions["category"]["select_value"], "FIN")
        self.assertEqual(suggestions["competence"]["select_value"], "DEPARTMENT:7")
        self.assertEqual(suggestions["workflow_template"]["select_value"], "12")
        self.assertEqual(suggestions["priority"]["value"], "URGENT")
        self.assertEqual(suggestions["confidentiality"]["value"], "SECRET")
        self.assertEqual(suggestions["due_date"]["value"], "2026-08-30")
        self.assertEqual(suggestions["document_reference"]["value"], "FIN/2026/44")
        self.assertEqual(result["privacy"], "LOCAL_ONLY")

    def test_outbound_analysis_uses_recipient_and_sent_date_fields(self):
        source = """
        الجهة المستلمة: وزارة الصحة
        الموضوع: إرسال التقرير المالي
        يرجى متابعة الصادر عبر مسار المالية.
        """.encode("utf-8")

        result = analyze_correspondence_attachment(
            source,
            "تقرير_صادر.txt",
            recipient_choices=[
                {"value": "وزارة الصحة", "label": "وزارة الصحة"},
            ],
            workflow_choices=[
                {"value": "12", "label": "مسار المالية"},
            ],
            sent_date="2026-08-17",
            direction="OUT",
        )

        suggestions = result["suggestions"]
        self.assertEqual(suggestions["sent_date"]["value"], "2026-08-17")
        self.assertEqual(suggestions["recipient"]["select_value"], "وزارة الصحة")
        self.assertTrue(suggestions["recipient"]["is_known"])
        self.assertEqual(suggestions["workflow_template"]["select_value"], "12")
        self.assertNotIn("received_date", suggestions)
        self.assertNotIn("sender", suggestions)

    def test_eml_uses_email_headers_and_ignores_attachment_payload(self):
        payload = (
            "From: Ministry Example <office@example.test>\r\n"
            "To: intake@example.test\r\n"
            "Subject: Follow-up request\r\n"
            "MIME-Version: 1.0\r\n"
            "Content-Type: multipart/mixed; boundary=part\r\n\r\n"
            "--part\r\nContent-Type: text/plain; charset=utf-8\r\n\r\n"
            "الموضوع: نص الرسالة\r\n"
            "--part\r\nContent-Type: application/octet-stream\r\n"
            "Content-Disposition: attachment; filename=secret.bin\r\n\r\n"
            "SHOULD_NOT_BE_EXTRACTED\r\n--part--\r\n"
        ).encode("utf-8")

        result = analyze_correspondence_attachment(
            payload,
            "message.eml",
            sender_choices=[{"value": "Ministry Example", "label": "Ministry Example"}],
        )

        self.assertEqual(result["format"], "Email")
        self.assertEqual(result["suggestions"]["subject"]["value"], "Follow-up request")
        self.assertEqual(result["suggestions"]["sender"]["select_value"], "Ministry Example")
        self.assertNotIn("SHOULD_NOT_BE_EXTRACTED", result["suggestions"]["body"]["value"])

    def test_image_returns_manual_ocr_warning_and_filename_subject(self):
        result = analyze_correspondence_attachment(b"not-a-real-image", "scan.jpg")

        self.assertEqual(result["format"], "Image")
        self.assertEqual(result["suggestions"]["subject"]["value"], "scan")
        self.assertNotIn("body", result["suggestions"])
        self.assertTrue(any("OCR" in warning for warning in result["warnings"]))

    def test_image_ocr_text_is_used_for_suggestions_when_local_engine_is_available(self):
        from PIL import Image

        image_stream = BytesIO()
        Image.new("RGB", (600, 300), "white").save(image_stream, format="PNG")
        ocr_config = OcrConfig(enabled=True, command="fake-tesseract")

        with patch(
            "services.correspondence_intake._resolve_tesseract_command",
            return_value="fake-tesseract",
        ), patch(
            "services.correspondence_intake._run_tesseract_png",
            return_value="الموضوع: طلب مستخرج من صورة\nعاجل",
        ):
            result = analyze_correspondence_attachment(
                image_stream.getvalue(),
                "scan.png",
                ocr_config=ocr_config,
            )

        self.assertEqual(
            result["suggestions"]["subject"]["value"],
            "طلب مستخرج من صورة",
        )
        self.assertEqual(result["suggestions"]["priority"]["value"], "URGENT")
        self.assertTrue(result["ocr"]["enabled"])
        self.assertTrue(result["ocr"]["available"])
        self.assertTrue(result["ocr"]["used"])

    def test_scanned_pdf_uses_local_ocr_when_regular_text_is_empty(self):
        import fitz

        pdf_document = fitz.open()
        pdf_document.new_page()
        payload = pdf_document.tobytes()
        pdf_document.close()

        with patch(
            "services.correspondence_intake._resolve_tesseract_command",
            return_value="fake-tesseract",
        ), patch(
            "services.correspondence_intake._run_tesseract_png",
            return_value="Subject: Scanned PDF request",
        ):
            result = analyze_correspondence_attachment(
                payload,
                "scanned.pdf",
                ocr_config=OcrConfig(enabled=True, command="fake-tesseract"),
            )

        self.assertEqual(
            result["suggestions"]["subject"]["value"],
            "Scanned PDF request",
        )
        self.assertTrue(result["ocr"]["used"])
        self.assertEqual(result["ocr"]["pages"], 1)

    def test_pdf_ocr_skips_pages_that_exceed_safe_pixel_limit(self):
        import fitz

        pdf_document = fitz.open()
        pdf_document.new_page(width=5000, height=5000)
        payload = pdf_document.tobytes()
        pdf_document.close()

        with patch(
            "services.correspondence_intake._resolve_tesseract_command",
            return_value="fake-tesseract",
        ), patch(
            "services.correspondence_intake._run_tesseract_png"
        ) as run_ocr:
            result = analyze_correspondence_attachment(
                payload,
                "oversized-page.pdf",
                ocr_config=OcrConfig(
                    enabled=True,
                    command="fake-tesseract",
                    max_image_pixels=1_000_000,
                ),
            )

        run_ocr.assert_not_called()
        self.assertFalse(result["ocr"]["used"])
        self.assertTrue(any("الحد الآمن" in warning for warning in result["warnings"]))

    def test_missing_tesseract_falls_back_to_manual_entry(self):
        with patch(
            "services.correspondence_intake._resolve_tesseract_command",
            return_value=None,
        ):
            result = analyze_correspondence_attachment(
                b"image-placeholder",
                "scan.jpg",
                ocr_config=OcrConfig(enabled=True),
            )

        self.assertFalse(result["ocr"]["available"])
        self.assertFalse(result["ocr"]["used"])
        self.assertTrue(any("Tesseract" in warning for warning in result["warnings"]))

    def test_legacy_doc_rtf_content_is_analyzed_without_external_office(self):
        source_text = "الموضوع: طلب شراء أجهزة قديمة\nيرجى مراجعة الطلب واعتماده."
        rtf_text = "{\\rtf1\\ansi\\ansicpg1256 " + "".join(
            (
                f"\\u{ord(character)}?"
                if ord(character) > 127
                else ("\\par " if character == "\n" else character)
            )
            for character in source_text
        ) + "}"

        result = analyze_workflow_attachment(
            rtf_text.encode("ascii"),
            "legacy-request.doc",
        )

        self.assertEqual(result["format"], "Word (DOC)")
        self.assertEqual(
            result["suggestions"]["title"]["value"],
            "طلب شراء أجهزة قديمة",
        )
        self.assertIn("مراجعة الطلب", result["suggestions"]["description"]["value"])

    def test_binary_legacy_doc_uses_bounded_direct_text_fallback(self):
        ole_signature = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
        expected_text = "الموضوع: طلب متابعة ملف قديم"
        payload = ole_signature + (b"\x00" * 128) + expected_text.encode("utf-16le") + b"\x01\x02"

        with patch(
            "services.correspondence_intake._resolve_legacy_word_converter",
            return_value=None,
        ):
            result = extract_attachment_text(payload, "legacy.doc")

        self.assertEqual(result["format"], "Word (DOC)")
        self.assertIn(expected_text, result["text"])
        self.assertTrue(any("الاحتياطية" in warning for warning in result["warnings"]))

    def test_binary_legacy_doc_filters_internal_noise_and_preserves_table_rows(self):
        ole_signature = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
        table_text = (
            "تقرير متابعة الأعمال – أغسطس 2026\r"
            "العمل / النشاط\x07التفاصيل\x07التاريخ\x07\x07"
            "متابعة نظام مسار\x07"
            "مراجعة الأعمال والملاحظات المرتبطة بالنظام خلال التشغيل والتسليم\x07"
            "25 أغسطس 2026\x07\x07"
            "تحديث بيانات الموظفين\x07"
            "إعداد وتنظيم بيانات الموظفين تمهيدًا لإدخالها إلى النظام\x07"
            "24 أغسطس 2026\x07\x07"
        )
        noisy_prefix = ("ے" * 300).encode("utf-16le")
        noisy_suffix = (
            "Heading 1 Default Paragraph Font Table Normal Table Grid xmlns "
            "㐀䈀䐀耀蠀ﳲﳲﳲﳲﳲﳲ"
        ).encode("utf-16le")
        payload = (
            ole_signature
            + noisy_prefix
            + b"\x00\x00"
            + table_text.encode("utf-16le")
            + b"\x00\x00"
            + noisy_suffix
        )

        with patch(
            "services.correspondence_intake._resolve_legacy_word_converter",
            return_value=None,
        ):
            result = extract_attachment_text(payload, "legacy-table.doc")

        self.assertIn("العمل / النشاط | التفاصيل | التاريخ", result["text"])
        self.assertIn(
            "متابعة نظام مسار | مراجعة الأعمال والملاحظات المرتبطة بالنظام خلال التشغيل والتسليم | 25 أغسطس 2026",
            result["text"],
        )
        self.assertNotIn("ےےے", result["text"])
        self.assertNotIn("Heading 1", result["text"])
        self.assertNotIn("Table Grid", result["text"])
        self.assertNotIn("㐀", result["text"])
        self.assertNotIn("ﳲ", result["text"])

    def test_unsupported_file_remains_available_only_for_normal_upload(self):
        with self.assertRaises(CorrespondenceIntakeError) as context:
            extract_attachment_text(b"data", "legacy.msg")

        self.assertEqual(context.exception.code, "UNSUPPORTED_FILE_TYPE")
        self.assertEqual(context.exception.status_code, 415)

    def test_upload_size_is_bounded_while_reading(self):
        upload = FileStorage(stream=BytesIO(b"123456"), filename="letter.txt")

        with self.assertRaises(CorrespondenceIntakeError) as context:
            read_limited_upload(upload, 5)

        self.assertEqual(context.exception.code, "FILE_TOO_LARGE")
        self.assertEqual(context.exception.status_code, 413)
        self.assertIn("5 بايت", context.exception.message)
        self.assertIn("حفظه كمرفق دون تحليله", context.exception.message)

    def test_word_excel_pdf_and_powerpoint_text_are_extracted_locally(self):
        from docx import Document
        from openpyxl import Workbook
        import fitz

        word_stream = BytesIO()
        word_document = Document()
        word_document.add_paragraph("الموضوع: كتاب Word تجريبي")
        word_document.save(word_stream)
        word_result = extract_attachment_text(word_stream.getvalue(), "letter.docx")
        self.assertIn("كتاب Word تجريبي", word_result["text"])

        excel_stream = BytesIO()
        workbook = Workbook()
        workbook.active.append(["الموضوع: جدول Excel تجريبي"])
        workbook.save(excel_stream)
        workbook.close()
        excel_result = extract_attachment_text(excel_stream.getvalue(), "table.xlsx")
        self.assertIn("جدول Excel تجريبي", excel_result["text"])

        pdf_document = fitz.open()
        page = pdf_document.new_page()
        page.insert_text((72, 72), "Subject: Local PDF intake")
        pdf_payload = pdf_document.tobytes()
        pdf_document.close()
        pdf_result = extract_attachment_text(pdf_payload, "letter.pdf")
        self.assertIn("Local PDF intake", pdf_result["text"])

        powerpoint_stream = BytesIO()
        with zipfile.ZipFile(powerpoint_stream, "w") as package:
            package.writestr(
                "ppt/slides/slide1.xml",
                '<p:sld xmlns:p="urn:p" xmlns:a="urn:a"><a:t>PowerPoint intake</a:t></p:sld>',
            )
        powerpoint_result = extract_attachment_text(
            powerpoint_stream.getvalue(), "slides.pptx"
        )
        self.assertIn("PowerPoint intake", powerpoint_result["text"])

    def test_inbound_form_exposes_review_and_save_start_actions(self):
        project_root = Path(__file__).resolve().parents[1]
        template = (
            project_root / "templates" / "portal" / "corr" / "inbound_new.html"
        ).read_text(encoding="utf-8")
        routes = (project_root / "portal" / "routes.py").read_text(encoding="utf-8")

        self.assertIn('id="analyzeInboundAttachment"', template)
        self.assertIn('id="applyInboundSuggestions"', template)
        self.assertLess(template.index('id="inboundFiles"'), template.index('id="received_date"'))
        self.assertIn('data-max-bytes="{{ intake_max_bytes }}"', template)
        self.assertIn('value="save_and_start"', template)
        self.assertIn('request.form.get("submit_action") == "save_and_start"', routes)
        self.assertIn('workflow_request = _start_corr_workflow(', routes)
        self.assertIn("الخطوة الاختيارية: ارفع المرفق وحلله", template)
        self.assertNotIn('id="inboundFiles" multiple required', template)

    def test_outbound_form_exposes_optional_attachment_analysis(self):
        project_root = Path(__file__).resolve().parents[1]
        template = (
            project_root / "templates" / "portal" / "corr" / "outbound_new.html"
        ).read_text(encoding="utf-8")
        routes = (project_root / "portal" / "routes.py").read_text(encoding="utf-8")
        view_template = (
            project_root / "templates" / "portal" / "corr" / "outbound_view.html"
        ).read_text(encoding="utf-8")

        self.assertIn('id="analyzeOutboundAttachment"', template)
        self.assertIn('id="applyOutboundSuggestions"', template)
        self.assertIn("الخطوة الاختيارية: ارفع المرفق وحلله", template)
        self.assertLess(template.index('id="outboundFiles"'), template.index('id="sent_date"'))
        self.assertIn('data-max-bytes="{{ intake_max_bytes }}"', template)
        self.assertNotIn('id="outboundFiles" multiple required', template)
        self.assertIn('route("/corr/outbound/analyze-attachment"', routes)
        self.assertIn('direction="OUT"', routes)
        self.assertIn('id="outboundParallelSelection"', view_template)
        self.assertIn('name="parallel_assignee_ids"', view_template)
        self.assertIn("workflow_parallel_candidates", view_template)
        self.assertIn(
            'initial_parallel_user_ids=request.form.getlist("parallel_assignee_ids")',
            routes,
        )

    def test_workflow_form_exposes_attachment_analysis_and_field_dump(self):
        project_root = Path(__file__).resolve().parents[1]
        template = (
            project_root / "templates" / "workflow" / "new_request.html"
        ).read_text(encoding="utf-8")
        routes = (project_root / "workflow" / "routes.py").read_text(encoding="utf-8")

        self.assertIn('id="analyzeWorkflowAttachment"', template)
        self.assertIn('id="applyWorkflowSuggestions"', template)
        self.assertIn("تفريغ البيانات على نموذج المسار", template)
        self.assertIn('name="files"', template)
        self.assertIn('data-max-bytes="{{ intake_max_bytes }}"', template)
        self.assertIn('route("/new/analyze-attachment"', routes)
        self.assertIn("analyze_workflow_attachment(", routes)
        self.assertIn('id="workflowAnalysisTools"', template)
        self.assertIn('data-bs-target="#workflowAnalysisTools"', template)
        self.assertIn('name="request_type_name"', template)
        self.assertIn('name="priority"', template)
        self.assertNotIn('name="request_type_id"', template)
        self.assertLess(
            template.index("<span class=\"badge text-bg-primary ms-1\">1</span> اختيار المسار"),
            template.index("<span class=\"badge text-bg-primary ms-1\">2</span> بيانات الطلب"),
        )
        self.assertLess(
            template.index("<span class=\"badge text-bg-primary ms-1\">2</span> بيانات الطلب"),
            template.index("<span class=\"badge text-bg-primary ms-1\">3</span> إرسال وحفظ"),
        )
        self.assertIn('id="dynamicHierarchySelect"', template)
        self.assertIn(
            'class="d-none" id="dynamicHierarchySelect"',
            template,
        )
        self.assertIn('id="dynamicHierarchyTree"', template)
        self.assertIn('class="dynamic-tree-node"', template)
        self.assertIn('id="dynamicPathPreviewPanel"', template)
        self.assertIn('عرض تحليل وتتبع المسار', template)
        self.assertIn('{{ node.name }}{% if node.total_user_count %}', template)
        self.assertNotIn('{{ node.type_name }}: {{ node.name }}', template)
        self.assertIn('#dynamicRouteBox .dynamic-contained-select', template)
        self.assertIn('overflow-wrap: anywhere', template)
        self.assertIn('dynamic-selected-actions', template)
        self.assertNotIn("dynamic-browser-tab", template)
        self.assertNotIn('id="dynamicTeamFilter"', template)
        self.assertIn("_find_or_create_request_type(request_type_name)", routes)

    def test_workflow_close_action_and_portal_return_button_layout_are_exposed(self):
        project_root = Path(__file__).resolve().parents[1]
        request_template = (
            project_root / "templates" / "workflow" / "view_request.html"
        ).read_text(encoding="utf-8")
        portal_layout = (
            project_root / "templates" / "portal" / "layout.html"
        ).read_text(encoding="utf-8")
        routes = (project_root / "workflow" / "routes.py").read_text(encoding="utf-8")

        self.assertIn("{% if can_close %}", request_template)
        self.assertIn("workflow.close_request", request_template)
        self.assertIn("إغلاق الطلب", request_template)
        self.assertIn('def close_request(request_id):', routes)
        self.assertIn('action="REQUEST_CLOSED"', routes)
        self.assertIn("portal-utilbar-layout", portal_layout)
        self.assertIn("grid-template-columns:max-content minmax(0, 1fr)", portal_layout)
        self.assertIn("portal-masar-return", portal_layout)


if __name__ == "__main__":
    unittest.main()
