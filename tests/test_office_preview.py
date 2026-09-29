import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import fitz
from docx import Document
from openpyxl import Workbook

from utils.office_preview import OfficePreviewError, convert_office_bytes_to_pdf, convert_office_to_pdf


class OfficePreviewTests(unittest.TestCase):
    def test_docx_is_returned_as_inline_pdf_without_libreoffice(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "document.docx"
            document = Document()
            document.add_heading("معاينة Word", level=1)
            document.add_paragraph("محتوى تجريبي للمعاينة")
            document.save(path)

            with (
                patch(
                    "utils.office_preview._convert_with_windows_office",
                    side_effect=OfficePreviewError("native converter unavailable"),
                ),
                patch("utils.office_preview.find_libreoffice_executable", return_value=None),
            ):
                pdf_bytes = convert_office_to_pdf(path, original_name=path.name)
                embedded_pdf_bytes = convert_office_bytes_to_pdf(
                    path.read_bytes(),
                    original_name=path.name,
                )

        self.assertTrue(pdf_bytes.startswith(b"%PDF"))
        self.assertTrue(embedded_pdf_bytes.startswith(b"%PDF"))
        with fitz.open(stream=pdf_bytes, filetype="pdf") as pdf:
            self.assertGreaterEqual(len(pdf), 1)

    def test_xlsx_is_returned_as_inline_pdf_without_libreoffice(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "sheet.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "البيانات"
            sheet.append(["الاسم", "القيمة"])
            sheet.append(["اختبار", 42])
            workbook.save(path)

            with (
                patch(
                    "utils.office_preview._convert_with_windows_office",
                    side_effect=OfficePreviewError("native converter unavailable"),
                ),
                patch("utils.office_preview.find_libreoffice_executable", return_value=None),
            ):
                pdf_bytes = convert_office_to_pdf(path, original_name=path.name)

        self.assertTrue(pdf_bytes.startswith(b"%PDF"))
        with fitz.open(stream=pdf_bytes, filetype="pdf") as pdf:
            self.assertGreaterEqual(len(pdf), 1)


if __name__ == "__main__":
    unittest.main()
