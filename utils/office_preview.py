"""Convert Word and Excel attachments to an inline PDF preview.

On Windows, installed Microsoft Word/Excel are preferred to preserve the
source document's layout. LibreOffice is the portable secondary converter, and
a local renderer remains available for modern ``.docx`` and ``.xlsx`` files.
"""

from __future__ import annotations

import base64
import logging
import os
import shutil
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape as xml_escape

from utils.file_uploads import clean_original_filename


logger = logging.getLogger(__name__)


OFFICE_PREVIEW_EXTENSIONS = frozenset({
    "doc", "docx", "docm", "xls", "xlsx", "xlsm",
})
OFFICE_PREVIEW_MIME_TYPES = frozenset({
    "application/msword",
    "application/vnd.ms-excel",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
})


class OfficePreviewError(RuntimeError):
    """Raised when an Office attachment cannot be converted to PDF."""


def _normalized_mime(mimetype: str | None) -> str:
    return (mimetype or "").split(";", 1)[0].strip().lower()


def office_extension(filename: str | None) -> str:
    name = clean_original_filename(filename)
    return Path(name).suffix.lower().lstrip(".")


def is_office_previewable(
    filename: str | None,
    mimetype: str | None = None,
) -> bool:
    """Return whether an attachment is a supported Word or Excel document."""
    return (
        office_extension(filename) in OFFICE_PREVIEW_EXTENSIONS
        or _normalized_mime(mimetype) in OFFICE_PREVIEW_MIME_TYPES
    )


def office_pdf_filename(filename: str | None) -> str:
    """Build a safe, user-friendly filename for the generated preview PDF."""
    name = clean_original_filename(filename)
    stem = Path(name).stem or "document"
    return f"{stem}.pdf"


def _configured_executable(value: str | None) -> str | None:
    candidate = str(value or "").strip().strip('"')
    if not candidate:
        return None
    expanded = os.path.expandvars(candidate)
    if os.path.isfile(expanded):
        return expanded
    return shutil.which(expanded)


def find_libreoffice_executable() -> str | None:
    """Find LibreOffice without invoking a shell or trusting user input."""
    candidates = [
        os.getenv("OFFICE_PREVIEW_LIBREOFFICE_CMD"),
        os.getenv("CORR_INTAKE_LIBREOFFICE_CMD"),
        shutil.which("soffice"),
        shutil.which("libreoffice"),
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    for candidate in candidates:
        executable = _configured_executable(candidate)
        if executable:
            return executable
    return None


def _office_application_available(extension: str) -> bool:
    """Return whether the locally installed Microsoft Office app can export PDF."""
    if os.name != "nt":
        return False

    executable = "EXCEL.EXE" if extension in {"xls", "xlsx", "xlsm"} else "WINWORD.EXE"
    program_files = [
        Path(os.environ.get("ProgramFiles", r"C:\Program Files")),
        Path(os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")),
    ]
    for root in program_files:
        for relative in (
            Path("Microsoft Office") / "Office16" / executable,
            Path("Microsoft Office") / "root" / "Office16" / executable,
        ):
            if (root / relative).is_file():
                return True
    return False


def _powershell_executable() -> str:
    """Return an absolute PowerShell path when Windows runs the web service.

    Windows services frequently have a reduced PATH.  Using SystemRoot avoids
    silently falling back to the simplified ReportLab renderer for that reason.
    """
    if os.name != "nt":
        return "powershell.exe"

    system_root = Path(os.environ.get("SystemRoot", r"C:\\Windows"))
    system_powershell = system_root / "System32" / "WindowsPowerShell" / "v1.0" / "powershell.exe"
    if system_powershell.is_file():
        return str(system_powershell)
    return shutil.which("powershell.exe") or "powershell.exe"


def _powershell_literal(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def _convert_with_windows_office(
    source_path: str | os.PathLike[str],
    original_name: str | None,
) -> bytes:
    """Export with locally installed Word/Excel, preserving Office formatting."""
    extension = office_extension(original_name) or Path(source_path).suffix.lower().lstrip(".")
    application = "Excel" if extension in {"xls", "xlsx", "xlsm"} else "Word"
    with tempfile.TemporaryDirectory(prefix="office_native_preview_") as temp_dir:
        temp_root = Path(temp_dir)
        input_path = temp_root / f"document.{extension}"
        output_path = temp_root / "document.pdf"
        shutil.copyfile(source_path, input_path)
        source_literal = _powershell_literal(str(input_path))
        output_literal = _powershell_literal(str(output_path))

        if application == "Excel":
            script = f"""
$ErrorActionPreference = 'Stop'
$excel = $null
$workbook = $null
try {{
  $excel = New-Object -ComObject Excel.Application
  $excel.Visible = $false
  $excel.DisplayAlerts = $false
  try {{ $excel.AutomationSecurity = 3 }} catch {{ }}
  $workbook = $excel.Workbooks.Open({source_literal}, 0, $true)
  $workbook.ExportAsFixedFormat(0, {output_literal}, 0, $true, $false)
}} finally {{
  if ($workbook) {{ $workbook.Close($false); [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($workbook) }}
  if ($excel) {{ $excel.Quit(); [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($excel) }}
}}
"""
        else:
            script = f"""
$ErrorActionPreference = 'Stop'
$word = $null
$document = $null
try {{
  $word = New-Object -ComObject Word.Application
  $word.Visible = $false
  $word.DisplayAlerts = 0
  try {{ $word.AutomationSecurity = 3 }} catch {{ }}
  $document = $word.Documents.Open({source_literal}, $false, $true, $false)
  $document.ExportAsFixedFormat({output_literal}, 17)
}} finally {{
  if ($document) {{ $document.Close($false); [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($document) }}
  if ($word) {{ $word.Quit(); [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($word) }}
}}
"""

        encoded_script = base64.b64encode(script.encode("utf-16-le")).decode("ascii")
        command = [
            _powershell_executable(),
            "-NoLogo",
            "-NoProfile",
            "-NonInteractive",
            "-EncodedCommand",
            encoded_script,
        ]
        try:
            completed = subprocess.run(
                command,
                capture_output=True,
                text=True,
                timeout=120,
                check=False,
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            )
        except (OSError, subprocess.SubprocessError) as exc:
            raise OfficePreviewError("تعذر تشغيل محوّل Microsoft Office") from exc

        if completed.returncode != 0 or not output_path.exists():
            detail = (completed.stderr or completed.stdout or "").strip()[-500:]
            message = f"فشل تحويل الملف بواسطة {application}"
            raise OfficePreviewError(f"{message}: {detail}" if detail else message)

        pdf_bytes = output_path.read_bytes()
        if not pdf_bytes.startswith(b"%PDF"):
            raise OfficePreviewError("نتيجة تحويل Microsoft Office ليست ملف PDF صالحاً")
        return pdf_bytes


def _convert_with_libreoffice(
    source_path: str | os.PathLike[str],
    original_name: str | None,
    executable: str,
) -> bytes:
    suffix = Path(clean_original_filename(original_name)).suffix.lower()
    if suffix not in {".doc", ".docx", ".docm", ".xls", ".xlsx", ".xlsm"}:
        suffix = Path(source_path).suffix.lower()

    with tempfile.TemporaryDirectory(prefix="office_preview_") as temp_dir:
        temp_root = Path(temp_dir)
        input_path = temp_root / f"document{suffix}"
        output_path = temp_root / "document.pdf"
        profile_path = temp_root / "libreoffice-profile"
        shutil.copyfile(source_path, input_path)

        command = [
            executable,
            "--headless",
            f"-env:UserInstallation={profile_path.resolve().as_uri()}",
            "--convert-to",
            "pdf",
            "--outdir",
            str(temp_root),
            str(input_path),
        ]
        run_kwargs = {
            "capture_output": True,
            "text": True,
            "timeout": 120,
            "check": False,
        }
        if os.name == "nt":
            run_kwargs["creationflags"] = getattr(subprocess, "CREATE_NO_WINDOW", 0)

        try:
            completed = subprocess.run(command, **run_kwargs)
        except (OSError, subprocess.SubprocessError) as exc:
            raise OfficePreviewError("تعذر تشغيل LibreOffice لتحويل الملف إلى PDF") from exc

        if completed.returncode != 0 or not output_path.exists():
            detail = (completed.stderr or completed.stdout or "").strip()
            detail = detail[-500:] if detail else ""
            message = "فشل تحويل ملف Office إلى PDF"
            if detail:
                message = f"{message}: {detail}"
            raise OfficePreviewError(message)

        pdf_bytes = output_path.read_bytes()
        if not pdf_bytes.startswith(b"%PDF"):
            raise OfficePreviewError("نتيجة تحويل Office ليست ملف PDF صالحاً")
        return pdf_bytes


def _shape_pdf_text(value: object) -> str:
    """Shape Arabic text for ReportLab while keeping its markup escaped."""
    raw = "" if value is None else str(value)
    try:
        import arabic_reshaper
        from bidi.algorithm import get_display

        shaped_lines = []
        for line in raw.splitlines() or [""]:
            candidate = line
            if any("\u0600" <= char <= "\u06ff" for char in candidate):
                candidate = get_display(arabic_reshaper.reshape(candidate))
            shaped_lines.append(xml_escape(candidate) or " ")
        return "<br/>".join(shaped_lines)
    except Exception:
        return xml_escape(raw).replace("\n", "<br/>") or " "


def _preview_pdf_fonts(pdfmetrics, ttfont) -> tuple[str, str]:
    project_root = Path(__file__).resolve().parents[1]
    windows_fonts = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
    candidates = [
        windows_fonts / "majalla.ttf",
        project_root / "assets" / "fonts" / "DejaVuSans.ttf",
        windows_fonts / "arial.ttf",
        windows_fonts / "tahoma.ttf",
    ]
    bold_candidates = [
        windows_fonts / "majallab.ttf",
        project_root / "assets" / "fonts" / "DejaVuSans-Bold.ttf",
        windows_fonts / "arialbd.ttf",
        windows_fonts / "tahomabd.ttf",
    ]

    regular_name = "OfficePreviewRegular"
    bold_name = "OfficePreviewBold"
    registered = set(pdfmetrics.getRegisteredFontNames())
    if regular_name not in registered:
        for path in candidates:
            if path.is_file():
                try:
                    pdfmetrics.registerFont(ttfont(regular_name, str(path)))
                    break
                except Exception:
                    continue
    registered = set(pdfmetrics.getRegisteredFontNames())
    if bold_name not in registered:
        for path in bold_candidates:
            if path.is_file():
                try:
                    pdfmetrics.registerFont(ttfont(bold_name, str(path)))
                    break
                except Exception:
                    continue
    registered = set(pdfmetrics.getRegisteredFontNames())
    regular = regular_name if regular_name in registered else "Helvetica"
    bold = bold_name if bold_name in registered else regular
    return regular, bold


def _docx_reportlab_blocks(path: str | os.PathLike[str]) -> list[tuple[str, object]]:
    try:
        from docx import Document
        document = Document(str(path))
    except ImportError as exc:
        raise OfficePreviewError("مكتبة قراءة ملفات Word غير متوفرة") from exc
    except Exception as exc:
        raise OfficePreviewError("تعذر قراءة ملف Word") from exc

    blocks: list[tuple[str, object]] = []
    for paragraph in document.paragraphs:
        text = paragraph.text or ""
        if text.strip():
            style_name = str(getattr(getattr(paragraph, "style", None), "name", "") or "").lower()
            blocks.append(("heading" if "heading" in style_name or "title" in style_name else "text", text))

    for table in document.tables:
        rows = [[cell.text or "" for cell in row.cells] for row in table.rows]
        rows = [row for row in rows if any(str(value).strip() for value in row)]
        if rows:
            blocks.append(("table", rows))
    return blocks


def _xlsx_reportlab_blocks(path: str | os.PathLike[str]) -> list[tuple[str, object]]:
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(str(path), read_only=True, data_only=False)
    except ImportError as exc:
        raise OfficePreviewError("مكتبة قراءة ملفات Excel غير متوفرة") from exc
    except Exception as exc:
        raise OfficePreviewError("تعذر قراءة ملف Excel") from exc

    blocks: list[tuple[str, object]] = []
    try:
        for sheet in workbook.worksheets:
            blocks.append(("heading", str(sheet.title)))
            max_rows = min(int(sheet.max_row or 0), 500)
            max_columns = min(int(sheet.max_column or 0), 20)
            rows = []
            if max_rows and max_columns:
                for row in sheet.iter_rows(
                    min_row=1,
                    max_row=max_rows,
                    max_col=max_columns,
                    values_only=True,
                ):
                    values = list(row)
                    while values and values[-1] in (None, ""):
                        values.pop()
                    if values:
                        rows.append(values)
            if rows:
                blocks.append(("table", rows))
            else:
                blocks.append(("text", "لا توجد بيانات قابلة للعرض في ورقة العمل."))
    finally:
        workbook.close()
    return blocks


def _convert_with_reportlab_fallback(
    source_path: str | os.PathLike[str],
    original_name: str | None,
) -> bytes:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER, TA_RIGHT
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:
        raise OfficePreviewError("مكتبة إنشاء PDF غير متوفرة") from exc

    ext = office_extension(original_name) or Path(source_path).suffix.lower().lstrip(".")
    title = clean_original_filename(original_name) or Path(source_path).name
    if ext in {"docx", "docm"}:
        blocks = _docx_reportlab_blocks(source_path)
        page_size = A4
        note = "تم إنشاء هذه المعاينة محلياً؛ قد لا تتطابق بعض التفاصيل البصرية مع Word."
    elif ext in {"xlsx", "xlsm"}:
        blocks = _xlsx_reportlab_blocks(source_path)
        page_size = landscape(A4)
        note = "تم إنشاء هذه المعاينة محلياً؛ قد لا تتطابق بعض الصيغ والتنسيقات مع Excel."
    else:
        raise OfficePreviewError(
            "يتطلب عرض ملفات Word/Excel القديمة تثبيت LibreOffice على الخادم"
        )

    regular_font, bold_font = _preview_pdf_fonts(pdfmetrics, TTFont)
    body_style = ParagraphStyle(
        "OfficePreviewBody",
        fontName=regular_font,
        fontSize=16,
        leading=21,
        alignment=TA_RIGHT,
        wordWrap="RTL",
        spaceAfter=5,
    )
    heading_style = ParagraphStyle(
        "OfficePreviewHeading",
        parent=body_style,
        fontName=bold_font,
        fontSize=18,
        leading=24,
        spaceBefore=8,
        spaceAfter=6,
    )
    title_style = ParagraphStyle(
        "OfficePreviewTitle",
        parent=body_style,
        fontName=bold_font,
        fontSize=20,
        leading=26,
        alignment=TA_CENTER,
        spaceAfter=8,
    )
    note_style = ParagraphStyle(
        "OfficePreviewNote",
        parent=body_style,
        fontSize=16,
        leading=21,
        textColor=colors.HexColor("#1e40af"),
        backColor=colors.HexColor("#eff6ff"),
        borderColor=colors.HexColor("#bfdbfe"),
        borderWidth=0.5,
        borderPadding=5,
        spaceAfter=10,
    )
    cell_style = ParagraphStyle(
        "OfficePreviewCell",
        parent=body_style,
        fontSize=16,
        leading=21,
        spaceAfter=0,
    )

    stream = BytesIO()
    margin = 12 * mm
    document = SimpleDocTemplate(
        stream,
        pagesize=page_size,
        rightMargin=margin,
        leftMargin=margin,
        topMargin=margin,
        bottomMargin=margin,
    )
    story = [
        Paragraph(_shape_pdf_text(title), title_style),
        Paragraph(_shape_pdf_text(note), note_style),
    ]
    content_width = page_size[0] - (2 * margin)
    saw_content = False
    for kind, value in blocks:
        if kind == "heading":
            if saw_content and ext in {"xlsx", "xlsm"}:
                story.append(Spacer(1, 8 * mm))
            story.append(Paragraph(_shape_pdf_text(value), heading_style))
            saw_content = True
        elif kind == "text":
            story.append(Paragraph(_shape_pdf_text(value), body_style))
            saw_content = True
        elif kind == "table":
            rows = list(value or [])
            if not rows:
                continue
            column_count = max(len(row) for row in rows)
            normalized_rows = [
                list(row) + [""] * (column_count - len(row))
                for row in rows
            ]
            table_data = [
                [Paragraph(_shape_pdf_text(cell), cell_style) for cell in row]
                for row in normalized_rows
            ]
            table = Table(
                table_data,
                colWidths=[content_width / column_count] * column_count,
                repeatRows=1,
                hAlign="RIGHT",
            )
            table.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ALIGN", (0, 0), (-1, -1), "RIGHT"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            story.extend([table, Spacer(1, 6 * mm)])
            saw_content = True

    if not saw_content:
        story.append(Paragraph(_shape_pdf_text("لا يوجد محتوى قابل للعرض في هذا الملف."), body_style))

    try:
        document.build(story)
    except Exception as exc:
        raise OfficePreviewError("تعذر إنشاء PDF للمعاينة المحلية") from exc
    pdf_bytes = stream.getvalue()
    if not pdf_bytes.startswith(b"%PDF"):
        raise OfficePreviewError("نتيجة المعاينة المحلية ليست ملف PDF صالحاً")
    return pdf_bytes


def _convert_with_python_fallback(
    source_path: str | os.PathLike[str],
    original_name: str | None,
) -> bytes:
    ext = office_extension(original_name) or Path(source_path).suffix.lower().lstrip(".")
    if ext in {"docx", "docm", "xlsx", "xlsm"}:
        return _convert_with_reportlab_fallback(source_path, original_name)
    raise OfficePreviewError(
        "يتطلب عرض ملفات Word/Excel القديمة تثبيت LibreOffice على الخادم"
    )


def convert_office_to_pdf(
    source_path: str | os.PathLike[str],
    *,
    original_name: str | None = None,
) -> bytes:
    """Return a PDF representation of a Word or Excel attachment."""
    if not (
        is_office_previewable(original_name)
        or is_office_previewable(str(source_path))
    ):
        raise OfficePreviewError("نوع الملف غير مدعوم للمعاينة بصيغة PDF")

    extension = office_extension(original_name) or Path(source_path).suffix.lower().lstrip(".")
    # Do not use the executable-path check as a gate.  Click-to-Run and
    # 32-bit Office installations can register Word/Excel COM correctly while
    # keeping their executable outside the two historical Office16 paths.
    # Attempting COM is the authoritative check and preserves the original
    # document formatting whenever Word/Excel is available to this process.
    if os.name == "nt":
        try:
            return _convert_with_windows_office(source_path, original_name)
        except OfficePreviewError as exc:
            # Service accounts without interactive Office automation can still
            # use LibreOffice or the local PDF fallback below.
            logger.warning("Native %s PDF preview unavailable; using fallback: %s", extension, exc)

    executable = find_libreoffice_executable()
    if executable:
        try:
            return _convert_with_libreoffice(source_path, original_name, executable)
        except OfficePreviewError:
            # A broken/partial LibreOffice installation should not make modern
            # DOCX/XLSX attachments completely unpreviewable.
            ext = office_extension(original_name) or Path(source_path).suffix.lower().lstrip(".")
            if ext not in {"docx", "docm", "xlsx", "xlsm"}:
                raise

    return _convert_with_python_fallback(source_path, original_name)


def convert_office_bytes_to_pdf(
    payload: bytes,
    *,
    original_name: str,
) -> bytes:
    """Convert an Office payload held in memory without keeping a copy."""
    if not is_office_previewable(original_name):
        raise OfficePreviewError("نوع الملف غير مدعوم للمعاينة بصيغة PDF")

    suffix = Path(clean_original_filename(original_name)).suffix.lower()
    if not suffix:
        raise OfficePreviewError("لا يمكن تحديد امتداد ملف Office")

    with tempfile.TemporaryDirectory(prefix="office_preview_bytes_") as temp_dir:
        source_path = Path(temp_dir) / f"document{suffix}"
        source_path.write_bytes(payload)
        return convert_office_to_pdf(source_path, original_name=original_name)


__all__ = [
    "OfficePreviewError",
    "convert_office_to_pdf",
    "convert_office_bytes_to_pdf",
    "find_libreoffice_executable",
    "is_office_previewable",
    "office_pdf_filename",
]
