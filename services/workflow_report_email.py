"""Scheduled and on-demand email delivery for Masar workflow reports."""

from __future__ import annotations

from datetime import datetime, time as dt_time
from email.message import EmailMessage
from email.utils import formataddr
from html import escape
from io import BytesIO
import json
import smtplib
import ssl
import uuid
from typing import Any

from flask import has_app_context

from extensions import db
from models import SystemSetting, User, WorkflowReportEmailDelivery
from services.attendance_report_email import (
    _as_bool,
    _parse_date,
    _parse_date_list,
    _parse_email_list,
    _parse_hhmm,
    _parse_hhmm_list,
    _parse_int_list,
    _schedule_is_due,
    _send_day_is_excluded,
)
from services.delivery_controls import email_delivery_enabled
from services.workflow_task_email import _mail_config, resolve_user_delivery_email
from utils.timezone import app_timezone, format_local_datetime
from utils.ui_labels import workflow_status_label
REPORT_KEY = "EMPLOYEE_WORKFLOW"
REPORT_NAME = "تقرير موظفي ومسارات مسار"
REPORT_FROM_NAME = "مسار"

SETTING_PREFIX = "WORKFLOW_REPORTS_EMAIL_"
SETTING_KEYS = {
    "enabled": SETTING_PREFIX + "ENABLED",
    "send_times": SETTING_PREFIX + "TIMES",
    "send_time": SETTING_PREFIX + "TIME",
    "frequency": SETTING_PREFIX + "FREQUENCY",
    "interval_days": SETTING_PREFIX + "INTERVAL_DAYS",
    "weekdays": SETTING_PREFIX + "WEEKDAYS",
    "skip_weekly_holidays": SETTING_PREFIX + "SKIP_WEEKLY_HOLIDAYS",
    "skip_official_holidays": SETTING_PREFIX + "SKIP_OFFICIAL_HOLIDAYS",
    "excluded_dates": SETTING_PREFIX + "EXCLUDED_DATES",
    "recipient_user_ids": SETTING_PREFIX + "RECIPIENT_USER_IDS",
    "recipient_emails": SETTING_PREFIX + "RECIPIENT_EMAILS",
    "start_date": SETTING_PREFIX + "START_DATE",
    "include_details": SETTING_PREFIX + "INCLUDE_DETAILS",
    "viewer_user_id": SETTING_PREFIX + "VIEWER_USER_ID",
}

SCHEDULER_STATE_KEYS = {
    "last_check_at": SETTING_PREFIX + "LAST_CHECK_AT",
    "last_status": SETTING_PREFIX + "LAST_STATUS",
    "last_error": SETTING_PREFIX + "LAST_ERROR",
    "last_due_times": SETTING_PREFIX + "LAST_DUE_TIMES",
}

DEFAULT_CONFIG = {
    "enabled": "0",
    "send_times": "10:00",
    "send_time": "10:00",
    "frequency": "DAILY",
    "interval_days": "1",
    "weekdays": "0,1,2,3,4,5,6",
    "skip_weekly_holidays": "0",
    "skip_official_holidays": "0",
    "excluded_dates": "",
    "recipient_user_ids": "[]",
    "recipient_emails": "",
    "start_date": "",
    "include_details": "1",
    "viewer_user_id": "",
}


def _setting(key: str, default: str | None = None) -> str | None:
    row = SystemSetting.query.filter_by(key=key).first()
    if row is None or row.value in (None, ""):
        return default
    return str(row.value)


def _set_setting(key: str, value: str | None) -> None:
    row = SystemSetting.query.filter_by(key=key).first()
    if row is None:
        db.session.add(SystemSetting(key=key, value=value or ""))
    else:
        row.value = value or ""


def _parse_optional_int(value: Any) -> int | None:
    try:
        parsed = int(str(value or "").strip())
    except (TypeError, ValueError):
        return None
    return parsed if parsed > 0 else None


def _can_build_report(user: User | None) -> bool:
    if not user:
        return False
    return user.has_perm("WORKFLOW_REPORTS_READ") or user.has_perm("WORKFLOW_REPORTS_EMAIL_MANAGE")


def get_report_email_config() -> dict[str, Any]:
    values = {
        name: _setting(key, DEFAULT_CONFIG[name])
        for name, key in SETTING_KEYS.items()
    }
    values["enabled"] = _as_bool(values["enabled"])
    values["skip_weekly_holidays"] = _as_bool(values["skip_weekly_holidays"])
    values["skip_official_holidays"] = _as_bool(values["skip_official_holidays"])
    values["include_details"] = _as_bool(values["include_details"])

    frequency = str(values["frequency"] or "DAILY").strip().upper()
    if frequency not in {"DAILY", "EVERY_N_DAYS", "WEEKLY"}:
        frequency = "DAILY"
    values["frequency"] = frequency

    try:
        interval_days = max(1, min(int(str(values["interval_days"] or "1")), 365))
    except (TypeError, ValueError):
        interval_days = 1
    values["interval_days"] = interval_days

    send_times = _parse_hhmm_list(values.get("send_times"))
    if not send_times:
        legacy_send_time = _parse_hhmm(values.get("send_time"))
        send_times = [legacy_send_time] if legacy_send_time else [dt_time(10, 0)]
    values["send_times_list"] = [value.strftime("%H:%M") for value in send_times]
    values["send_times"] = "\n".join(values["send_times_list"])
    values["send_time"] = values["send_times_list"][0]
    values["weekdays_list"] = _parse_int_list(values["weekdays"], 0, 6)
    values["excluded_dates_list"] = _parse_date_list(values["excluded_dates"])
    values["recipient_user_ids_list"] = _parse_int_list(values["recipient_user_ids"], 1, None)
    values["recipient_emails_list"] = _parse_email_list(values["recipient_emails"])
    values["start_date"] = _parse_date(values["start_date"])
    values["viewer_user_id"] = _parse_optional_int(values["viewer_user_id"])
    values["scheduler_last_check_at"] = _setting(SCHEDULER_STATE_KEYS["last_check_at"], "") or ""
    values["scheduler_last_status"] = _setting(SCHEDULER_STATE_KEYS["last_status"], "") or ""
    values["scheduler_last_error"] = _setting(SCHEDULER_STATE_KEYS["last_error"], "") or ""
    values["scheduler_last_due_times"] = _setting(SCHEDULER_STATE_KEYS["last_due_times"], "") or ""
    return values


def save_report_email_config(payload: dict[str, Any]) -> dict[str, Any]:
    send_times = _parse_hhmm_list(payload.get("send_times"))
    if not send_times:
        send_times = _parse_hhmm_list(payload.get("send_time"))
    if not send_times:
        send_times = [dt_time(10, 0)]

    frequency = str(payload.get("frequency") or "DAILY").strip().upper()
    if frequency not in {"DAILY", "EVERY_N_DAYS", "WEEKLY"}:
        frequency = "DAILY"
    try:
        interval_days = max(1, min(int(payload.get("interval_days") or 1), 365))
    except (TypeError, ValueError):
        interval_days = 1

    weekdays = sorted(set(_parse_int_list(payload.get("weekdays"), 0, 6)))
    if frequency == "WEEKLY" and not weekdays:
        weekdays = [0, 1, 2, 3, 4]
    excluded_dates = sorted(set(_parse_date_list(payload.get("excluded_dates"))))
    start_date = _parse_date(payload.get("start_date"))
    user_ids = sorted(set(_parse_int_list(payload.get("recipient_user_ids"), 1, None)))
    if user_ids:
        valid_users = {
            int(user.id)
            for user in User.query.filter(User.id.in_(user_ids)).all()
            if resolve_user_delivery_email(user)
        }
        user_ids = [user_id for user_id in user_ids if user_id in valid_users]
    emails = _parse_email_list(payload.get("recipient_emails"))
    viewer_user_id = _parse_optional_int(payload.get("viewer_user_id"))
    viewer = db.session.get(User, viewer_user_id) if viewer_user_id else None
    if not _can_build_report(viewer):
        viewer_user_id = None

    values = {
        "enabled": "1" if _as_bool(payload.get("enabled")) else "0",
        "send_times": "\n".join(value.strftime("%H:%M") for value in send_times),
        "send_time": send_times[0].strftime("%H:%M"),
        "frequency": frequency,
        "interval_days": str(interval_days),
        "weekdays": ",".join(str(value) for value in weekdays),
        "skip_weekly_holidays": "1" if _as_bool(payload.get("skip_weekly_holidays")) else "0",
        "skip_official_holidays": "1" if _as_bool(payload.get("skip_official_holidays")) else "0",
        "excluded_dates": "\n".join(excluded_dates),
        "recipient_user_ids": json.dumps(user_ids, ensure_ascii=False),
        "recipient_emails": "\n".join(emails),
        "start_date": start_date.isoformat() if start_date else "",
        "include_details": "1" if _as_bool(payload.get("include_details")) else "0",
        "viewer_user_id": str(viewer_user_id or ""),
    }
    for name, key in SETTING_KEYS.items():
        _set_setting(key, values[name])
    db.session.commit()
    return get_report_email_config()


def _configured_send_times(config: dict[str, Any]) -> list[dt_time]:
    values = _parse_hhmm_list(config.get("send_times_list") or config.get("send_times"))
    return values or [_parse_hhmm(config.get("send_time")) or dt_time(10, 0)]


def _report_viewer(config: dict[str, Any], preferred: User | None = None) -> User:
    candidates = []
    if preferred:
        candidates.append(preferred)
    configured_id = config.get("viewer_user_id")
    if configured_id:
        configured = db.session.get(User, int(configured_id))
        if configured:
            candidates.append(configured)

    all_users = User.query.order_by(User.id.asc()).all()
    candidates.extend(all_users)
    seen = set()
    for user in candidates:
        if not user or int(user.id) in seen:
            continue
        seen.add(int(user.id))
        try:
            if _can_build_report(user):
                return user
        except Exception:
            continue
    raise RuntimeError("لا يوجد مستخدم مخول لقراءة تقرير مسار وإرساله.")


def _recipients(config: dict[str, Any]) -> list[str]:
    result = list(config.get("recipient_emails_list") or [])
    user_ids = config.get("recipient_user_ids_list") or []
    if user_ids:
        users = User.query.filter(User.id.in_(user_ids)).all()
        result.extend(resolve_user_delivery_email(user) for user in users)
    return sorted({email.strip() for email in result if email}, key=str.casefold)


def _datetime_text(value) -> str:
    return format_local_datetime(value) if value else ""


def _summary_rows(report: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for row in report.get("summary_rows") or []:
        user = row["user"]
        status_text = "، ".join(
            f"{workflow_status_label(status)}: {count}"
            for status, count in sorted((row.get("status_counts") or {}).items())
        )
        rows.append({
            "employee": getattr(user, "full_name", None) or getattr(user, "email", None) or f"المستخدم #{user.id}",
            "job_title": getattr(user, "job_title", None) or getattr(user, "role", None) or "",
            "department": getattr(row.get("department"), "name_ar", None) or "",
            "email": getattr(user, "email", None) or getattr(user, "username", None) or "",
            "total_paths": row.get("total_paths") or 0,
            "active_paths": row.get("active_paths") or 0,
            "statuses": status_text,
            "last_assigned_at": _datetime_text(row.get("last_assigned_at")),
            "last_opened_at": _datetime_text(row.get("last_opened_at")),
        })
    return rows


def _detail_rows(report: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for detail in report.get("report_details") or []:
        user = detail.get("user")
        request_row = detail["request"]
        requester = getattr(request_row, "requester", None)
        rows.append({
            "employee": getattr(user, "full_name", None) or getattr(user, "email", None) or f"المستخدم #{detail['user_id']}",
            "request_id": request_row.id,
            "path": getattr(request_row, "title", None) or "",
            "status": workflow_status_label(getattr(request_row, "status", None)),
            "requester": getattr(requester, "full_name", None) if requester else "",
            "steps": detail.get("step_orders_label") or "",
            "active": "نعم" if detail.get("active") else "لا",
            "created_at": _datetime_text(getattr(request_row, "created_at", None)),
            "assignment_at": _datetime_text(detail.get("assignment_at")),
            "opened_at": _datetime_text(detail.get("opened_at")),
        })
    return rows


def build_workflow_report_data(viewer: User, *, include_details: bool = True) -> dict[str, Any]:
    # Import lazily so the standalone background worker can load this service
    # without initializing workflow.routes while that blueprint is importing.
    from workflow.reports import build_employee_workflow_report

    report = build_employee_workflow_report(
        viewer,
        page=1,
        per_page=None,
        include_all_details=include_details,
    )
    return {
        "summary": _summary_rows(report),
        "details": _detail_rows(report) if include_details else [],
        "total_paths": report.get("total_paths") or 0,
        "total_assignments": report.get("total_assignments") or 0,
        "active_assignments": report.get("active_assignments") or 0,
        "opened_assignments": report.get("opened_assignments") or 0,
    }


def _html_table(rows: list[dict[str, Any]], columns: list[tuple[str, str]]) -> str:
    if not rows:
        return '<p style="margin:0;color:#6b7280">لا توجد سجلات ضمن هذا التقرير.</p>'
    header = "".join(
        f'<th style="padding:8px;border:1px solid #e5e7eb;background:#f8fafc">{escape(label)}</th>'
        for _key, label in columns
    )
    body = []
    for row in rows:
        cells = "".join(
            f'<td style="padding:7px;border:1px solid #e5e7eb">{escape(str(row.get(key, "") or ""))}</td>'
            for key, _label in columns
        )
        body.append(f"<tr>{cells}</tr>")
    return (
        '<table style="width:100%;border-collapse:collapse;font-size:13px">'
        f"<thead><tr>{header}</tr></thead><tbody>{''.join(body)}</tbody></table>"
    )


def build_email_content(data: dict[str, Any]) -> tuple[str, str, str]:
    scheduled_time = str(data.get("scheduled_time") or "").strip()
    subject = REPORT_NAME
    if scheduled_time:
        subject += f" — {scheduled_time}"
    summary_columns = [
        ("employee", "الموظف"),
        ("job_title", "المسمى/الدور"),
        ("department", "الدائرة"),
        ("total_paths", "كل المسارات"),
        ("active_paths", "بانتظار إجراء"),
        ("statuses", "الحالات"),
        ("last_assigned_at", "آخر تكليف"),
        ("last_opened_at", "آخر فتح"),
    ]
    detail_columns = [
        ("employee", "الموظف"),
        ("request_id", "#"),
        ("path", "المسار"),
        ("status", "الحالة"),
        ("requester", "مقدم الطلب"),
        ("steps", "الخطوات"),
        ("active", "نشط"),
        ("created_at", "أنشئ في"),
        ("assignment_at", "وقت التكليف"),
        ("opened_at", "وقت الفتح"),
    ]
    summary = data.get("summary") or []
    details = data.get("details") or []
    sections = [
        ("ملخص المسارات حسب الموظف", "#2F5597", summary, summary_columns),
    ]
    if data.get("include_details"):
        sections.append(("تفاصيل المسارات", "#2F7D6D", details, detail_columns))

    html_sections = []
    text_sections = []
    for title, color, rows, columns in sections:
        html_sections.append(
            f'<section style="margin:0 0 22px;border:1px solid #e5e7eb;border-radius:8px;overflow:hidden">'
            f'<h3 style="margin:0;padding:10px 14px;background:{color};color:#fff;font-size:16px">{escape(title)} ({len(rows)})</h3>'
            f'<div style="padding:12px">{_html_table(rows, columns)}</div></section>'
        )
        text_sections.append(
            title + f" ({len(rows)})\n"
            + ("\n".join(" | ".join(str(row.get(key, "") or "") for key, _ in columns) for row in rows) or "لا توجد سجلات")
        )

    html_body = (
        '<html><body dir="rtl" style="font-family:Arial,sans-serif;color:#1f2937;line-height:1.7">'
        f'<div style="max-width:1400px;margin:auto"><h2 style="margin-bottom:4px;color:#5a4a86">{escape(REPORT_NAME)}</h2>'
        f'<p style="margin-top:0;color:#6b7280">تاريخ ووقت إنشاء التقرير: {escape(_datetime_text(datetime.utcnow()))}'
        + (f" — وقت الإرسال: {escape(scheduled_time)}" if scheduled_time else "")
        + "</p>"
        + "".join(html_sections)
        + '<p style="font-size:12px;color:#6b7280">هذه رسالة آلية من نظام مسار.</p></div></body></html>'
    )
    text_body = (
        f"{REPORT_NAME}\nوقت إنشاء التقرير: {_datetime_text(datetime.utcnow())}\n\n"
        + "\n\n".join(text_sections)
        + "\n\nهذه رسالة آلية من نظام مسار."
    )
    return subject, text_body, html_body


def _xlsx_bytes(data: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()

    def add_sheet(title: str, headers: list[str], rows: list[list[Any]], color: str) -> None:
        sheet = workbook.active if workbook.active.title == "Sheet" else workbook.create_sheet()
        sheet.title = title
        sheet.sheet_view.rightToLeft = True
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        for cell in sheet[1]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=color)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in sheet.columns:
            letter = column[0].column_letter
            width = max(12, min(42, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[letter].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    summary = data.get("summary") or []
    add_sheet(
        "ملخص الموظفين",
        ["الموظف", "المسمى/الدور", "الدائرة", "كل المسارات", "بانتظار إجراء", "الحالات", "آخر تكليف", "آخر فتح"],
        [
            [row.get(key, "") for key in ("employee", "job_title", "department", "total_paths", "active_paths", "statuses", "last_assigned_at", "last_opened_at")]
            for row in summary
        ],
        "2F5597",
    )
    if data.get("include_details"):
        details = data.get("details") or []
        add_sheet(
            "تفاصيل المسارات",
            ["الموظف", "#", "المسار", "الحالة", "مقدم الطلب", "الخطوات", "نشط", "أنشئ في", "وقت التكليف", "وقت الفتح"],
            [
                [row.get(key, "") for key in ("employee", "request_id", "path", "status", "requester", "steps", "active", "created_at", "assignment_at", "opened_at")]
                for row in details
            ],
            "2F7D6D",
        )
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _send_email(config: dict[str, Any], recipients: list[str], subject: str, text_body: str, html_body: str, attachment: bytes) -> None:
    message = EmailMessage()
    message["Subject"] = subject
    message["From"] = formataddr((REPORT_FROM_NAME, config["from_email"]))
    message["To"] = ", ".join(recipients)
    if config.get("reply_to"):
        message["Reply-To"] = config["reply_to"]
    message.set_content(text_body)
    message.add_alternative(html_body, subtype="html")
    message.add_attachment(
        attachment,
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="masar_employee_workflow_report.xlsx",
    )
    if config["security"] == "ssl":
        smtp = smtplib.SMTP_SSL(config["host"], config["port"], timeout=20, context=ssl.create_default_context())
    else:
        smtp = smtplib.SMTP(config["host"], config["port"], timeout=20)
    try:
        smtp.ehlo()
        if config["security"] == "starttls":
            smtp.starttls(context=ssl.create_default_context())
            smtp.ehlo()
        if config["username"]:
            smtp.login(config["username"], config["password"])
        smtp.send_message(message, from_addr=config["from_email"], to_addrs=recipients)
    finally:
        try:
            smtp.quit()
        except Exception:
            pass


def _record_scheduler_state(local_now: datetime, status: str, *, due_times: list[dt_time] | None = None, error: str | None = None) -> None:
    if not has_app_context():
        return
    try:
        local_value = local_now.astimezone(app_timezone()) if local_now.tzinfo else local_now.replace(tzinfo=app_timezone())
        values = {
            "last_check_at": local_value.isoformat(timespec="seconds"),
            "last_status": str(status or "UNKNOWN").strip().upper()[:40],
            "last_error": str(error).strip()[:2000] if error else ("" if str(status or "").upper() in {"SENT", "ALREADY_SENT"} else (_setting(SCHEDULER_STATE_KEYS["last_error"], "") or "")),
            "last_due_times": ", ".join(value.strftime("%H:%M") for value in (due_times or [])),
        }
        for name, key in SCHEDULER_STATE_KEYS.items():
            _set_setting(key, values[name])
        db.session.commit()
    except Exception:
        db.session.rollback()


def _slot_key(run_date: str, scheduled_time: str) -> str:
    return f"{REPORT_KEY}:{run_date}:{scheduled_time}"


def _send_slot(local_now: datetime, config: dict[str, Any], scheduled_time: dt_time) -> dict[str, Any]:
    run_date = local_now.date().isoformat()
    scheduled_time_text = scheduled_time.strftime("%H:%M")
    slot_key = _slot_key(run_date, scheduled_time_text)
    delivery = WorkflowReportEmailDelivery.query.filter_by(slot_key=slot_key).first()
    if delivery and delivery.status == "SENT":
        return {"status": "already_sent", "run_date": run_date, "scheduled_time": scheduled_time_text}
    if delivery:
        return {"status": "already_attempted", "run_date": run_date, "scheduled_time": scheduled_time_text, "attempt_status": delivery.status, "error": delivery.last_error or ""}
    delivery = WorkflowReportEmailDelivery(
        report_key=REPORT_KEY,
        slot_key=slot_key,
        run_date=run_date,
        scheduled_time=scheduled_time_text,
        trigger="SCHEDULED",
        status="SENDING",
        attempt_count=1,
    )
    db.session.add(delivery)
    try:
        db.session.commit()
    except Exception:
        db.session.rollback()
        existing = WorkflowReportEmailDelivery.query.filter_by(slot_key=slot_key).first()
        if existing:
            return {"status": "already_sent" if existing.status == "SENT" else "already_attempted", "run_date": run_date, "scheduled_time": scheduled_time_text, "attempt_status": existing.status, "error": existing.last_error or ""}
        raise

    try:
        recipients = _recipients(config)
        if not recipients:
            raise RuntimeError("لم يتم تحديد أي مستلم صالح لتقارير مسار.")
        if not email_delivery_enabled():
            raise RuntimeError("تم تعطيل إرسال البريد الإلكتروني من إعدادات النظام.")
        mail_config = _mail_config()
        if not mail_config.get("ready"):
            raise RuntimeError("إعدادات SMTP غير مكتملة أو البريد غير مفعّل.")
        viewer = _report_viewer(config)
        data = build_workflow_report_data(viewer, include_details=config["include_details"])
        data["include_details"] = config["include_details"]
        data["scheduled_time"] = scheduled_time_text
        subject, text_body, html_body = build_email_content(data)
        _send_email(mail_config, recipients, subject, text_body, html_body, _xlsx_bytes(data))
        delivery.status = "SENT"
        delivery.sent_at = datetime.utcnow()
        delivery.recipient_count = len(recipients)
        delivery.employee_count = len(data["summary"])
        delivery.assignment_count = int(data["total_assignments"] or 0)
        delivery.last_error = None
        delivery.updated_at = datetime.utcnow()
        db.session.commit()
        return {"status": "sent", "run_date": run_date, "scheduled_time": scheduled_time_text, "recipients": len(recipients), "employees": len(data["summary"]), "assignments": data["total_assignments"]}
    except Exception as exc:
        db.session.rollback()
        delivery = WorkflowReportEmailDelivery.query.filter_by(slot_key=slot_key).first()
        if delivery:
            delivery.status = "FAILED"
            delivery.attempt_count = max(1, int(delivery.attempt_count or 0))
            delivery.last_error = str(exc)[:2000]
            delivery.updated_at = datetime.utcnow()
            db.session.commit()
        return {"status": "failed", "run_date": run_date, "scheduled_time": scheduled_time_text, "error": str(exc)[:500]}


def run_workflow_report_email_cycle(now: datetime | None = None, *, force: bool = False) -> dict[str, Any]:
    config = get_report_email_config()
    local_now = now or datetime.now(app_timezone())
    if local_now.tzinfo is None:
        local_now = local_now.replace(tzinfo=app_timezone())
    if not force and not config["enabled"]:
        _record_scheduler_state(local_now, "DISABLED")
        return {"status": "disabled"}

    configured_times = _configured_send_times(config)
    if force:
        due_times = configured_times[:1]
    else:
        current_time = local_now.time().replace(second=0, microsecond=0)
        if not _schedule_is_due(local_now.date(), config) or _send_day_is_excluded(local_now.date(), config):
            _record_scheduler_state(local_now, "EXCLUDED", due_times=configured_times)
            return {"status": "excluded"}
        due_times = [scheduled_time for scheduled_time in configured_times if scheduled_time == current_time]
        if not due_times:
            _record_scheduler_state(local_now, "NOT_DUE", due_times=configured_times)
            return {"status": "not_due"}

    _record_scheduler_state(local_now, "DUE", due_times=due_times)
    results = [_send_slot(local_now, config, scheduled_time) for scheduled_time in due_times]
    if len(results) == 1:
        result = results[0]
        _record_scheduler_state(local_now, str(result.get("status") or "UNKNOWN"), due_times=due_times, error=result.get("error"))
        return result

    sent = [result for result in results if result.get("status") == "sent"]
    failed = [result for result in results if result.get("status") == "failed"]
    if sent:
        status = "sent"
    elif failed:
        status = "failed"
    elif results and all(result.get("status") == "already_sent" for result in results):
        status = "already_sent"
    else:
        status = "already_attempted"
    errors = "; ".join(str(result.get("error") or "").strip() for result in failed if result.get("error"))
    _record_scheduler_state(local_now, status, due_times=due_times, error=errors)
    return {"status": status, "run_date": local_now.date().isoformat(), "scheduled_times": [value.strftime("%H:%M") for value in due_times], "runs": results}


def send_workflow_report_email_now(viewer: User | None = None) -> dict[str, Any]:
    config = get_report_email_config()
    recipients = _recipients(config)
    if not recipients:
        raise RuntimeError("حدد مستلمًا واحدًا على الأقل قبل الإرسال.")
    if not email_delivery_enabled():
        raise RuntimeError("تم تعطيل إرسال البريد الإلكتروني من إعدادات النظام.")
    mail_config = _mail_config()
    if not mail_config.get("ready"):
        raise RuntimeError("إعدادات SMTP غير مكتملة أو البريد غير مفعّل.")

    local_now = datetime.now(app_timezone())
    manual_key = f"{REPORT_KEY}:ON_DEMAND:{uuid.uuid4().hex}"
    delivery = WorkflowReportEmailDelivery(
        report_key=REPORT_KEY,
        slot_key=manual_key,
        run_date=local_now.date().isoformat(),
        scheduled_time="عند الطلب",
        trigger="ON_DEMAND",
        status="SENDING",
        attempt_count=1,
    )
    db.session.add(delivery)
    db.session.commit()
    try:
        report_viewer = _report_viewer(config, preferred=viewer)
        data = build_workflow_report_data(report_viewer, include_details=config["include_details"])
        data["include_details"] = config["include_details"]
        data["scheduled_time"] = "عند الطلب"
        subject, text_body, html_body = build_email_content(data)
        _send_email(mail_config, recipients, f"[إرسال عند الطلب] {subject}", text_body, html_body, _xlsx_bytes(data))
        delivery.status = "SENT"
        delivery.sent_at = datetime.utcnow()
        delivery.recipient_count = len(recipients)
        delivery.employee_count = len(data["summary"])
        delivery.assignment_count = int(data["total_assignments"] or 0)
        delivery.updated_at = datetime.utcnow()
        db.session.commit()
        return {"status": "sent", "recipients": len(recipients), "employees": len(data["summary"]), "assignments": data["total_assignments"]}
    except Exception as exc:
        db.session.rollback()
        delivery = WorkflowReportEmailDelivery.query.filter_by(slot_key=manual_key).first()
        if delivery:
            delivery.status = "FAILED"
            delivery.last_error = str(exc)[:2000]
            delivery.updated_at = datetime.utcnow()
            db.session.commit()
        raise
