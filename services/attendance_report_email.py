"""Scheduled HR attendance reports delivered by email.

The service deliberately keeps the report generation independent from a Flask
request.  It is used by the in-process email worker and by the HR settings
page's test-send action.  All persisted values are stored in ``SystemSetting``
so this feature can be enabled on existing SQLite installations without a
manual migration.
"""

from __future__ import annotations

from datetime import date, datetime, timedelta, time as dt_time
from email.message import EmailMessage
from email.utils import formataddr
from html import escape
import json
import re
import smtplib
import ssl
from io import BytesIO
from typing import Any

from sqlalchemy import func

from extensions import db
from models import (
    AttendanceDailySummary,
    AttendanceEvent,
    EmployeeFile,
    HRAttendanceEmailReportDelivery,
    HRLeaveRequest,
    HROfficialOccasion,
    HROfficialOccasionRange,
    SystemSetting,
    User,
)
from services.delivery_controls import email_delivery_enabled
from services.workflow_task_email import _mail_config, resolve_user_delivery_email
from utils.timezone import app_timezone


REPORT_NAME = "تقارير الدوام للموارد البشرية"
REPORT_FROM_NAME = REPORT_NAME

SETTING_PREFIX = "HR_ATTENDANCE_EMAIL_REPORTS_"
SETTING_KEYS = {
    "enabled": SETTING_PREFIX + "ENABLED",
    "send_times": SETTING_PREFIX + "TIMES",
    # Keep the original key as a compatibility bridge for existing installs.
    "send_time": SETTING_PREFIX + "TIME",
    "frequency": SETTING_PREFIX + "FREQUENCY",
    "interval_days": SETTING_PREFIX + "INTERVAL_DAYS",
    "weekdays": SETTING_PREFIX + "WEEKDAYS",
    "skip_weekly_holidays": SETTING_PREFIX + "SKIP_WEEKLY_HOLIDAYS",
    "skip_official_holidays": SETTING_PREFIX + "SKIP_OFFICIAL_HOLIDAYS",
    "excluded_dates": SETTING_PREFIX + "EXCLUDED_DATES",
    "recipient_user_ids": SETTING_PREFIX + "RECIPIENT_USER_IDS",
    "recipient_emails": SETTING_PREFIX + "RECIPIENT_EMAILS",
    "start_date": SETTING_PREFIX + "START_DATE",
    "report_day_mode": SETTING_PREFIX + "REPORT_DAY_MODE",
}

DEFAULT_CONFIG = {
    "enabled": "0",
    # Empty means "use the legacy TIME setting" on older installations.
    "send_times": "",
    "send_time": "09:30",
    "frequency": "DAILY",
    "interval_days": "2",
    "weekdays": "0,1,2,3,4",
    "skip_weekly_holidays": "1",
    "skip_official_holidays": "1",
    "excluded_dates": "",
    "recipient_user_ids": "[]",
    "recipient_emails": "",
    "start_date": "",
    "report_day_mode": "TODAY",
}

APPROVED_LEAVE_STATUSES = {"APPROVED", "CONFIRMED", "APPROVED_BY_MANAGER"}
PENDING_LEAVE_STATUSES = {"DRAFT", "PENDING", "SUBMITTED", "IN_REVIEW"}
ATTENDANCE_AUTO_LEAVE_SOURCE = "ATTENDANCE_AUTO"
MAX_RETRY_ATTEMPTS = 5


def _setting(key: str, default: str | None = None) -> str | None:
    row = SystemSetting.query.filter_by(key=key).first()
    if row is None or row.value in (None, ""):
        return default
    return str(row.value)


def _set_setting(key: str, value: str | None) -> None:
    row = SystemSetting.query.filter_by(key=key).first()
    if row is None:
        db.session.add(SystemSetting(key=key, value=value or ""))
    else:
        row.value = value or ""


def get_report_email_config() -> dict[str, Any]:
    """Return validated settings for the UI and background worker."""
    values = {
        name: _setting(key, DEFAULT_CONFIG[name])
        for name, key in SETTING_KEYS.items()
    }
    values["enabled"] = str(values["enabled"]).strip().lower() in {"1", "true", "yes", "on"}
    values["skip_weekly_holidays"] = str(values["skip_weekly_holidays"]).strip().lower() in {"1", "true", "yes", "on"}
    values["skip_official_holidays"] = str(values["skip_official_holidays"]).strip().lower() in {"1", "true", "yes", "on"}

    frequency = str(values["frequency"] or "DAILY").strip().upper()
    if frequency not in {"DAILY", "EVERY_N_DAYS", "WEEKLY"}:
        frequency = "DAILY"
    values["frequency"] = frequency

    report_day_mode = str(values["report_day_mode"] or "TODAY").strip().upper()
    if report_day_mode not in {"TODAY", "PREVIOUS_WORKING_DAY"}:
        report_day_mode = "TODAY"
    values["report_day_mode"] = report_day_mode

    try:
        interval = max(1, min(int(str(values["interval_days"] or "2")), 365))
    except (TypeError, ValueError):
        interval = 2
    values["interval_days"] = interval

    send_times = _parse_hhmm_list(values.get("send_times"))
    if not send_times:
        legacy_send_time = _parse_hhmm(values.get("send_time"))
        send_times = [legacy_send_time] if legacy_send_time else [dt_time(9, 30)]
    values["send_times_list"] = [value.strftime("%H:%M") for value in send_times]
    values["send_times"] = "\n".join(values["send_times_list"])
    # Existing callers and old templates can continue to read the first slot.
    values["send_time"] = values["send_times_list"][0]
    values["weekdays_list"] = _parse_int_list(values["weekdays"], 0, 6)
    values["excluded_dates_list"] = _parse_date_list(values["excluded_dates"])
    values["recipient_user_ids_list"] = _parse_int_list(values["recipient_user_ids"], 1, None)
    values["recipient_emails_list"] = _parse_email_list(values["recipient_emails"])
    values["start_date"] = _parse_date(values["start_date"])
    return values


def save_report_email_config(payload: dict[str, Any]) -> dict[str, Any]:
    """Validate and persist the settings submitted by the HR form."""
    send_times = _parse_hhmm_list(payload.get("send_times"))
    if not send_times:
        send_times = _parse_hhmm_list(payload.get("send_time"))
    if not send_times:
        send_times = [dt_time(9, 30)]
    frequency = str(payload.get("frequency") or "DAILY").strip().upper()
    if frequency not in {"DAILY", "EVERY_N_DAYS", "WEEKLY"}:
        frequency = "DAILY"
    try:
        interval_days = max(1, min(int(payload.get("interval_days") or 2), 365))
    except (TypeError, ValueError):
        interval_days = 2

    weekdays = sorted(set(_parse_int_list(payload.get("weekdays"), 0, 6)))
    if frequency == "WEEKLY" and not weekdays:
        weekdays = [0, 1, 2, 3, 4]

    excluded_dates = sorted(set(_parse_date_list(payload.get("excluded_dates"))))
    start_date = _parse_date(payload.get("start_date")) or date.today()
    report_day_mode = str(payload.get("report_day_mode") or "TODAY").strip().upper()
    if report_day_mode not in {"TODAY", "PREVIOUS_WORKING_DAY"}:
        report_day_mode = "TODAY"

    user_ids = sorted(set(_parse_int_list(payload.get("recipient_user_ids"), 1, None)))
    # Only persist ids that still exist and have a valid delivery address.
    if user_ids:
        valid_users = {
            int(user.id)
            for user in User.query.filter(User.id.in_(user_ids)).all()
            if resolve_user_delivery_email(user)
        }
        user_ids = [user_id for user_id in user_ids if user_id in valid_users]

    emails = _parse_email_list(payload.get("recipient_emails"))
    values = {
        "enabled": "1" if _as_bool(payload.get("enabled")) else "0",
        "send_times": "\n".join(value.strftime("%H:%M") for value in send_times),
        # Keep the legacy setting synchronized for older workers during a
        # rolling deployment.
        "send_time": send_times[0].strftime("%H:%M"),
        "frequency": frequency,
        "interval_days": str(interval_days),
        "weekdays": ",".join(str(value) for value in weekdays),
        "skip_weekly_holidays": "1" if _as_bool(payload.get("skip_weekly_holidays")) else "0",
        "skip_official_holidays": "1" if _as_bool(payload.get("skip_official_holidays")) else "0",
        "excluded_dates": ",".join(excluded_dates),
        "recipient_user_ids": json.dumps(user_ids, ensure_ascii=False),
        "recipient_emails": "\n".join(emails),
        "start_date": start_date.isoformat(),
        "report_day_mode": report_day_mode,
    }
    for name, key in SETTING_KEYS.items():
        _set_setting(key, values[name])
    db.session.commit()
    return get_report_email_config()


def _as_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value or "").strip().lower() in {"1", "true", "yes", "on", "checked"}


def _parse_hhmm(value: Any) -> dt_time | None:
    raw = str(value or "").strip()
    if not re.fullmatch(r"\d{2}:\d{2}", raw):
        return None
    try:
        return datetime.strptime(raw, "%H:%M").time()
    except ValueError:
        return None


def _parse_hhmm_list(value: Any) -> list[dt_time]:
    """Parse and normalize one or more HH:MM values."""
    if isinstance(value, (list, tuple, set)):
        raw_values = list(value)
    else:
        raw = str(value or "").strip()
        if raw.startswith("["):
            try:
                decoded = json.loads(raw)
                raw_values = decoded if isinstance(decoded, list) else []
            except Exception:
                raw_values = []
        else:
            raw_values = re.split(r"[,;\n\r\s]+", raw) if raw else []

    parsed = {
        parsed_value
        for item in raw_values
        if (parsed_value := _parse_hhmm(item)) is not None
    }
    return sorted(parsed)


def _parse_date(value: Any) -> date | None:
    try:
        return date.fromisoformat(str(value or "").strip())
    except (TypeError, ValueError):
        return None


def _parse_int_list(value: Any, minimum: int, maximum: int | None) -> list[int]:
    if isinstance(value, (list, tuple, set)):
        raw_values = list(value)
    else:
        raw = str(value or "").strip()
        if raw.startswith("["):
            try:
                decoded = json.loads(raw)
                raw_values = decoded if isinstance(decoded, list) else []
            except Exception:
                raw_values = []
        else:
            raw_values = re.split(r"[,;\n\r\s]+", raw) if raw else []
    result = []
    for raw_item in raw_values:
        try:
            number = int(str(raw_item).strip())
        except (TypeError, ValueError):
            continue
        if number < minimum or (maximum is not None and number > maximum):
            continue
        result.append(number)
    return sorted(set(result))


def _parse_date_list(value: Any) -> list[str]:
    if isinstance(value, (list, tuple, set)):
        raw_values = list(value)
    else:
        raw = str(value or "").strip()
        if raw.startswith("["):
            try:
                decoded = json.loads(raw)
                raw_values = decoded if isinstance(decoded, list) else []
            except Exception:
                raw_values = []
        else:
            raw_values = re.split(r"[,;\n\r\s]+", raw) if raw else []
    result = []
    for raw_item in raw_values:
        parsed = _parse_date(raw_item)
        if parsed:
            result.append(parsed.isoformat())
    return sorted(set(result))


def _parse_email_list(value: Any) -> list[str]:
    raw = "\n".join(str(item) for item in value) if isinstance(value, (list, tuple, set)) else str(value or "")
    result = []
    for candidate in re.split(r"[,;\n\r\s]+", raw):
        candidate = candidate.strip()
        if re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", candidate):
            result.append(candidate)
    return sorted(set(result), key=str.casefold)


def _official_holiday(day_value: date) -> bool:
    day_text = day_value.isoformat()
    try:
        if HROfficialOccasion.query.filter_by(day=day_text, is_day_off=True).first():
            return True
    except Exception:
        pass
    try:
        return bool(
            HROfficialOccasionRange.query.filter(
                HROfficialOccasionRange.start_day <= day_text,
                HROfficialOccasionRange.end_day >= day_text,
                HROfficialOccasionRange.is_day_off.is_(True),
            ).first()
        )
    except Exception:
        return False


def _weekly_holiday(day_value: date) -> bool:
    try:
        from portal.routes import _weekly_mask

        return bool(int(_weekly_mask()) & (1 << day_value.weekday()))
    except Exception:
        return False


def _send_day_is_excluded(day_value: date, config: dict[str, Any]) -> bool:
    if day_value.isoformat() in set(config.get("excluded_dates_list") or []):
        return True
    if config.get("skip_weekly_holidays") and _weekly_holiday(day_value):
        return True
    if config.get("skip_official_holidays") and _official_holiday(day_value):
        return True
    return False


def _schedule_is_due(day_value: date, config: dict[str, Any]) -> bool:
    frequency = config.get("frequency") or "DAILY"
    if frequency == "WEEKLY":
        return day_value.weekday() in set(config.get("weekdays_list") or [])
    if frequency == "EVERY_N_DAYS":
        anchor = config.get("start_date") or day_value
        return (day_value - anchor).days >= 0 and (day_value - anchor).days % int(config.get("interval_days") or 2) == 0
    return True


def _previous_working_day(day_value: date, config: dict[str, Any]) -> date:
    candidate = day_value - timedelta(days=1)
    for _ in range(370):
        if not _send_day_is_excluded(candidate, config):
            return candidate
        candidate -= timedelta(days=1)
    return day_value - timedelta(days=1)


def _report_day(now_local: datetime, config: dict[str, Any]) -> date:
    if config.get("report_day_mode") == "PREVIOUS_WORKING_DAY":
        return _previous_working_day(now_local.date(), config)
    return now_local.date()


def _employee_label(user: User) -> str:
    employee_file = getattr(user, "employee_file", None)
    return (
        getattr(employee_file, "full_name_quad", None)
        or getattr(user, "full_name", None)
        or getattr(user, "name", None)
        or getattr(user, "email", None)
        or f"المستخدم #{user.id}"
    ).strip()


def _time_text(value: datetime | None) -> str:
    return value.strftime("%H:%M") if value else ""


def _attendance_exempt_user_ids() -> set[int]:
    """Return users excluded from attendance reports when that module exists.

    The exemption feature is deployed independently on some installations.
    Resolve its model lazily so older deployments keep sending reports, while
    installations with the HR exemption table automatically honour it.
    """
    try:
        import models as app_models

        exemption_model = getattr(app_models, "HRAttendanceExemption", None)
        if exemption_model is None:
            return set()
        rows = db.session.query(exemption_model.user_id).all()
        return {int(user_id) for (user_id,) in rows if user_id}
    except Exception:
        db.session.rollback()
        return set()


def _schedule_info(user_id: int, day_value: date):
    """Resolve effective schedule without making report delivery depend on UI state."""
    try:
        from portal.routes import (
            _effective_schedule_for_user,
            _effective_work_policy_for_user,
        )

        schedule = _effective_schedule_for_user(int(user_id), day_value.isoformat())
        if not schedule:
            return None, False
        # A fixed work-policy mask can turn an otherwise active schedule into
        # a scheduled day off for this employee.  Keep that distinction out
        # of the absence report.
        try:
            policy = _effective_work_policy_for_user(int(user_id), day_value.isoformat())
            if (
                policy
                and (getattr(policy, "days_policy", None) or "FIXED").upper() == "FIXED"
                and getattr(policy, "fixed_days_mask", None) is not None
                and not (int(policy.fixed_days_mask) & (1 << day_value.weekday()))
            ):
                return schedule, False
        except Exception:
            pass
        kind = (getattr(schedule, "kind", None) or "").strip().upper()
        if kind == "SHIFT":
            day_row = next(
                (
                    item
                    for item in (getattr(schedule, "days", None) or [])
                    if int(getattr(item, "weekday", -1)) == day_value.weekday()
                ),
                None,
            )
            if not day_row:
                # Approved schedule-day proxies do not expose ``days``; use
                # their resolved times directly instead of treating them as
                # an unscheduled holiday.
                return schedule, bool(
                    getattr(schedule, "start_time", None)
                    or getattr(schedule, "end_time", None)
                )
            return schedule, bool(getattr(day_row, "start_time", None) or getattr(day_row, "end_time", None))
        return schedule, kind != "OFF"
    except Exception:
        return None, True


def _load_attendance_rows(report_day: date) -> tuple[list[dict], list[dict], list[dict]]:
    """Build (overall, late, absence) rows for one local calendar day."""
    day_text = report_day.isoformat()
    start_dt = datetime.combine(report_day, dt_time.min)
    end_dt = datetime.combine(report_day + timedelta(days=1), dt_time.min)
    users = (
        User.query.join(EmployeeFile, EmployeeFile.user_id == User.id)
        .order_by(func.coalesce(EmployeeFile.full_name_quad, User.name, User.email).asc(), User.id.asc())
        .all()
    )
    exempt_ids = _attendance_exempt_user_ids()
    if exempt_ids:
        users = [user for user in users if int(user.id) not in exempt_ids]
    user_ids = [int(user.id) for user in users]
    if not user_ids:
        return [], [], []

    summaries = (
        AttendanceDailySummary.query
        .filter(AttendanceDailySummary.user_id.in_(user_ids), AttendanceDailySummary.day == day_text)
        .all()
    )
    summary_map = {int(row.user_id): row for row in summaries}
    events = (
        AttendanceEvent.query
        .filter(
            AttendanceEvent.user_id.in_(user_ids),
            AttendanceEvent.event_dt >= start_dt,
            AttendanceEvent.event_dt < end_dt,
        )
        .order_by(AttendanceEvent.event_dt.asc())
        .all()
    )
    events_map: dict[int, list[AttendanceEvent]] = {}
    for event in events:
        events_map.setdefault(int(event.user_id), []).append(event)

    # Project only the columns needed here.  Long-lived SQLite databases can
    # temporarily lag behind the full HRLeaveRequest model while runtime
    # schema synchronization is catching up; selecting the entity would make
    # an otherwise usable report fail on an unrelated newly-added column.
    approved = (
        db.session.query(
            HRLeaveRequest.id,
            HRLeaveRequest.user_id,
            HRLeaveRequest.leave_type_id,
            HRLeaveRequest.start_date,
            HRLeaveRequest.end_date,
            HRLeaveRequest.status,
            HRLeaveRequest.source,
            HRLeaveRequest.replaced_at,
        )
        .filter(
            HRLeaveRequest.user_id.in_(user_ids),
            HRLeaveRequest.start_date <= day_text,
            HRLeaveRequest.end_date >= day_text,
            HRLeaveRequest.status.in_(sorted(APPROVED_LEAVE_STATUSES)),
        )
        .order_by(HRLeaveRequest.id.desc())
        .all()
    )
    pending = (
        db.session.query(
            HRLeaveRequest.id,
            HRLeaveRequest.user_id,
            HRLeaveRequest.leave_type_id,
            HRLeaveRequest.start_date,
            HRLeaveRequest.end_date,
            HRLeaveRequest.status,
            HRLeaveRequest.source,
            HRLeaveRequest.replaced_at,
        )
        .filter(
            HRLeaveRequest.user_id.in_(user_ids),
            HRLeaveRequest.start_date <= day_text,
            HRLeaveRequest.end_date >= day_text,
            HRLeaveRequest.status.in_(sorted(PENDING_LEAVE_STATUSES)),
        )
        .order_by(HRLeaveRequest.id.desc())
        .all()
    )
    approved_map = {}
    pending_map = {}
    for row in approved:
        # The attendance reconciliation job creates an APPROVED annual leave
        # row for an unpunched workday.  It is a system charge, not a leave
        # submitted through Masar, so it must remain visible as an absence.
        if (
            (getattr(row, "source", None) or "").strip().upper()
            == ATTENDANCE_AUTO_LEAVE_SOURCE
            and not getattr(row, "replaced_at", None)
        ):
            continue
        approved_map.setdefault(int(row.user_id), row)
    for row in pending:
        pending_map.setdefault(int(row.user_id), row)

    overall = []
    late_rows = []
    absence_rows = []
    for user in users:
        user_id = int(user.id)
        summary = summary_map.get(user_id)
        user_events = events_map.get(user_id, [])
        incoming = [event.event_dt for event in user_events if (event.event_type or "").strip().upper() in {"I", "IN", "ENTRY", "CHECKIN"}]
        outgoing = [event.event_dt for event in user_events if (event.event_type or "").strip().upper() in {"O", "OUT", "EXIT", "CHECKOUT"}]
        first_in = getattr(summary, "first_in", None) if summary else None
        last_out = getattr(summary, "last_out", None) if summary else None
        first_in = first_in or (min(incoming) if incoming else None)
        last_out = last_out or (max(outgoing) if outgoing else None)
        has_attendance = bool(first_in or last_out or user_events)

        schedule, scheduled_work = _schedule_info(user_id, report_day)
        schedule_kind = (getattr(schedule, "kind", None) or "").strip().upper() if schedule else ""
        approved_leave = approved_map.get(user_id)
        pending_leave = pending_map.get(user_id)
        if approved_leave:
            status = "إجازة معتمدة"
        elif pending_leave and not has_attendance:
            status = "إجازة قيد المسار"
        elif schedule_kind == "REMOTE":
            status = "عن بُعد"
        elif not scheduled_work:
            status = "عطلة / لا يوجد دوام"
        elif has_attendance:
            status = "حاضر"
        else:
            status = "غياب دون إجازة"

        late_minutes = int(getattr(summary, "late_minutes", 0) or 0) if summary else 0
        if late_minutes <= 0 and first_in and schedule and schedule_kind != "REMOTE":
            configured_start = getattr(schedule, "start_time", None)
            configured_grace = getattr(schedule, "start_grace_minutes", None)
            if schedule_kind == "SHIFT":
                shift_day = next(
                    (
                        item
                        for item in (getattr(schedule, "days", None) or [])
                        if int(getattr(item, "weekday", -1)) == report_day.weekday()
                    ),
                    None,
                )
                if shift_day:
                    configured_start = getattr(shift_day, "start_time", None) or configured_start
                    configured_grace = getattr(shift_day, "grace_minutes", None)
            try:
                start_minutes = int(str(configured_start).split(":")[0]) * 60 + int(str(configured_start).split(":")[1])
                actual_minutes = first_in.hour * 60 + first_in.minute
                grace_value = (
                    configured_grace
                    if configured_grace is not None
                    else getattr(schedule, "grace_minutes", 0)
                )
                grace = int(grace_value or 0)
                late_minutes = max(0, actual_minutes - start_minutes - grace)
            except Exception:
                pass

        employee_file = getattr(user, "employee_file", None)
        row = {
            "user_id": user_id,
            "date": day_text,
            "employee_no": getattr(employee_file, "employee_no", "") or "",
            "name": _employee_label(user),
            "email": resolve_user_delivery_email(user),
            "status": status,
            "first_in": _time_text(first_in),
            "last_out": _time_text(last_out),
            "late_minutes": late_minutes,
            "work_minutes": int(getattr(summary, "work_minutes", 0) or 0) if summary else 0,
            "note": (
                "إجازة معتمدة"
            ) if approved_leave else "",
        }
        overall.append(row)
        if late_minutes > 0 and has_attendance:
            late_rows.append(row.copy())
        if status == "غياب دون إجازة":
            absence_rows.append(row.copy())

    late_rows.sort(key=lambda row: (-int(row["late_minutes"]), row["name"]))
    absence_rows.sort(key=lambda row: row["name"])
    return overall, late_rows, absence_rows


def _load_departure_rows(report_day: date, overall_rows: list[dict]) -> list[dict]:
    """Build the departures list using the same reconciliation as the events page."""
    try:
        from portal.routes import (
            _departure_display_lines,
            _departure_source_label,
            _departure_time_range_label,
            _reconciled_departure_records,
        )

        # The events page includes system-only departure requests, so include
        # all users here rather than relying only on employees with a daily
        # attendance summary.  The exemption list is applied before querying
        # the reconciliation helper, matching the HR exceptions page rules.
        exempt_ids = _attendance_exempt_user_ids()
        all_user_ids = {
            int(user_id)
            for (user_id,) in db.session.query(User.id).all()
            if user_id and int(user_id) not in exempt_ids
        }
        all_user_ids.update(
            int(row["user_id"])
            for row in (overall_rows or [])
            if row.get("user_id") and int(row["user_id"]) not in exempt_ids
        )
        if not all_user_ids:
            return []

        day_text = report_day.isoformat()
        records = _reconciled_departure_records(
            sorted(all_user_ids),
            day_text,
            day_text,
            include_pending=True,
        )
        if not records:
            return []

        users = {
            int(user.id): user
            for user in User.query.filter(User.id.in_(sorted(all_user_ids))).all()
        }
        rows = []
        for record in records:
            user = users.get(int(record.get("user_id") or 0))
            employee_file = getattr(user, "employee_file", None) if user else None
            display_lines = _departure_display_lines([record])
            time_range = " | ".join(
                f"{line['source_label']}: {line['time_range']}"
                for line in display_lines
            )
            if not time_range:
                time_range = _departure_time_range_label(
                    record.get("from_dt"),
                    record.get("to_dt"),
                )

            approval_status = str(record.get("approval_status") or "APPROVED").upper()
            if approval_status in {"SUBMITTED", "PENDING"}:
                status = "قيد الاعتماد"
            elif not record.get("complete"):
                status = "غير مكتملة"
            elif record.get("source") in {"SYSTEM", "CLOCK_SYSTEM"}:
                status = "معتمدة"
            else:
                status = "مسجلة"

            rows.append({
                "date": record.get("day") or day_text,
                "employee_no": getattr(employee_file, "employee_no", "") or "",
                "name": _employee_label(user) if user else f"المستخدم #{record.get('user_id')}",
                "departure_type": record.get("label") or (
                    "مغادرة رسمية" if record.get("kind") == "OFFICIAL" else "مغادرة شخصية"
                ),
                "permission_name": record.get("permission_name") or "",
                "source": _departure_source_label(record.get("source")),
                "time_range": time_range,
                "minutes": int(record.get("display_minutes") or 0),
                "status": status,
            })

        return sorted(
            rows,
            key=lambda row: (
                row.get("name") or "",
                row.get("date") or "",
                row.get("time_range") or "",
            ),
        )
    except Exception:
        # A report email must remain useful on installations that are midway
        # through an HR schema upgrade.  Attendance rows can still be sent if
        # the optional departures reconciliation is unavailable.
        db.session.rollback()
        return []


def build_attendance_report_data(report_day: date) -> dict[str, Any]:
    overall, late_rows, absence_rows = _load_attendance_rows(report_day)
    departures = _load_departure_rows(report_day, overall)
    return {
        "report_day": report_day,
        "overall": overall,
        "late": late_rows,
        "absence": absence_rows,
        "departures": departures,
    }


def _xlsx_bytes(
    title: str,
    headers: list[str],
    rows: list[list[Any]],
    fill_color: str,
    report_day: str | None = None,
) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "تقرير"
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "A4"
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(headers)))
    sheet.cell(1, 1, title)
    sheet.cell(1, 1).font = Font(name="Arial", size=15, bold=True, color="FFFFFF")
    sheet.cell(1, 1).fill = PatternFill("solid", fgColor=fill_color)
    sheet.cell(1, 1).alignment = Alignment(horizontal="center")
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(1, len(headers)))
    sheet.cell(2, 1, f"تاريخ التقرير: {report_day or date.today().isoformat()}")
    sheet.cell(2, 1).font = Font(name="Arial", italic=True, color="666666")
    sheet.append(headers)
    thin = Side(style="thin", color="D9E2F3")
    for cell in sheet[3]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=fill_color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in rows:
        sheet.append(row)
    for row in sheet.iter_rows(min_row=4):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    sheet.auto_filter.ref = f"A3:{get_column_letter(max(1, len(headers)))}{max(3, sheet.max_row)}"
    for index, column in enumerate(sheet.columns, 1):
        values = [str(cell.value or "") for cell in column]
        sheet.column_dimensions[get_column_letter(index)].width = min(42, max(12, max(len(value) for value in values) + 2))
    sheet.row_dimensions[1].height = 28
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _attachments(data: dict[str, Any]) -> list[tuple[str, str, bytes]]:
    report_day = data["report_day"].isoformat()
    overall_headers = ["التاريخ", "الرقم الوظيفي", "الموظف", "الحالة", "أول دخول", "آخر خروج", "دقائق التأخير", "ساعات العمل", "ملاحظة"]
    delay_headers = ["التاريخ", "الرقم الوظيفي", "الموظف", "وقت الدخول", "مدة التأخير (دقيقة)", "مدة التأخير (ساعات)"]
    absence_headers = ["التاريخ", "الرقم الوظيفي", "الموظف", "الحالة", "أول دخول", "آخر خروج"]
    departure_headers = [
        "التاريخ",
        "الرقم الوظيفي",
        "الموظف",
        "نوع المغادرة",
        "نوع الطلب",
        "المصدر",
        "مصدر وأوقات المغادرة",
        "المدة (دقيقة)",
        "الحالة",
    ]
    overall_rows = [
        [row["date"], row["employee_no"], row["name"], row["status"], row["first_in"], row["last_out"], row["late_minutes"], round(row["work_minutes"] / 60, 2), row["note"]]
        for row in data["overall"]
    ]
    delay_rows = [
        [row["date"], row["employee_no"], row["name"], row["first_in"], row["late_minutes"], round(row["late_minutes"] / 60, 2)]
        for row in data["late"]
    ]
    absence_rows = [
        [row["date"], row["employee_no"], row["name"], row["status"], row["first_in"], row["last_out"]]
        for row in data["absence"]
    ]
    departure_rows = [
        [
            row["date"],
            row["employee_no"],
            row["name"],
            row["departure_type"],
            row["permission_name"],
            row["source"],
            row["time_range"],
            row["minutes"],
            row["status"],
        ]
        for row in data.get("departures", [])
    ]
    return [
        (f"attendance_daily_{report_day}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", _xlsx_bytes("تقرير الدوام اليومي", overall_headers, overall_rows, "2F5597", report_day)),
        (f"attendance_morning_delay_{report_day}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", _xlsx_bytes("تقرير التأخير الصباحي", delay_headers, delay_rows, "BF7E00", report_day)),
        (f"attendance_absence_without_leave_{report_day}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", _xlsx_bytes("تقرير الغياب دون إجازة", absence_headers, absence_rows, "B42318", report_day)),
        (f"attendance_departures_{report_day}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", _xlsx_bytes("قائمة المغادرات", departure_headers, departure_rows, "2F7D6D", report_day)),
    ]


def _html_table(rows: list[dict], columns: list[tuple[str, str]]) -> str:
    if not rows:
        return '<p style="margin:0;color:#6b7280">لا توجد سجلات ضمن هذا التقرير.</p>'
    header = "".join(f'<th style="padding:8px;border:1px solid #e5e7eb;background:#f8fafc">{escape(label)}</th>' for key, label in columns)
    body = []
    for row in rows:
        cells = "".join(
            f'<td style="padding:7px;border:1px solid #e5e7eb">{escape(str(row.get(key, "") or ""))}</td>'
            for key, _label in columns
        )
        body.append(f"<tr>{cells}</tr>")
    return f'<table style="width:100%;border-collapse:collapse;font-size:13px"><thead><tr>{header}</tr></thead><tbody>{"".join(body)}</tbody></table>'


def build_email_content(data: dict[str, Any]) -> tuple[str, str, str]:
    report_day = data["report_day"].isoformat()
    subject = f"{REPORT_NAME} — {report_day}"
    scheduled_time = str(data.get("scheduled_time") or "").strip()
    if scheduled_time:
        subject += f" — {scheduled_time}"
    sections = [
        ("تقرير الدوام اليومي العام", "#2F5597", data["overall"], [("employee_no", "الرقم الوظيفي"), ("name", "الموظف"), ("status", "الحالة"), ("first_in", "أول دخول"), ("last_out", "آخر خروج"), ("late_minutes", "التأخير (دقيقة)")]),
        ("تقرير التأخير الصباحي", "#BF7E00", data["late"], [("employee_no", "الرقم الوظيفي"), ("name", "الموظف"), ("first_in", "وقت الدخول"), ("late_minutes", "مدة التأخير (دقيقة)")]),
        ("تقرير الغياب دون إجازة عبر مسار", "#B42318", data["absence"], [("employee_no", "الرقم الوظيفي"), ("name", "الموظف"), ("status", "الحالة")]),
        ("قائمة المغادرات", "#2F7D6D", data.get("departures", []), [("employee_no", "الرقم الوظيفي"), ("name", "الموظف"), ("departure_type", "نوع المغادرة"), ("permission_name", "نوع الطلب"), ("source", "المصدر"), ("time_range", "مصدر وأوقات المغادرة"), ("minutes", "المدة (دقيقة)"), ("status", "الحالة")]),
    ]
    html_sections = []
    text_sections = []
    for title, color, rows, columns in sections:
        html_sections.append(
            f'<section style="margin:0 0 22px;border:1px solid #e5e7eb;border-radius:8px;overflow:hidden">'
            f'<h3 style="margin:0;padding:10px 14px;background:{color};color:#fff;font-size:16px">{escape(title)} ({len(rows)})</h3>'
            f'<div style="padding:12px">{_html_table(rows, columns)}</div></section>'
        )
        text_sections.append(title + f" ({len(rows)})\n" + ("\n".join(" | ".join(str(row.get(key, "") or "") for key, _ in columns) for row in rows) or "لا توجد سجلات"))
    html_body = (
        '<html><body dir="rtl" style="font-family:Arial,sans-serif;color:#1f2937;line-height:1.7">'
        f'<div style="max-width:1100px;margin:auto"><h2 style="margin-bottom:4px;color:#5a4a86">{escape(REPORT_NAME)}</h2>'
        f'<p style="margin-top:0;color:#6b7280">تاريخ التقرير: {escape(report_day)}'
        + (f" — وقت الإرسال: {escape(scheduled_time)}" if scheduled_time else "")
        + "</p>"
        + "".join(html_sections)
        + '<p style="font-size:12px;color:#6b7280">هذه رسالة آلية من نظام مسار.</p></div></body></html>'
    )
    text_body = f"{REPORT_NAME}\nتاريخ التقرير: {report_day}\n\n" + "\n\n".join(text_sections) + "\n\nهذه رسالة آلية من نظام مسار."
    return subject, text_body, html_body


def _send_email(config: dict[str, Any], recipients: list[str], subject: str, text_body: str, html_body: str, attachments: list[tuple[str, str, bytes]]) -> None:
    message = EmailMessage()
    message["Subject"] = subject
    message["From"] = formataddr((REPORT_FROM_NAME, config["from_email"]))
    message["To"] = ", ".join(recipients)
    if config.get("reply_to"):
        message["Reply-To"] = config["reply_to"]
    message.set_content(text_body)
    message.add_alternative(html_body, subtype="html")
    for filename, mime_type, payload in attachments:
        maintype, subtype = (mime_type.split("/", 1) + ["octet-stream"])[:2]
        message.add_attachment(payload, maintype=maintype, subtype=subtype, filename=filename)

    if config["security"] == "ssl":
        smtp = smtplib.SMTP_SSL(config["host"], config["port"], timeout=20, context=ssl.create_default_context())
    else:
        smtp = smtplib.SMTP(config["host"], config["port"], timeout=20)
    try:
        smtp.ehlo()
        if config["security"] == "starttls":
            smtp.starttls(context=ssl.create_default_context())
            smtp.ehlo()
        if config["username"]:
            smtp.login(config["username"], config["password"])
        smtp.send_message(message, from_addr=config["from_email"], to_addrs=recipients)
    finally:
        try:
            smtp.quit()
        except Exception:
            pass


def _recipients(config: dict[str, Any]) -> list[str]:
    result = list(config.get("recipient_emails_list") or [])
    if config.get("recipient_user_ids_list"):
        users = User.query.filter(User.id.in_(config["recipient_user_ids_list"])).all()
        result.extend(resolve_user_delivery_email(user) for user in users)
    return sorted({email.strip() for email in result if email}, key=str.casefold)


def _configured_send_times(config: dict[str, Any]) -> list[dt_time]:
    values = _parse_hhmm_list(config.get("send_times_list") or config.get("send_times"))
    if values:
        return values
    return [_parse_hhmm(config.get("send_time")) or dt_time(9, 30)]


def _run_attendance_report_email_slot(
    local_now: datetime,
    config: dict[str, Any],
    scheduled_time: dt_time,
    *,
    force: bool = False,
) -> dict[str, Any]:
    """Send one idempotent day/time slot of the attendance report."""
    run_date = local_now.date().isoformat()
    scheduled_time_text = scheduled_time.strftime("%H:%M")
    report_day = _report_day(local_now, config)
    delivery = HRAttendanceEmailReportDelivery.query.filter_by(
        run_date=run_date,
        scheduled_time=scheduled_time_text,
    ).first()
    if delivery and delivery.status == "SENT":
        return {"status": "already_sent", "run_date": run_date, "scheduled_time": scheduled_time_text}
    if delivery and delivery.status == "SENDING":
        age = datetime.utcnow() - (delivery.updated_at or delivery.created_at or datetime.utcnow())
        if age < timedelta(minutes=10):
            return {"status": "in_progress", "run_date": run_date, "scheduled_time": scheduled_time_text}
    if delivery and int(delivery.attempt_count or 0) >= MAX_RETRY_ATTEMPTS and not force:
        return {"status": "retry_limit", "run_date": run_date, "scheduled_time": scheduled_time_text}
    if delivery is None:
        delivery = HRAttendanceEmailReportDelivery(
            run_date=run_date,
            scheduled_time=scheduled_time_text,
            report_day=report_day.isoformat(),
            status="SENDING",
            attempt_count=0,
        )
        db.session.add(delivery)
    else:
        delivery.report_day = report_day.isoformat()
        delivery.scheduled_time = scheduled_time_text
        delivery.status = "SENDING"
        delivery.updated_at = datetime.utcnow()
    try:
        db.session.commit()
    except Exception:
        db.session.rollback()
        existing = HRAttendanceEmailReportDelivery.query.filter_by(
            run_date=run_date,
            scheduled_time=scheduled_time_text,
        ).first()
        if existing and existing.status in {"SENT", "SENDING"}:
            return {"status": "in_progress", "run_date": run_date, "scheduled_time": scheduled_time_text}
        raise

    try:
        recipients = _recipients(config)
        if not recipients:
            raise RuntimeError("لم يتم تحديد أي مستلم صالح لتقارير الدوام.")
        if not email_delivery_enabled():
            raise RuntimeError("تم تعطيل إرسال البريد الإلكتروني من إعدادات النظام.")
        mail_config = _mail_config()
        if not mail_config.get("ready"):
            raise RuntimeError("إعدادات SMTP غير مكتملة أو البريد غير مفعّل.")
        data = build_attendance_report_data(report_day)
        data["scheduled_time"] = scheduled_time_text
        subject, text_body, html_body = build_email_content(data)
        _send_email(mail_config, recipients, subject, text_body, html_body, _attachments(data))
        delivery.status = "SENT"
        delivery.sent_at = datetime.utcnow()
        delivery.recipient_count = len(recipients)
        delivery.last_error = None
        delivery.updated_at = datetime.utcnow()
        db.session.commit()
        return {
            "status": "sent",
            "run_date": run_date,
            "scheduled_time": scheduled_time_text,
            "report_day": report_day.isoformat(),
            "recipients": len(recipients),
            "counts": {
                "overall": len(data["overall"]),
                "late": len(data["late"]),
                "absence": len(data["absence"]),
                "departures": len(data.get("departures", [])),
            },
        }
    except Exception as exc:
        db.session.rollback()
        delivery = HRAttendanceEmailReportDelivery.query.filter_by(
            run_date=run_date,
            scheduled_time=scheduled_time_text,
        ).first()
        if delivery:
            delivery.status = "FAILED"
            delivery.attempt_count = int(delivery.attempt_count or 0) + 1
            delivery.last_error = str(exc)[:2000]
            delivery.updated_at = datetime.utcnow()
            db.session.commit()
        return {
            "status": "failed",
            "run_date": run_date,
            "scheduled_time": scheduled_time_text,
            "error": str(exc)[:500],
        }


def run_attendance_report_email_cycle(now: datetime | None = None, *, force: bool = False) -> dict[str, Any]:
    """Send every due attendance-report time slot for the current day."""
    config = get_report_email_config()
    local_now = now or datetime.now(app_timezone())
    if local_now.tzinfo is None:
        local_now = local_now.replace(tzinfo=app_timezone())
    if not force and not config["enabled"]:
        return {"status": "disabled"}

    configured_times = _configured_send_times(config)
    if force:
        due_times = configured_times[:1]
    else:
        current_time = local_now.time().replace(second=0, microsecond=0)
        if current_time < configured_times[0]:
            return {"status": "not_due"}
        if not _schedule_is_due(local_now.date(), config) or _send_day_is_excluded(local_now.date(), config):
            return {"status": "excluded"}
        due_times = [scheduled_time for scheduled_time in configured_times if scheduled_time <= current_time]

    results = [
        _run_attendance_report_email_slot(local_now, config, scheduled_time, force=force)
        for scheduled_time in due_times
    ]
    if len(results) == 1:
        return results[0]

    sent_results = [result for result in results if result.get("status") == "sent"]
    failed_results = [result for result in results if result.get("status") == "failed"]
    if sent_results:
        status = "sent"
    elif failed_results:
        status = "failed"
    elif results and all(result.get("status") == "already_sent" for result in results):
        status = "already_sent"
    else:
        status = "in_progress"
    return {
        "status": status,
        "run_date": local_now.date().isoformat(),
        "scheduled_times": [scheduled_time.strftime("%H:%M") for scheduled_time in due_times],
        "runs": results,
        "sent_slots": [result.get("scheduled_time") for result in sent_results],
    }


def test_send_attendance_report_email() -> dict[str, Any]:
    """Send a manual preview using today's data without affecting the schedule marker."""
    now_local = datetime.now(app_timezone())
    config = get_report_email_config()
    recipients = _recipients(config)
    if not recipients:
        raise RuntimeError("حدد مستلماً واحداً على الأقل قبل إرسال الاختبار.")
    mail_config = _mail_config()
    if not mail_config.get("ready"):
        raise RuntimeError("إعدادات SMTP غير مكتملة أو البريد غير مفعّل.")
    data = build_attendance_report_data(now_local.date())
    subject, text_body, html_body = build_email_content(data)
    _send_email(mail_config, recipients, f"[تجريبي] {subject}", text_body, html_body, _attachments(data))
    return {"recipients": len(recipients), "report_day": now_local.date().isoformat()}
