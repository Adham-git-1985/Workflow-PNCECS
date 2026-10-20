from __future__ import annotations

import csv
import hashlib
import io
import re
import secrets
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import qrcode
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer
from flask import (
    abort,
    current_app,
    flash,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from flask_login import current_user, login_required
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image
from sqlalchemy import case, func, or_
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import joinedload, load_only, noload

from extensions import db
from models import (
    AuditLog,
    Committee,
    InvFixedAsset,
    InvFixedAssetCycle,
    InvFixedAssetCycleMember,
    InvFixedAssetMobileDevice,
    InvFixedAssetEntry,
    InvFixedAssetScanLog,
    InvAssetDocument,
    InvAssetDocumentLine,
    InvItem,
    InvItemCategory,
    InvRoom,
    InvWarehouse,
    SystemSetting,
    User,
)
from services.fixed_asset_inventory import (
    ASSET_CONDITIONS,
    ASSET_LIFECYCLE_STATUSES,
    CYCLE_STATUSES,
    INVENTORY_RESULT_COLORS,
    INVENTORY_RESULT_LABELS,
    close_inventory_cycle,
    cycle_result_counts,
    cycle_result_counts_for_cycles,
    find_asset_by_identifier,
    normalize_asset_identifier,
    populate_cycle_entries,
    record_asset_scan,
    reopen_inventory_cycle,
)

from . import portal_bp
from .routes import STORE_EXPORT, STORE_MANAGE, STORE_READ, _perm_any


PREVIOUS_RESULT_LABELS = {
    **INVENTORY_RESULT_LABELS,
    "NOT_IN_PREVIOUS": "لم يكن ضمن الجرد السابق",
    "NOT_IN_SCOPE": "خارج نطاق الدورة",
    "MASTER": "السجل الأساسي",
    None: "لا توجد دورة سابقة",
}


MOBILE_PAIRING_TOKEN_MAX_AGE = 15 * 60
MOBILE_PAIRING_TOKEN_SALT = "fixed-asset-mobile-pair-v1"
MOBILE_DEVICE_INVITE_MAX_AGE = 60 * 60
MOBILE_DEVICE_INVITE_SALT = "fixed-asset-mobile-device-invite-v1"
MOBILE_DEVICE_COOKIE = "fixed_asset_mobile_device"
MOBILE_DEVICE_COOKIE_MAX_AGE = 365 * 24 * 60 * 60
MOBILE_DEVICE_STATUS_LABELS = {
    "PENDING": "بانتظار موافقة الأدمن",
    "APPROVED": "معتمد وقادر على المسح",
    "REVOKED": "موقوف",
}
FIXED_ASSET_IMAGE_MAX_BYTES = 8 * 1024 * 1024
FIXED_ASSET_IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}
FIXED_ASSET_IMAGE_FORMAT_EXTENSIONS = {
    "JPEG": ".jpg",
    "PNG": ".png",
    "WEBP": ".webp",
}

IMPORT_HEADERS = [
    ("asset_tag", "رقم الأصل"),
    ("name", "اسم الأصل"),
    ("item", "كود/اسم الصنف"),
    ("category", "التصنيف"),
    ("serial_number", "الرقم التسلسلي"),
    ("manufacturer", "الشركة المصنعة"),
    ("model", "الموديل"),
    ("description", "الوصف"),
    ("acquisition_date", "تاريخ الشراء"),
    ("purchase_cost", "التكلفة"),
    ("asset_condition", "حالة الأصل"),
    ("lifecycle_status", "وضع الأصل"),
    ("warehouse", "المستودع"),
    ("room", "الغرفة"),
    ("custodian", "الموظف المسؤول"),
    ("note", "ملاحظات"),
]

IMPORT_ALIASES = {
    "asset_tag": ("رقمالأصل", "رقمالاصل", "كودالأصل", "كودالاصل", "assettag", "tag", "assetcode"),
    "name": ("اسمالأصل", "اسمالاصل", "الأصل", "الاصل", "name", "assetname"),
    "item": ("كوداسمالصنف", "كودالصنف", "اسمالصنف", "item", "itemcode"),
    "category": ("التصنيف", "الفئة", "category"),
    "serial_number": ("الرقمالتسلسلي", "السيريال", "serial", "serialnumber"),
    "manufacturer": ("الشركةالمصنعة", "الشركة", "manufacturer", "brand"),
    "model": ("الموديل", "الطراز", "model"),
    "description": ("الوصف", "البيان", "description"),
    "acquisition_date": ("تاريخالشراء", "تاريخالاقتناء", "acquisitiondate", "purchasedate"),
    "purchase_cost": ("التكلفة", "سعرالشراء", "cost", "purchasecost"),
    "asset_condition": ("حالةالأصل", "حالةالاصل", "condition", "assetcondition"),
    "lifecycle_status": ("وضعالأصل", "وضعالاصل", "الحالةالإدارية", "status", "lifecyclestatus"),
    "warehouse": ("المستودع", "المخزن", "warehouse"),
    "room": ("الغرفة", "المكتب", "room"),
    "custodian": ("الموظفالمسؤول", "صاحبالعهدة", "العهدة", "custodian", "employee"),
    "note": ("ملاحظات", "ملاحظة", "notes", "note"),
}


@portal_bp.app_context_processor
def inject_fixed_asset_labels():
    return {
        "fixed_asset_condition_labels": ASSET_CONDITIONS,
        "fixed_asset_lifecycle_labels": ASSET_LIFECYCLE_STATUSES,
        "fixed_asset_cycle_labels": CYCLE_STATUSES,
        "fixed_asset_result_labels": INVENTORY_RESULT_LABELS,
        "fixed_asset_result_colors": INVENTORY_RESULT_COLORS,
        "fixed_asset_previous_labels": PREVIOUS_RESULT_LABELS,
    }


def _optional_int(value) -> int | None:
    try:
        return int(value) if str(value or "").strip() else None
    except (TypeError, ValueError):
        return None


def _clean_text(value, limit: int | None = None) -> str | None:
    text_value = str(value or "").strip()
    if not text_value:
        return None
    return text_value[:limit] if limit else text_value


def _valid_iso_date(value, field_label: str) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text_value = str(value).strip()
    try:
        return datetime.strptime(text_value, "%Y-%m-%d").date().isoformat()
    except ValueError as exc:
        raise ValueError(f"{field_label} يجب أن يكون بصيغة YYYY-MM-DD.") from exc


def _decimal_or_none(value) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        parsed = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError) as exc:
        raise ValueError("التكلفة غير صالحة.") from exc
    if parsed < 0:
        raise ValueError("التكلفة لا يمكن أن تكون سالبة.")
    return parsed.quantize(Decimal("0.01"))


def _entity_or_none(model, entity_id: int | None, label: str):
    if entity_id is None:
        return None
    entity = db.session.get(model, entity_id)
    if entity is None:
        raise ValueError(f"{label} المحدد غير موجود.")
    return entity


def _fixed_asset_image_directory() -> Path:
    directory = Path(current_app.instance_path) / "uploads" / "fixed_assets"
    directory.mkdir(parents=True, exist_ok=True)
    return directory


def _fixed_asset_image_path(filename: str | None) -> Path | None:
    safe_name = str(filename or "").strip()
    if not safe_name or Path(safe_name).name != safe_name:
        return None
    if Path(safe_name).suffix.lower() not in FIXED_ASSET_IMAGE_EXTENSIONS:
        return None
    return _fixed_asset_image_directory() / safe_name


def _save_fixed_asset_image(upload) -> str | None:
    if upload is None or not (getattr(upload, "filename", None) or "").strip():
        return None

    try:
        data = upload.read(FIXED_ASSET_IMAGE_MAX_BYTES + 1)
    except (OSError, ValueError) as exc:
        raise ValueError("تعذر قراءة صورة الأصل.") from exc
    if len(data) > FIXED_ASSET_IMAGE_MAX_BYTES:
        raise ValueError("حجم صورة الأصل يجب ألا يتجاوز 8 ميغابايت.")

    try:
        with Image.open(io.BytesIO(data)) as image:
            image_format = (image.format or "").upper()
            image.verify()
    except Exception as exc:
        raise ValueError("الملف المرفوع ليس صورة صالحة.") from exc

    extension = FIXED_ASSET_IMAGE_FORMAT_EXTENSIONS.get(image_format)
    if extension is None:
        raise ValueError("صيغة الصورة المدعومة هي JPG أو PNG أو WEBP.")

    filename = f"{secrets.token_hex(20)}{extension}"
    path = _fixed_asset_image_path(filename)
    if path is None:
        raise ValueError("تعذر تجهيز صورة الأصل.")
    try:
        path.write_bytes(data)
    except OSError as exc:
        raise ValueError("تعذر حفظ صورة الأصل.") from exc
    return filename


def _delete_fixed_asset_image(filename: str | None):
    path = _fixed_asset_image_path(filename)
    if path is None:
        return
    try:
        path.unlink(missing_ok=True)
    except OSError:
        current_app.logger.warning("Unable to remove fixed asset image %s", path)


def _audit(action: str, note: str, target_type: str, target_id: int | None):
    db.session.add(AuditLog(
        user_id=getattr(current_user, "id", None),
        action=action,
        note=note,
        target_type=target_type,
        target_id=target_id,
        created_at=datetime.utcnow(),
    ))


def _user_choices() -> list[User]:
    """Load only the columns needed by fixed-asset selectors.

    ``User.permissions`` and ``User.employee_file`` are select-in relationships.
    Loading every user for a committee/custodian selector used to hydrate both
    collections for the whole organization, even though the form only renders
    the user's display name and username.
    """
    return (
        User.query
        .options(
            load_only(User.id, User.name, User.username, User.email),
            noload(User.permissions),
            noload(User.employee_file),
        )
        .order_by(func.coalesce(User.name, User.username, User.email).asc(), User.id.asc())
        .all()
    )


def _choices(*, include_items: bool = False, include_categories: bool = True, include_users: bool = True) -> dict:
    """Return only the lookup lists required by the current fixed-asset page.

    Item catalogs can be very large.  Callers that only need warehouse/room or
    committee selectors must not pay the cost of loading the catalog (or the
    user's eager relationships) just to render a form.
    """
    choices = {
        "items": [],
        "categories": [],
        "warehouses": InvWarehouse.query
        .filter(InvWarehouse.is_active.is_(True))
        .order_by(InvWarehouse.name.asc())
        .all(),
        "rooms": InvRoom.query
        .filter(InvRoom.is_active.is_(True))
        .order_by(InvRoom.name.asc())
        .all(),
        "users": _user_choices() if include_users else [],
    }
    if include_items:
        choices["items"] = (
            InvItem.query
            .options(noload(InvItem.attributes))
            .filter(InvItem.is_active.is_(True))
            .order_by(InvItem.name.asc())
            .all()
        )
    if include_categories:
        choices["categories"] = (
            InvItemCategory.query
            .filter(InvItemCategory.is_active.is_(True))
            .order_by(InvItemCategory.name.asc())
            .all()
        )
    return choices


def _selected_item(item_id: int | None) -> InvItem | None:
    """Load one selected item without hydrating its attribute collection."""
    if not item_id:
        return None
    return (
        InvItem.query
        .options(noload(InvItem.attributes))
        .filter(InvItem.id == item_id)
        .first()
    )


def _cycle_entry_load_options() -> list:
    """Keep the committee table query focused on fields it actually renders."""
    asset_loader = (
        joinedload(InvFixedAssetEntry.asset)
        .load_only(
            InvFixedAsset.id,
            InvFixedAsset.asset_tag,
            InvFixedAsset.name,
            InvFixedAsset.serial_number,
        )
        .options(
            noload(InvFixedAsset.item),
            noload(InvFixedAsset.category),
            noload(InvFixedAsset.warehouse),
            noload(InvFixedAsset.room),
            noload(InvFixedAsset.custodian),
            noload(InvFixedAsset.created_by),
            noload(InvFixedAsset.updated_by),
        )
    )
    user_fields = (User.id, User.name, User.username, User.email)
    user_options = (noload(User.permissions), noload(User.employee_file))
    return [
        asset_loader,
        noload(InvFixedAssetEntry.scans),
        joinedload(InvFixedAssetEntry.expected_warehouse).load_only(
            InvWarehouse.id,
            InvWarehouse.name,
            InvWarehouse.code,
        ),
        joinedload(InvFixedAssetEntry.expected_room)
        .load_only(InvRoom.id, InvRoom.name, InvRoom.code)
        .options(noload(InvRoom.unit)),
        joinedload(InvFixedAssetEntry.expected_custodian)
        .load_only(*user_fields)
        .options(*user_options),
        joinedload(InvFixedAssetEntry.observed_warehouse).load_only(
            InvWarehouse.id,
            InvWarehouse.name,
            InvWarehouse.code,
        ),
        joinedload(InvFixedAssetEntry.observed_room)
        .load_only(InvRoom.id, InvRoom.name, InvRoom.code)
        .options(noload(InvRoom.unit)),
        joinedload(InvFixedAssetEntry.observed_custodian)
        .load_only(*user_fields)
        .options(*user_options),
        joinedload(InvFixedAssetEntry.scanned_by)
        .load_only(*user_fields)
        .options(*user_options),
    ]


def _cycle_scan_load_options() -> list:
    """Avoid loading unused asset/organization relationships in scan history."""
    scan_asset_loader = (
        joinedload(InvFixedAssetScanLog.asset)
        .load_only(
            InvFixedAsset.id,
            InvFixedAsset.asset_tag,
            InvFixedAsset.name,
        )
        .options(
            noload(InvFixedAsset.item),
            noload(InvFixedAsset.category),
            noload(InvFixedAsset.warehouse),
            noload(InvFixedAsset.room),
            noload(InvFixedAsset.custodian),
            noload(InvFixedAsset.created_by),
            noload(InvFixedAsset.updated_by),
        )
    )
    scan_user_loader = (
        joinedload(InvFixedAssetScanLog.scanned_by)
        .load_only(User.id, User.name, User.username, User.email)
        .options(noload(User.permissions), noload(User.employee_file))
    )
    return [
        noload(InvFixedAssetScanLog.cycle),
        scan_asset_loader,
        scan_user_loader,
        noload(InvFixedAssetScanLog.observed_warehouse),
        noload(InvFixedAssetScanLog.observed_room),
        noload(InvFixedAssetScanLog.observed_custodian),
    ]


def _system_setting(key: str, default: str) -> str:
    row = SystemSetting.query.filter_by(key=key).first()
    return (row.value or default).strip() if row and row.value is not None else default


def _mobile_base_url() -> str:
    return _system_setting("INV_FIXED_ASSET_QR_BASE_URL", "").rstrip("/") or request.url_root.rstrip("/")


def _mobile_pairing_serializer() -> URLSafeTimedSerializer:
    return URLSafeTimedSerializer(current_app.config["SECRET_KEY"], salt=MOBILE_PAIRING_TOKEN_SALT)


def _mobile_device_invitation_serializer() -> URLSafeTimedSerializer:
    return URLSafeTimedSerializer(
        current_app.config["SECRET_KEY"],
        salt=MOBILE_DEVICE_INVITE_SALT,
    )


def _mobile_pairing_token(cycle: InvFixedAssetCycle) -> str:
    return _mobile_pairing_serializer().dumps({
        "purpose": "FIXED_ASSET_MOBILE_PAIR",
        "cycle_id": cycle.id,
    })


def _mobile_pairing_target(cycle: InvFixedAssetCycle) -> str:
    return f"{request.url_root.rstrip('/')}{url_for('portal.fixed_asset_mobile_pair', token=_mobile_pairing_token(cycle))}"


def _decode_mobile_pairing_token(token: str) -> dict | None:
    try:
        payload = _mobile_pairing_serializer().loads(token, max_age=MOBILE_PAIRING_TOKEN_MAX_AGE)
    except (BadSignature, SignatureExpired):
        return None
    if not isinstance(payload, dict) or payload.get("purpose") != "FIXED_ASSET_MOBILE_PAIR":
        return None
    return payload


def _mobile_device_invitation_token() -> str:
    return _mobile_device_invitation_serializer().dumps({
        "purpose": "FIXED_ASSET_MOBILE_DEVICE_INVITE",
        "nonce": secrets.token_urlsafe(12),
    })


def _decode_mobile_device_invitation(token: str) -> dict | None:
    try:
        payload = _mobile_device_invitation_serializer().loads(
            token,
            max_age=MOBILE_DEVICE_INVITE_MAX_AGE,
        )
    except (BadSignature, SignatureExpired):
        return None
    if not isinstance(payload, dict) or payload.get("purpose") != "FIXED_ASSET_MOBILE_DEVICE_INVITE":
        return None
    return payload


def _mobile_device_pairing_target(token: str) -> str:
    return f"{request.url_root.rstrip('/')}{url_for('portal.fixed_asset_mobile_device_pair', token=token)}"


def _mobile_device_token_hash(raw_token: str) -> str:
    return hashlib.sha256(raw_token.encode("utf-8")).hexdigest()


def _mobile_device_label() -> str:
    parts = [
        value
        for value in (request.user_agent.platform, request.user_agent.browser)
        if value
    ]
    return "هاتف" + (f" ({' / '.join(parts)})" if parts else "")


def _mobile_device_from_request() -> InvFixedAssetMobileDevice | None:
    raw_token = (request.cookies.get(MOBILE_DEVICE_COOKIE) or "").strip()
    if not raw_token or not current_user.is_authenticated:
        return None
    device = InvFixedAssetMobileDevice.query.filter_by(
        token_hash=_mobile_device_token_hash(raw_token),
    ).first()
    if device is None or device.user_id != current_user.id:
        return None
    return device


def _require_approved_mobile_device() -> InvFixedAssetMobileDevice:
    device = _mobile_device_from_request()
    if device is None:
        abort(403, description="هذا الهاتف غير مسجل في النظام. امسح QR ربط جهاز جديد أولاً.")
    if device.status != "APPROVED":
        abort(403, description="لم تتم الموافقة على هذا الهاتف بعد.")
    return device


def _register_mobile_device():
    raw_token = (request.cookies.get(MOBILE_DEVICE_COOKIE) or "").strip()
    device = _mobile_device_from_request()
    if device is None or device.status == "REVOKED":
        raw_token = secrets.token_urlsafe(32)
        device = InvFixedAssetMobileDevice(
            token_hash=_mobile_device_token_hash(raw_token),
            user_id=current_user.id,
            device_label=_mobile_device_label(),
            status="PENDING",
            requested_at=datetime.utcnow(),
            last_ip=request.remote_addr,
            user_agent=request.user_agent.string[:255],
        )
        db.session.add(device)
        db.session.flush()
    else:
        device.last_seen_at = datetime.utcnow()
        device.last_ip = request.remote_addr
        device.user_agent = request.user_agent.string[:255]
    session["fixed_asset_mobile_device_id"] = device.id
    session.modified = True
    db.session.commit()
    response = redirect(url_for("portal.fixed_asset_mobile_device_status"))
    response.set_cookie(
        MOBILE_DEVICE_COOKIE,
        raw_token,
        max_age=MOBILE_DEVICE_COOKIE_MAX_AGE,
        httponly=True,
        samesite="Lax",
        secure=False,
    )
    return response


def _asset_qr_target(asset: InvFixedAsset) -> str:
    path = url_for("portal.fixed_asset_qr_resolve", token=asset.qr_token)
    configured_base = _system_setting("INV_FIXED_ASSET_QR_BASE_URL", "").rstrip("/")
    if configured_base:
        return f"{configured_base}{path}"
    return url_for("portal.fixed_asset_qr_resolve", token=asset.qr_token, _external=True)


def _can_scan_cycle(cycle):
    return current_user.has_perm(STORE_MANAGE) or any(member.user_id == current_user.id for member in cycle.members)


def _find_asset_by_system_qr(raw_value: str | None) -> InvFixedAsset | None:
    """Resolve only the random QR token issued by this system, never an asset tag."""
    identifier = normalize_asset_identifier(raw_value)
    if not identifier:
        return None
    # ``qr_token`` is generated by the system and is case-sensitive.  Query it
    # directly so the unique index is used even when the asset table is large.
    return InvFixedAsset.query.filter(InvFixedAsset.qr_token == identifier).first()


def _require_asset_access(asset):
    if current_user.has_perm(STORE_READ) or current_user.has_perm(STORE_MANAGE) or asset.custodian_user_id == current_user.id:
        return
    member = InvFixedAssetCycleMember.query.join(InvFixedAssetCycle).filter(InvFixedAssetCycleMember.user_id == current_user.id, InvFixedAssetCycle.status == "ACTIVE").first()
    if member:
        return  # Committee members can identify unexpected assets during field scans.
    abort(403)


def _mobile_scan_defaults(cycle: InvFixedAssetCycle, asset: InvFixedAsset) -> dict:
    entry = InvFixedAssetEntry.query.filter_by(cycle_id=cycle.id, asset_id=asset.id).first()

    def value(observed_name: str, expected_name: str, fallback):
        if entry is not None and entry.scanned_at:
            observed = getattr(entry, observed_name)
            if observed is not None:
                return observed
        if entry is not None:
            expected = getattr(entry, expected_name)
            if expected is not None:
                return expected
        return fallback

    return {
        "warehouse_id": value("observed_warehouse_id", "expected_warehouse_id", asset.warehouse_id),
        "room_id": value("observed_room_id", "expected_room_id", asset.room_id),
        "custodian_user_id": value(
            "observed_custodian_user_id",
            "expected_custodian_user_id",
            asset.custodian_user_id,
        ),
        "condition": value("observed_condition", "expected_condition", asset.asset_condition or "GOOD"),
        "lifecycle_status": value(
            "observed_lifecycle_status",
            "expected_lifecycle_status",
            asset.lifecycle_status or "ACTIVE",
        ),
    }


def _record_mobile_qr_scan(cycle: InvFixedAssetCycle, asset: InvFixedAsset):
    defaults = _mobile_scan_defaults(cycle, asset)
    entry = record_asset_scan(
        cycle,
        asset,
        scanned_by_id=current_user.id,
        observed_warehouse_id=defaults["warehouse_id"],
        observed_room_id=defaults["room_id"],
        observed_custodian_user_id=defaults["custodian_user_id"],
        observed_condition=defaults["condition"] or "GOOD",
        observed_lifecycle_status=defaults["lifecycle_status"] or "ACTIVE",
        source="PHONE_QR",
        device_info=request.user_agent.string,
    )
    _audit(
        "INV_FIXED_ASSET_PHONE_QR_SCAN",
        f"قراءة هاتفية للأصل {asset.asset_tag}: {INVENTORY_RESULT_LABELS.get(entry.result_status, entry.result_status)}",
        "INV_FIXED_ASSET_CYCLE",
        cycle.id,
    )
    db.session.commit()
    return entry


@portal_bp.route("/inventory/fixed-assets/mobile/device/pair/<token>")
@login_required
def fixed_asset_mobile_device_pair(token: str):
    if _decode_mobile_device_invitation(token) is None:
        abort(404)
    return _register_mobile_device()


@portal_bp.route("/inventory/fixed-assets/mobile/pair/<token>")
@login_required
def fixed_asset_mobile_pair(token: str):
    """Keep old printed pairing labels useful, but route them through approval."""
    payload = _decode_mobile_pairing_token(token)
    cycle_id = _optional_int(payload.get("cycle_id")) if payload else None
    cycle = db.session.get(InvFixedAssetCycle, cycle_id) if cycle_id is not None else None
    if cycle is None or cycle.status != "ACTIVE" or not _can_scan_cycle(cycle):
        abort(404)
    session["fixed_asset_requested_cycle_id"] = cycle.id
    session.modified = True
    return _register_mobile_device()


@portal_bp.route("/inventory/fixed-assets/mobile/device/status")
@login_required
def fixed_asset_mobile_device_status():
    device = _mobile_device_from_request()
    return render_template(
        "portal/inventory/fixed_assets/mobile_device_status.html",
        device=device,
        status_labels=MOBILE_DEVICE_STATUS_LABELS,
    )


@portal_bp.route("/inventory/fixed-assets/mobile/device/status.json")
@login_required
def fixed_asset_mobile_device_status_json():
    device = _mobile_device_from_request()
    if device is None:
        return jsonify({"ok": False, "message": "الهاتف غير مسجل."}), 404
    return jsonify({
        "ok": True,
        "status": device.status,
        "label": MOBILE_DEVICE_STATUS_LABELS.get(device.status, device.status),
        "mobile_url": url_for("portal.fixed_asset_mobile"),
    })


@portal_bp.route("/inventory/fixed-assets/mobile")
@login_required
def fixed_asset_mobile():
    device = _mobile_device_from_request()
    if device is None or device.status != "APPROVED":
        return redirect(url_for("portal.fixed_asset_mobile_device_status"))
    session["fixed_asset_mobile_mode"] = True
    session["fixed_asset_mobile_device_id"] = device.id
    session.modified = True
    query = InvFixedAssetCycle.query.filter_by(status="ACTIVE")
    if not current_user.has_perm(STORE_MANAGE):
        query = query.filter(InvFixedAssetCycle.members.any(user_id=current_user.id))
    return render_template(
        "portal/inventory/fixed_assets/mobile.html",
        cycles=query.order_by(InvFixedAssetCycle.id.desc()).all(),
        code=request.args.get("code", ""),
    )


@portal_bp.route("/inventory/fixed-assets/mobile-settings")
@login_required
@_perm_any(STORE_MANAGE)
def fixed_asset_mobile_settings():
    invite_token = _mobile_device_invitation_token()
    devices = InvFixedAssetMobileDevice.query.order_by(
        InvFixedAssetMobileDevice.requested_at.desc(),
        InvFixedAssetMobileDevice.id.desc(),
    ).all()
    return render_template(
        "portal/inventory/fixed_assets/mobile_settings.html",
        invite_token=invite_token,
        pair_url=_mobile_device_pairing_target(invite_token),
        pairing_minutes=MOBILE_DEVICE_INVITE_MAX_AGE // 60,
        devices=devices,
        device_status_labels=MOBILE_DEVICE_STATUS_LABELS,
    )


@portal_bp.route("/inventory/fixed-assets/mobile-entry.png")
@login_required
@_perm_any(STORE_MANAGE)
def fixed_asset_mobile_entry_qr():
    invite_token = (request.args.get("token") or "").strip()
    if invite_token and _decode_mobile_device_invitation(invite_token) is None:
        abort(404)
    invite_token = invite_token or _mobile_device_invitation_token()
    target = _mobile_device_pairing_target(invite_token)
    output = io.BytesIO()
    qrcode.make(target).save(output, format="PNG")
    output.seek(0)
    response = send_file(output, mimetype="image/png")
    response.headers["Cache-Control"] = "no-store"
    return response


@portal_bp.route("/inventory/fixed-assets/mobile-devices/<int:device_id>/approve", methods=["POST"])
@login_required
@_perm_any(STORE_MANAGE)
def fixed_asset_mobile_device_approve(device_id: int):
    device = InvFixedAssetMobileDevice.query.get_or_404(device_id)
    device.status = "APPROVED"
    device.approved_at = datetime.utcnow()
    device.approved_by_id = current_user.id
    _audit(
        "INV_FIXED_ASSET_MOBILE_DEVICE_APPROVE",
        f"اعتماد هاتف الجرد #{device.id}",
        "INV_FIXED_ASSET_MOBILE_DEVICE",
        device.id,
    )
    db.session.commit()
    flash("تم اعتماد الهاتف، ويمكنه الآن مسح QR أصول النظام.", "success")
    return redirect(url_for("portal.fixed_asset_mobile_settings"))


@portal_bp.route("/inventory/fixed-assets/mobile-devices/<int:device_id>/revoke", methods=["POST"])
@login_required
@_perm_any(STORE_MANAGE)
def fixed_asset_mobile_device_revoke(device_id: int):
    device = InvFixedAssetMobileDevice.query.get_or_404(device_id)
    device.status = "REVOKED"
    _audit(
        "INV_FIXED_ASSET_MOBILE_DEVICE_REVOKE",
        f"إيقاف هاتف الجرد #{device.id}",
        "INV_FIXED_ASSET_MOBILE_DEVICE",
        device.id,
    )
    db.session.commit()
    flash("تم إيقاف الهاتف.", "success")
    return redirect(url_for("portal.fixed_asset_mobile_settings"))


def _next_asset_tag() -> str:
    prefix = re.sub(r"[^A-Za-z0-9_-]", "", _system_setting("INV_FIXED_ASSET_PREFIX", "FA")) or "FA"
    # Counting is enough to choose a first candidate and avoids materializing
    # every historical tag into Python.  The existence check below handles
    # gaps and preserves the old case-insensitive uniqueness behavior.
    sequence = int(db.session.query(func.count(InvFixedAsset.id)).scalar() or 0) + 1
    while True:
        candidate = f"{prefix.upper()}-{sequence:06d}"
        if not db.session.query(InvFixedAsset.id).filter(
            func.lower(InvFixedAsset.asset_tag) == candidate.casefold()
        ).first():
            return candidate
        sequence += 1


def _next_cycle_code() -> str:
    prefix = f"FA-INV-{date.today().year}"
    sequence = int(db.session.query(func.count(InvFixedAssetCycle.id)).scalar() or 0) + 1
    while True:
        candidate = f"{prefix}-{sequence:03d}"
        if not db.session.query(InvFixedAssetCycle.id).filter(
            func.lower(InvFixedAssetCycle.code) == candidate.casefold()
        ).first():
            return candidate
        sequence += 1


def _asset_values(source) -> dict:
    item_id = _optional_int(source.get("item_id"))
    category_id = _optional_int(source.get("category_id"))
    warehouse_id = _optional_int(source.get("warehouse_id"))
    room_id = _optional_int(source.get("room_id"))
    custodian_user_id = _optional_int(source.get("custodian_user_id"))

    item = _entity_or_none(InvItem, item_id, "الصنف")
    _entity_or_none(InvItemCategory, category_id, "التصنيف")
    _entity_or_none(InvWarehouse, warehouse_id, "المستودع")
    _entity_or_none(InvRoom, room_id, "الغرفة")
    _entity_or_none(User, custodian_user_id, "الموظف المسؤول")

    asset_tag = (_clean_text(source.get("asset_tag"), 80) or _next_asset_tag()).upper()
    name = _clean_text(source.get("name"), 255) or (item.name if item else None)
    if not name:
        raise ValueError("اسم الأصل مطلوب.")
    if category_id is None and item is not None:
        category_id = item.category_id

    condition = (_clean_text(source.get("asset_condition"), 30) or "GOOD").upper()
    lifecycle_status = (_clean_text(source.get("lifecycle_status"), 30) or "ACTIVE").upper()
    if condition not in ASSET_CONDITIONS:
        raise ValueError("حالة الأصل غير صالحة.")
    if lifecycle_status not in ASSET_LIFECYCLE_STATUSES:
        raise ValueError("وضع الأصل غير صالح.")

    return {
        "asset_tag": asset_tag,
        "name": name,
        "item_id": item_id,
        "category_id": category_id,
        "serial_number": _clean_text(source.get("serial_number"), 200),
        "manufacturer": _clean_text(source.get("manufacturer"), 150),
        "model": _clean_text(source.get("model"), 150),
        "description": _clean_text(source.get("description")),
        "acquisition_date": _valid_iso_date(source.get("acquisition_date"), "تاريخ الشراء"),
        "purchase_cost": _decimal_or_none(source.get("purchase_cost")),
        "asset_condition": condition,
        "lifecycle_status": lifecycle_status,
        "warehouse_id": warehouse_id,
        "room_id": room_id,
        "custodian_user_id": custodian_user_id,
        "note": _clean_text(source.get("note")),
        "is_active": bool(source.get("is_active")),
    }


def _apply_asset_values(asset: InvFixedAsset, values: dict):
    if asset.id and InvAssetDocumentLine.query.filter_by(asset_id=asset.id).first():
        if any(values.get(key) != getattr(asset, key) for key in ("custodian_user_id", "lifecycle_status", "is_active")):
            raise ValueError("الأصل مرتبط بمستندات عهدة؛ استخدم حركة عهدة لتغيير الموظف أو إسقاط الأصل.")
    duplicate = (
        InvFixedAsset.query
        .filter(func.lower(InvFixedAsset.asset_tag) == values["asset_tag"].casefold())
        .filter(InvFixedAsset.id != (asset.id or 0))
        .first()
    )
    if duplicate:
        raise ValueError("رقم الأصل مستخدم مسبقاً.")
    for field_name, value in values.items():
        setattr(asset, field_name, value)
    asset.updated_by_id = current_user.id
    asset.updated_at = datetime.utcnow()


def _filtered_assets(args):
    query = InvFixedAsset.query
    search = (args.get("q") or "").strip()
    category_id = _optional_int(args.get("category_id"))
    warehouse_id = _optional_int(args.get("warehouse_id"))
    room_id = _optional_int(args.get("room_id"))
    lifecycle_status = (args.get("status") or "").strip().upper()
    condition = (args.get("condition") or "").strip().upper()
    active = (args.get("active") or "1").strip()

    if search:
        like = f"%{search}%"
        query = query.filter(or_(
            InvFixedAsset.asset_tag.ilike(like),
            InvFixedAsset.name.ilike(like),
            InvFixedAsset.serial_number.ilike(like),
            InvFixedAsset.manufacturer.ilike(like),
            InvFixedAsset.model.ilike(like),
            InvFixedAsset.description.ilike(like),
            InvFixedAsset.note.ilike(like),
            InvFixedAsset.item.has(InvItem.name.ilike(like)),
        ))
    if category_id:
        query = query.filter(InvFixedAsset.category_id == category_id)
    if warehouse_id:
        query = query.filter(InvFixedAsset.warehouse_id == warehouse_id)
    if room_id:
        query = query.filter(InvFixedAsset.room_id == room_id)
    if lifecycle_status in ASSET_LIFECYCLE_STATUSES:
        query = query.filter(InvFixedAsset.lifecycle_status == lifecycle_status)
    if condition in ASSET_CONDITIONS:
        query = query.filter(InvFixedAsset.asset_condition == condition)
    if active == "1":
        query = query.filter(InvFixedAsset.is_active.is_(True))
    elif active == "0":
        query = query.filter(InvFixedAsset.is_active.is_(False))

    selected = {
        "q": search,
        "category_id": category_id,
        "warehouse_id": warehouse_id,
        "room_id": room_id,
        "status": lifecycle_status,
        "condition": condition,
        "active": active,
    }
    return query, selected


def _item_lookup_label(item: InvItem) -> str:
    """Build a compact item label without loading the full attribute list."""
    base = (item.name or "").strip() or (item.code or "").strip() or str(item.id)
    variant = (item.variant or "").strip()
    code = (item.code or "").strip()
    if variant:
        base = f"{base} — {variant}"
    if code and code not in base:
        base = f"{base} ({code})"
    return base


@portal_bp.route("/inventory/fixed-assets/lookups/items")
@login_required
@_perm_any(STORE_READ, STORE_MANAGE)
def fixed_asset_item_lookup():
    """Return a small, server-filtered item list for the asset form."""
    search = (request.args.get("q") or "").strip()
    if len(search) < 2:
        return jsonify([])

    like = f"%{search}%"
    prefix = f"{search}%"
    filters = [
        InvItem.name.ilike(like),
        InvItem.code.ilike(like),
        InvItem.variant.ilike(like),
    ]
    numeric_id = _optional_int(search)
    if numeric_id is not None:
        filters.append(InvItem.id == numeric_id)

    rows = (
        InvItem.query
        .options(noload(InvItem.attributes))
        .filter(InvItem.is_active.is_(True))
        .filter(or_(*filters))
        .order_by(
            case(
                (InvItem.name.ilike(prefix), 0),
                (InvItem.code.ilike(prefix), 1),
                else_=2,
            ),
            InvItem.name.asc(),
            InvItem.id.asc(),
        )
        .limit(30)
        .all()
    )
    return jsonify([
        {"id": item.id, "label": _item_lookup_label(item)}
        for item in rows
    ])


def _style_worksheet(worksheet, widths: dict[int, int] | None = None):
    worksheet.sheet_view.rightToLeft = True
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    widths = widths or {}
    for index in range(1, worksheet.max_column + 1):
        worksheet.column_dimensions[get_column_letter(index)].width = widths.get(index, 18)
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)


def _workbook_response(workbook: Workbook, filename: str):
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )


@portal_bp.route("/inventory/fixed-assets")
@login_required
@_perm_any(STORE_READ, STORE_MANAGE)
def fixed_assets_dashboard():
    total_assets = InvFixedAsset.query.count()
    active_assets = InvFixedAsset.query.filter(InvFixedAsset.is_active.is_(True)).count()
    status_counts = {
        status: count
        for status, count in (
            db.session.query(InvFixedAsset.lifecycle_status, func.count(InvFixedAsset.id))
            .filter(InvFixedAsset.is_active.is_(True))
            .group_by(InvFixedAsset.lifecycle_status)
            .all()
        )
    }
    active_cycles = (
        InvFixedAssetCycle.query
        .options(noload(InvFixedAssetCycle.entries), noload(InvFixedAssetCycle.manual_previous_lines))
        .filter(InvFixedAssetCycle.status == "ACTIVE")
        .order_by(InvFixedAssetCycle.inventory_date.desc(), InvFixedAssetCycle.id.desc())
        .all()
    )
    active_cycle_counts = cycle_result_counts_for_cycles(cycle.id for cycle in active_cycles)
    cycle_cards = [(cycle, active_cycle_counts[cycle.id]) for cycle in active_cycles]
    recent_cycles = (
        InvFixedAssetCycle.query
        .options(noload(InvFixedAssetCycle.entries), noload(InvFixedAssetCycle.manual_previous_lines))
        .order_by(InvFixedAssetCycle.inventory_date.desc(), InvFixedAssetCycle.id.desc())
        .limit(8)
        .all()
    )
    recent_scans = (
        InvFixedAssetScanLog.query
        .order_by(InvFixedAssetScanLog.scanned_at.desc())
        .limit(10)
        .all()
    )
    return render_template(
        "portal/inventory/fixed_assets/dashboard.html",
        total_assets=total_assets,
        active_assets=active_assets,
        inactive_assets=total_assets - active_assets,
        status_counts=status_counts,
        cycle_cards=cycle_cards,
        recent_cycles=recent_cycles,
        recent_scans=recent_scans,
        can_manage=current_user.has_perm(STORE_MANAGE),
        can_export=current_user.has_perm(STORE_EXPORT) or current_user.has_perm(STORE_MANAGE),
    )


@portal_bp.route("/inventory/fixed-assets/assets")
@login_required
@_perm_any(STORE_READ, STORE_MANAGE)
def fixed_asset_list():
    query, selected = _filtered_assets(request.args)
    page = max(_optional_int(request.args.get("page")) or 1, 1)
    pagination = query.order_by(InvFixedAsset.asset_tag.asc(), InvFixedAsset.id.asc()).paginate(
        page=page,
        per_page=50,
        error_out=False,
    )
    return render_template(
        "portal/inventory/fixed_assets/asset_list.html",
        rows=pagination.items,
        pagination=pagination,
        selected=selected,
        can_manage=current_user.has_perm(STORE_MANAGE),
        can_export=current_user.has_perm(STORE_EXPORT) or current_user.has_perm(STORE_MANAGE),
        **_choices(include_items=False, include_users=False),
    )


@portal_bp.route("/inventory/fixed-assets/assets/new", methods=["GET", "POST"])
@login_required
@_perm_any(STORE_MANAGE)
def fixed_asset_new():
    asset = InvFixedAsset(
        asset_tag=_next_asset_tag(),
        asset_condition="GOOD",
        lifecycle_status="ACTIVE",
        is_active=True,
    )
    new_image_filename = None
    if request.method == "POST":
        try:
            values = _asset_values(request.form)
            _apply_asset_values(asset, values)
            new_image_filename = _save_fixed_asset_image(request.files.get("asset_image"))
            if new_image_filename:
                asset.image_filename = new_image_filename
            asset.created_by_id = current_user.id
            asset.created_at = datetime.utcnow()
            db.session.add(asset)
            db.session.flush()
            _audit("INV_FIXED_ASSET_CREATE", f"إضافة أصل ثابت {asset.asset_tag}", "INV_FIXED_ASSET", asset.id)
            db.session.commit()
            flash("تمت إضافة الأصل وإنشاء رمز QR الخاص به.", "success")
            return redirect(url_for("portal.fixed_asset_view", asset_id=asset.id))
        except (ValueError, IntegrityError) as exc:
            db.session.rollback()
            if new_image_filename:
                _delete_fixed_asset_image(new_image_filename)
            message = str(exc) if isinstance(exc, ValueError) else "تعذر الحفظ بسبب رقم أصل مكرر."
            flash(message, "danger")
            asset = InvFixedAsset(**{
                key: value for key, value in request.form.items()
                if key in {"asset_tag", "name", "serial_number", "manufacturer", "model", "description", "acquisition_date", "note"}
            })
    selected_item_id = (
        _optional_int(request.form.get("item_id"))
        if request.method == "POST"
        else asset.item_id
    )
    selected_item = _selected_item(selected_item_id)
    return render_template(
        "portal/inventory/fixed_assets/asset_form.html",
        asset=asset,
        selected_item=selected_item,
        is_new=True,
        **_choices(include_items=False),
    )


@portal_bp.route("/inventory/fixed-assets/assets/<int:asset_id>")
@login_required
def fixed_asset_view(asset_id: int):
    asset = InvFixedAsset.query.get_or_404(asset_id)
    _require_asset_access(asset)
    history = (
        InvFixedAssetEntry.query
        .filter(InvFixedAssetEntry.asset_id == asset.id)
        .order_by(InvFixedAssetEntry.id.desc())
        .all()
    )
    active_cycles = (
        InvFixedAssetCycle.query
        .filter(InvFixedAssetCycle.status == "ACTIVE")
        .order_by(InvFixedAssetCycle.inventory_date.desc(), InvFixedAssetCycle.id.desc())
        .all()
    )
    return render_template(
        "portal/inventory/fixed_assets/asset_view.html",
        asset=asset,
        history=history,
        active_cycles=active_cycles,
        can_manage=current_user.has_perm(STORE_MANAGE),
        custody_history=InvAssetDocument.query.join(InvAssetDocumentLine).filter(InvAssetDocumentLine.asset_id == asset.id).order_by(InvAssetDocument.id.desc()).all(),
    )


@portal_bp.route("/inventory/fixed-assets/assets/<int:asset_id>/image")
@login_required
def fixed_asset_image(asset_id: int):
    asset = InvFixedAsset.query.get_or_404(asset_id)
    _require_asset_access(asset)
    path = _fixed_asset_image_path(asset.image_filename)
    if path is None or not path.is_file():
        abort(404)
    response = send_file(path, conditional=True)
    response.headers["Cache-Control"] = "private, max-age=86400"
    return response


@portal_bp.route("/inventory/fixed-assets/assets/<int:asset_id>/edit", methods=["GET", "POST"])
@login_required
@_perm_any(STORE_MANAGE)
def fixed_asset_edit(asset_id: int):
    asset = InvFixedAsset.query.get_or_404(asset_id)
    old_image_filename = asset.image_filename
    new_image_filename = None
    if request.method == "POST":
        try:
            old_tag = asset.asset_tag
            _apply_asset_values(asset, _asset_values(request.form))
            new_image_filename = _save_fixed_asset_image(request.files.get("asset_image"))
            if new_image_filename:
                asset.image_filename = new_image_filename
            elif request.form.get("remove_image"):
                asset.image_filename = None
            _audit(
                "INV_FIXED_ASSET_UPDATE",
                f"تعديل الأصل {old_tag} إلى {asset.asset_tag}",
                "INV_FIXED_ASSET",
                asset.id,
            )
            db.session.commit()
            if old_image_filename and old_image_filename != asset.image_filename:
                _delete_fixed_asset_image(old_image_filename)
            flash("تم تحديث بيانات الأصل.", "success")
            return redirect(url_for("portal.fixed_asset_view", asset_id=asset.id))
        except (ValueError, IntegrityError) as exc:
            db.session.rollback()
            if new_image_filename:
                _delete_fixed_asset_image(new_image_filename)
            flash(str(exc) if isinstance(exc, ValueError) else "رقم الأصل مستخدم مسبقاً.", "danger")
    selected_item_id = (
        _optional_int(request.form.get("item_id"))
        if request.method == "POST"
        else asset.item_id
    )
    selected_item = _selected_item(selected_item_id)
    return render_template(
        "portal/inventory/fixed_assets/asset_form.html",
        asset=asset,
        selected_item=selected_item,
        is_new=False,
        **_choices(include_items=False),
    )


@portal_bp.route("/inventory/fixed-assets/assets/<int:asset_id>/toggle", methods=["POST"])
@login_required
@_perm_any(STORE_MANAGE)
def fixed_asset_toggle(asset_id: int):
    asset = InvFixedAsset.query.get_or_404(asset_id)
    if InvAssetDocumentLine.query.filter_by(asset_id=asset.id).first():
        flash("الأصل مرتبط بعهدة موثقة؛ استخدم حركة إسقاط العهدة.", "warning")
        return redirect(url_for("portal.fixed_asset_view", asset_id=asset.id))
    asset.is_active = not asset.is_active
    asset.updated_by_id = current_user.id
    asset.updated_at = datetime.utcnow()
    _audit(
        "INV_FIXED_ASSET_TOGGLE",
        f"{'تفعيل' if asset.is_active else 'أرشفة'} الأصل {asset.asset_tag}",
        "INV_FIXED_ASSET",
        asset.id,
    )
    db.session.commit()
    flash("تم تحديث حالة السجل.", "success")
    return redirect(request.referrer or url_for("portal.fixed_asset_list"))


@portal_bp.route("/inventory/fixed-assets/assets/<int:asset_id>/qr.png")
@login_required
def fixed_asset_qr_png(asset_id: int):
    asset = InvFixedAsset.query.get_or_404(asset_id)
    _require_asset_access(asset)
    target = _asset_qr_target(asset)
    qr_code = qrcode.QRCode(
        version=None,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=10,
        border=3,
    )
    qr_code.add_data(target)
    qr_code.make(fit=True)
    image = qr_code.make_image(fill_color="black", back_color="white")
    output = io.BytesIO()
    image.save(output, format="PNG")
    output.seek(0)
    response = send_file(output, mimetype="image/png", download_name=f"{asset.asset_tag}.png")
    response.headers["Cache-Control"] = "private, max-age=86400"
    return response


@portal_bp.route("/inventory/fixed-assets/qr/<token>")
@login_required
def fixed_asset_qr_resolve(token: str):
    asset = _find_asset_by_system_qr(token)
    if asset is None:
        abort(404)
    _require_asset_access(asset)
    cycle_id = _optional_int(session.get("fixed_asset_active_cycle_id"))
    cycle = db.session.get(InvFixedAssetCycle, cycle_id) if cycle_id else None
    if cycle and cycle.status == "ACTIVE" and _can_scan_cycle(cycle):
        if session.get("fixed_asset_mobile_mode"):
            _require_approved_mobile_device()
            return redirect(url_for(
                "portal.fixed_asset_cycle_scan",
                cycle_id=cycle.id,
                code=asset.qr_token,
                mobile="1",
            ))
        return redirect(url_for("portal.fixed_asset_cycle_scan", cycle_id=cycle.id, code=asset.qr_token))
    if current_user.has_perm(STORE_MANAGE) or InvFixedAssetCycleMember.query.filter_by(user_id=current_user.id).first():
        return redirect(url_for("portal.fixed_asset_mobile", code=asset.qr_token))
    return redirect(url_for("portal.fixed_asset_view", asset_id=asset.id))


@portal_bp.route("/inventory/fixed-assets/labels")
@login_required
@_perm_any(STORE_READ, STORE_MANAGE)
def fixed_asset_labels():
    raw_ids = (request.args.get("ids") or "").strip()
    if raw_ids:
        asset_ids = [int(value) for value in raw_ids.split(",") if value.strip().isdigit()]
        assets = InvFixedAsset.query.filter(InvFixedAsset.id.in_(asset_ids)).order_by(InvFixedAsset.asset_tag.asc()).all()
    else:
        query, _ = _filtered_assets(request.args)
        assets = query.order_by(InvFixedAsset.asset_tag.asc()).limit(500).all()
    if not assets:
        flash("لا توجد أصول لطباعة الملصقات.", "warning")
        return redirect(url_for("portal.fixed_asset_list"))
    return render_template(
        "portal/inventory/fixed_assets/labels.html",
        assets=assets,
        organization_name=_system_setting("ORGANIZATION_NAME", "اللجنة الوطنية الفلسطينية للتربية والثقافة والعلوم"),
    )


@portal_bp.route("/inventory/fixed-assets/assets/export.xlsx")
@login_required
@_perm_any(STORE_EXPORT, STORE_MANAGE)
def fixed_asset_export():
    query, _ = _filtered_assets(request.args)
    assets = query.order_by(InvFixedAsset.asset_tag.asc()).all()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "الأصول الثابتة"
    headers = [title for _, title in IMPORT_HEADERS] + ["رابط QR", "فعال", "تاريخ الإضافة", "آخر تحديث"]
    sheet.append(headers)
    for asset in assets:
        sheet.append([
            asset.asset_tag,
            asset.name,
            asset.item.code if asset.item and asset.item.code else (asset.item.name if asset.item else ""),
            asset.category.name if asset.category else "",
            asset.serial_number or "",
            asset.manufacturer or "",
            asset.model or "",
            asset.description or "",
            asset.acquisition_date or "",
            float(asset.purchase_cost) if asset.purchase_cost is not None else "",
            ASSET_CONDITIONS.get(asset.asset_condition, asset.asset_condition),
            ASSET_LIFECYCLE_STATUSES.get(asset.lifecycle_status, asset.lifecycle_status),
            asset.warehouse.label if asset.warehouse else "",
            asset.room.label if asset.room else "",
            asset.custodian.full_name if asset.custodian else "",
            asset.note or "",
            _asset_qr_target(asset),
            "نعم" if asset.is_active else "لا",
            asset.created_at,
            asset.updated_at,
        ])
    _style_worksheet(sheet, {1: 18, 2: 28, 8: 35, 16: 35, 17: 55})
    return _workbook_response(workbook, f"fixed-assets-{date.today().isoformat()}.xlsx")


def _normalize_header(value) -> str:
    text_value = str(value or "").strip().casefold()
    return re.sub(r"[^\w\u0600-\u06ff]", "", text_value)


def _read_import_rows(uploaded_file) -> tuple[dict[int, str], list[tuple[int, list]]]:
    suffix = Path(uploaded_file.filename or "").suffix.lower()
    if suffix == ".xlsx":
        workbook = load_workbook(uploaded_file.stream, read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        header_values = next(iterator, None)
        if not header_values:
            return {}, []
        rows = [(row_number, list(row)) for row_number, row in enumerate(iterator, start=2)]
    elif suffix == ".csv":
        content = uploaded_file.read()
        try:
            decoded = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            decoded = content.decode("cp1256")
        iterator = csv.reader(io.StringIO(decoded))
        header_values = next(iterator, None)
        if not header_values:
            return {}, []
        rows = [(row_number, list(row)) for row_number, row in enumerate(iterator, start=2)]
    else:
        raise ValueError("الصيغ المدعومة هي XLSX وCSV فقط.")

    alias_to_key = {
        alias: key
        for key, aliases in IMPORT_ALIASES.items()
        for alias in aliases
    }
    column_map = {}
    for index, raw_header in enumerate(header_values):
        normalized = _normalize_header(raw_header)
        key = alias_to_key.get(normalized)
        if key:
            column_map[index] = key
    return column_map, rows


def _coded_import_value(value, choices: dict[str, str], field_label: str, default: str) -> str:
    if value in (None, ""):
        return default
    text_value = str(value).strip()
    upper_value = text_value.upper()
    if upper_value in choices:
        return upper_value
    for code, label in choices.items():
        if text_value.casefold() == label.casefold():
            return code
    raise ValueError(f"{field_label} غير معروفة: {text_value}")


def _find_named_entity(model, value, *columns):
    text_value = str(value or "").strip()
    if not text_value:
        return None
    conditions = [func.lower(column) == text_value.casefold() for column in columns]
    return model.query.filter(or_(*conditions)).first()


def _find_import_user(value):
    text_value = str(value or "").strip()
    if not text_value:
        return None
    folded = text_value.casefold()
    return User.query.filter(or_(
        func.lower(User.username) == folded,
        func.lower(User.email) == folded,
        User.name == text_value,
    )).first()


@portal_bp.route("/inventory/fixed-assets/assets/import", methods=["GET", "POST"])
@login_required
@_perm_any(STORE_MANAGE)
def fixed_asset_import():
    result = None
    if request.method == "POST":
        uploaded_file = request.files.get("file")
        if not uploaded_file or not uploaded_file.filename:
            flash("اختر ملفاً للاستيراد.", "warning")
            return redirect(url_for("portal.fixed_asset_import"))
        if request.content_length and request.content_length > 10 * 1024 * 1024:
            flash("حجم الملف يتجاوز 10 ميغابايت.", "danger")
            return redirect(url_for("portal.fixed_asset_import"))

        update_existing = bool(request.form.get("update_existing"))
        create_categories = bool(request.form.get("create_categories"))
        created = 0
        updated = 0
        errors = []
        seen_tags = set()
        try:
            column_map, rows = _read_import_rows(uploaded_file)
            if "name" not in column_map.values() and "item" not in column_map.values():
                raise ValueError("يجب أن يحتوي الملف على عمود اسم الأصل أو الصنف.")
            if len(rows) > 5000:
                raise ValueError("الحد الأقصى للاستيراد هو 5000 صف في المرة الواحدة.")

            for row_number, raw_row in rows:
                row_data = {
                    key: raw_row[index] if index < len(raw_row) else None
                    for index, key in column_map.items()
                }
                if not any(value not in (None, "") for value in row_data.values()):
                    continue
                try:
                    raw_tag = _clean_text(row_data.get("asset_tag"), 80)
                    asset_tag = (raw_tag or _next_asset_tag()).upper()
                    if asset_tag.casefold() in seen_tags:
                        raise ValueError("رقم الأصل مكرر داخل الملف.")
                    seen_tags.add(asset_tag.casefold())

                    item = None
                    item_value = row_data.get("item")
                    if item_value not in (None, ""):
                        item = _find_named_entity(InvItem, item_value, InvItem.code, InvItem.name)
                        if item is None:
                            raise ValueError(f"الصنف غير موجود: {item_value}")

                    category = None
                    category_value = row_data.get("category")
                    if category_value not in (None, ""):
                        category = _find_named_entity(InvItemCategory, category_value, InvItemCategory.name)
                        if category is None and create_categories:
                            category = InvItemCategory(
                                name=str(category_value).strip(),
                                is_active=True,
                                created_at=datetime.utcnow(),
                            )
                            db.session.add(category)
                            db.session.flush()
                        elif category is None:
                            raise ValueError(f"التصنيف غير موجود: {category_value}")
                    elif item and item.category:
                        category = item.category

                    warehouse = None
                    if row_data.get("warehouse") not in (None, ""):
                        warehouse = _find_named_entity(
                            InvWarehouse,
                            row_data["warehouse"],
                            InvWarehouse.code,
                            InvWarehouse.name,
                        )
                        if warehouse is None:
                            raise ValueError(f"المستودع غير موجود: {row_data['warehouse']}")

                    room = None
                    if row_data.get("room") not in (None, ""):
                        room = _find_named_entity(InvRoom, row_data["room"], InvRoom.code, InvRoom.name)
                        if room is None:
                            raise ValueError(f"الغرفة غير موجودة: {row_data['room']}")

                    custodian = None
                    if row_data.get("custodian") not in (None, ""):
                        custodian = _find_import_user(row_data["custodian"])
                        if custodian is None:
                            raise ValueError(f"الموظف غير موجود: {row_data['custodian']}")

                    name = _clean_text(row_data.get("name"), 255) or (item.name if item else None)
                    if not name:
                        raise ValueError("اسم الأصل مطلوب.")

                    values = {
                        "asset_tag": asset_tag,
                        "name": name,
                        "item_id": item.id if item else None,
                        "category_id": category.id if category else None,
                        "serial_number": _clean_text(row_data.get("serial_number"), 200),
                        "manufacturer": _clean_text(row_data.get("manufacturer"), 150),
                        "model": _clean_text(row_data.get("model"), 150),
                        "description": _clean_text(row_data.get("description")),
                        "acquisition_date": _valid_iso_date(row_data.get("acquisition_date"), "تاريخ الشراء"),
                        "purchase_cost": _decimal_or_none(row_data.get("purchase_cost")),
                        "asset_condition": _coded_import_value(row_data.get("asset_condition"), ASSET_CONDITIONS, "حالة الأصل", "GOOD"),
                        "lifecycle_status": _coded_import_value(row_data.get("lifecycle_status"), ASSET_LIFECYCLE_STATUSES, "وضع الأصل", "ACTIVE"),
                        "warehouse_id": warehouse.id if warehouse else None,
                        "room_id": room.id if room else None,
                        "custodian_user_id": custodian.id if custodian else None,
                        "note": _clean_text(row_data.get("note")),
                        "is_active": True,
                    }
                    asset = InvFixedAsset.query.filter(func.lower(InvFixedAsset.asset_tag) == asset_tag.casefold()).first()
                    if asset and not update_existing:
                        raise ValueError("رقم الأصل موجود مسبقاً؛ فعّل خيار تحديث السجلات الموجودة.")
                    if asset:
                        _apply_asset_values(asset, values)
                        updated += 1
                    else:
                        asset = InvFixedAsset(
                            created_by_id=current_user.id,
                            created_at=datetime.utcnow(),
                        )
                        _apply_asset_values(asset, values)
                        db.session.add(asset)
                        created += 1
                except ValueError as exc:
                    errors.append({"row": row_number, "message": str(exc)})

            if created or updated:
                db.session.flush()
                _audit(
                    "INV_FIXED_ASSET_IMPORT",
                    f"استيراد أصول ثابتة: جديد {created}، محدث {updated}، أخطاء {len(errors)}",
                    "INV_FIXED_ASSET",
                    None,
                )
                db.session.commit()
            else:
                db.session.rollback()
            result = {"created": created, "updated": updated, "errors": errors}
            if created or updated:
                flash("اكتمل الاستيراد.", "success")
        except (ValueError, IntegrityError) as exc:
            db.session.rollback()
            flash(str(exc) if isinstance(exc, ValueError) else "تعذر الاستيراد بسبب بيانات مكررة.", "danger")

    return render_template("portal/inventory/fixed_assets/import.html", result=result)


@portal_bp.route("/inventory/fixed-assets/assets/import/template.xlsx")
@login_required
@_perm_any(STORE_MANAGE)
def fixed_asset_import_template():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "الأصول الثابتة"
    sheet.append([title for _, title in IMPORT_HEADERS])
    sheet.append([
        "FA-000001",
        "كرسي مكتب",
        "",
        "أثاث",
        "CH-1001",
        "",
        "",
        "كرسي دوار",
        date.today().isoformat(),
        250,
        "جيد",
        "قيد الاستخدام",
        "المستودع الرئيسي",
        "غرفة الإدارة",
        "employee@example.org",
        "صف تجريبي يُحذف قبل الاستيراد",
    ])
    _style_worksheet(sheet, {1: 18, 2: 28, 8: 35, 16: 38})

    lists = workbook.create_sheet("قوائم")
    for row_index, label in enumerate(ASSET_CONDITIONS.values(), start=1):
        lists.cell(row=row_index, column=1, value=label)
    for row_index, label in enumerate(ASSET_LIFECYCLE_STATUSES.values(), start=1):
        lists.cell(row=row_index, column=2, value=label)
    condition_validation = DataValidation(
        type="list",
        formula1=f"'قوائم'!$A$1:$A${len(ASSET_CONDITIONS)}",
        allow_blank=True,
    )
    lifecycle_validation = DataValidation(
        type="list",
        formula1=f"'قوائم'!$B$1:$B${len(ASSET_LIFECYCLE_STATUSES)}",
        allow_blank=True,
    )
    sheet.add_data_validation(condition_validation)
    sheet.add_data_validation(lifecycle_validation)
    condition_validation.add("K2:K5001")
    lifecycle_validation.add("L2:L5001")
    lists.sheet_state = "hidden"

    instructions = workbook.create_sheet("تعليمات")
    instructions.sheet_view.rightToLeft = True
    instructions.append(["تعليمات الاستيراد"])
    instructions.append(["رقم الأصل اختياري؛ إذا تُرك فارغاً ينشئ النظام رقماً تلقائياً."])
    instructions.append(["يجب أن تطابق أسماء المستودعات والغرف والموظفين البيانات المسجلة في النظام."])
    instructions.append(["يمكن تفعيل إنشاء التصنيفات المفقودة من شاشة الاستيراد."])
    instructions.column_dimensions["A"].width = 95
    instructions["A1"].font = Font(bold=True, size=14)
    return _workbook_response(workbook, "fixed-assets-import-template.xlsx")


@portal_bp.route("/inventory/fixed-assets/cycles")
@login_required
@_perm_any(STORE_READ, STORE_MANAGE)
def fixed_asset_cycles():
    status = (request.args.get("status") or "").strip().upper()
    search = (request.args.get("q") or "").strip()
    query = InvFixedAssetCycle.query
    if status in CYCLE_STATUSES:
        query = query.filter(InvFixedAssetCycle.status == status)
    if search:
        like = f"%{search}%"
        query = query.filter(or_(
            InvFixedAssetCycle.code.ilike(like),
            InvFixedAssetCycle.name.ilike(like),
            InvFixedAssetCycle.committee_name.ilike(like),
        ))
    cycles = query.options(
        noload(InvFixedAssetCycle.entries),
        noload(InvFixedAssetCycle.manual_previous_lines),
    ).order_by(
        InvFixedAssetCycle.inventory_date.desc(), InvFixedAssetCycle.id.desc()
    ).limit(250).all()
    cycle_counts = cycle_result_counts_for_cycles(cycle.id for cycle in cycles)
    rows = [(cycle, cycle_counts[cycle.id]) for cycle in cycles]
    return render_template(
        "portal/inventory/fixed_assets/cycle_list.html",
        rows=rows,
        selected={"q": search, "status": status},
        can_manage=current_user.has_perm(STORE_MANAGE),
    )


@portal_bp.route("/inventory/fixed-assets/cycles/new", methods=["GET", "POST"])
@login_required
@_perm_any(STORE_MANAGE)
def fixed_asset_cycle_new():
    form_values = {
        "code": _next_cycle_code(),
        "name": f"الجرد السنوي للأصول الثابتة {date.today().year}",
        "inventory_date": date.today().isoformat(),
        "committee_id": None,
        "committee_name": "لجنة الجرد",
        "previous_cycle_id": None,
        "scope_warehouse_id": None,
        "scope_room_id": None,
        "chair_user_id": None,
        "member_ids": [],
        "note": "",
    }
    previous_cycles = (
        InvFixedAssetCycle.query
        .options(noload(InvFixedAssetCycle.entries), noload(InvFixedAssetCycle.manual_previous_lines))
        .filter(InvFixedAssetCycle.status == "CLOSED")
        .order_by(InvFixedAssetCycle.inventory_date.desc(), InvFixedAssetCycle.id.desc())
        .all()
    )
    if previous_cycles:
        form_values["previous_cycle_id"] = previous_cycles[0].id

    if request.method == "POST":
        form_values.update({
            "code": _clean_text(request.form.get("code"), 80) or _next_cycle_code(),
            "name": _clean_text(request.form.get("name"), 255),
            "inventory_date": request.form.get("inventory_date") or "",
            "committee_id": _optional_int(request.form.get("committee_id")),
            "committee_name": _clean_text(request.form.get("committee_name"), 255),
            "previous_cycle_id": _optional_int(request.form.get("previous_cycle_id")),
            "scope_warehouse_id": _optional_int(request.form.get("scope_warehouse_id")),
            "scope_room_id": _optional_int(request.form.get("scope_room_id")),
            "chair_user_id": _optional_int(request.form.get("chair_user_id")),
            "member_ids": [_optional_int(value) for value in request.form.getlist("member_ids")],
            "note": _clean_text(request.form.get("note")) or "",
        })
        try:
            if not form_values["name"]:
                raise ValueError("اسم دورة الجرد مطلوب.")
            inventory_date = _valid_iso_date(form_values["inventory_date"], "تاريخ الجرد")
            _entity_or_none(InvWarehouse, form_values["scope_warehouse_id"], "المستودع")
            _entity_or_none(InvRoom, form_values["scope_room_id"], "الغرفة")
            previous_cycle = _entity_or_none(InvFixedAssetCycle, form_values["previous_cycle_id"], "الدورة السابقة")
            committee = _entity_or_none(Committee, form_values["committee_id"], "لجنة الجرد")
            _entity_or_none(User, form_values["chair_user_id"], "رئيس اللجنة")

            duplicate = InvFixedAssetCycle.query.filter(
                func.lower(InvFixedAssetCycle.code) == form_values["code"].casefold()
            ).first()
            if duplicate:
                raise ValueError("رقم دورة الجرد مستخدم مسبقاً.")

            cycle = InvFixedAssetCycle(
                code=form_values["code"].upper(),
                name=form_values["name"],
                inventory_date=inventory_date,
                committee_id=committee.id if committee else None,
                committee_name=form_values["committee_name"],
                previous_cycle_id=previous_cycle.id if previous_cycle else None,
                scope_warehouse_id=form_values["scope_warehouse_id"],
                scope_room_id=form_values["scope_room_id"],
                status="ACTIVE",
                note=form_values["note"] or None,
                started_at=datetime.utcnow(),
                created_by_id=current_user.id,
                created_at=datetime.utcnow(),
            )
            db.session.add(cycle)
            db.session.flush()

            member_ids = {member_id for member_id in form_values["member_ids"] if member_id}
            chair_user_id = form_values["chair_user_id"]
            if chair_user_id:
                member_ids.add(chair_user_id)
            for member_id in sorted(member_ids):
                if db.session.get(User, member_id) is None:
                    raise ValueError(f"عضو اللجنة رقم {member_id} غير موجود.")
                db.session.add(InvFixedAssetCycleMember(
                    cycle_id=cycle.id,
                    user_id=member_id,
                    member_role="CHAIR" if member_id == chair_user_id else "MEMBER",
                    created_at=datetime.utcnow(),
                ))

            entry_count = populate_cycle_entries(cycle)
            _audit(
                "INV_FIXED_ASSET_CYCLE_CREATE",
                f"إنشاء دورة جرد {cycle.code} بعدد {entry_count} أصل",
                "INV_FIXED_ASSET_CYCLE",
                cycle.id,
            )
            db.session.commit()
            flash(f"بدأت دورة الجرد وتم تثبيت خط أساس لعدد {entry_count} أصل.", "success")
            return redirect(url_for("portal.fixed_asset_cycle_view", cycle_id=cycle.id))
        except (ValueError, IntegrityError) as exc:
            db.session.rollback()
            flash(str(exc) if isinstance(exc, ValueError) else "تعذر إنشاء الدورة بسبب بيانات مكررة.", "danger")

    choices = _choices(include_items=False, include_categories=False)
    return render_template(
        "portal/inventory/fixed_assets/cycle_form.html",
        form=form_values,
        previous_cycles=previous_cycles,
        committees=Committee.query.filter_by(is_active=True).order_by(Committee.name_ar.asc()).all(),
        is_edit=False,
        **choices,
    )


@portal_bp.route("/inventory/fixed-assets/cycles/<int:cycle_id>/edit", methods=["GET", "POST"])
@login_required
@_perm_any(STORE_MANAGE)
def fixed_asset_cycle_edit(cycle_id: int):
    cycle = InvFixedAssetCycle.query.options(
        noload(InvFixedAssetCycle.entries),
        noload(InvFixedAssetCycle.manual_previous_lines),
    ).get_or_404(cycle_id)
    if cycle.finalized_at is not None:
        flash("اعتمد الأمين العام تقرير هذه الدورة؛ لا يمكن تعديلها.", "warning")
        return redirect(url_for("portal.fixed_asset_cycle_view", cycle_id=cycle.id))
    chair = next((member for member in cycle.members if member.member_role == "CHAIR"), None)
    form_values = {
        "code": cycle.code,
        "name": cycle.name,
        "inventory_date": cycle.inventory_date,
        "committee_id": cycle.committee_id,
        "committee_name": cycle.committee_name or "",
        "previous_cycle_id": cycle.previous_cycle_id,
        "scope_warehouse_id": cycle.scope_warehouse_id,
        "scope_room_id": cycle.scope_room_id,
        "chair_user_id": chair.user_id if chair else None,
        "member_ids": [member.user_id for member in cycle.members if member.member_role != "CHAIR"],
        "note": cycle.note or "",
    }
    if request.method == "POST":
        form_values.update({
            "code": _clean_text(request.form.get("code"), 80),
            "name": _clean_text(request.form.get("name"), 255),
            "inventory_date": request.form.get("inventory_date") or "",
            "committee_id": _optional_int(request.form.get("committee_id")),
            "committee_name": _clean_text(request.form.get("committee_name"), 255),
            "chair_user_id": _optional_int(request.form.get("chair_user_id")),
            "member_ids": [_optional_int(value) for value in request.form.getlist("member_ids")],
            "note": _clean_text(request.form.get("note")) or "",
        })
        try:
            if not form_values["code"] or not form_values["name"]:
                raise ValueError("رقم الدورة واسمها مطلوبان.")
            inventory_date = _valid_iso_date(form_values["inventory_date"], "تاريخ الجرد")
            duplicate = (
                InvFixedAssetCycle.query
                .filter(func.lower(InvFixedAssetCycle.code) == form_values["code"].casefold())
                .filter(InvFixedAssetCycle.id != cycle.id)
                .first()
            )
            if duplicate:
                raise ValueError("رقم دورة الجرد مستخدم مسبقاً.")

            cycle.code = form_values["code"].upper()
            cycle.name = form_values["name"]
            cycle.inventory_date = inventory_date
            committee = _entity_or_none(Committee, form_values["committee_id"], "لجنة الجرد")
            cycle.committee_id = committee.id if committee else None
            cycle.committee_name = form_values["committee_name"]
            cycle.note = form_values["note"] or None
            cycle.members.clear()
            db.session.flush()

            member_ids = {member_id for member_id in form_values["member_ids"] if member_id}
            chair_user_id = form_values["chair_user_id"]
            if chair_user_id:
                member_ids.add(chair_user_id)
            for member_id in sorted(member_ids):
                if db.session.get(User, member_id) is None:
                    raise ValueError(f"عضو اللجنة رقم {member_id} غير موجود.")
                cycle.members.append(InvFixedAssetCycleMember(
                    user_id=member_id,
                    member_role="CHAIR" if member_id == chair_user_id else "MEMBER",
                    created_at=datetime.utcnow(),
                ))
            _audit(
                "INV_FIXED_ASSET_CYCLE_UPDATE",
                f"تحديث بيانات دورة الجرد {cycle.code}",
                "INV_FIXED_ASSET_CYCLE",
                cycle.id,
            )
            db.session.commit()
            flash("تم تحديث بيانات دورة الجرد واللجنة.", "success")
            return redirect(url_for("portal.fixed_asset_cycle_view", cycle_id=cycle.id))
        except (ValueError, IntegrityError) as exc:
            db.session.rollback()
            flash(str(exc) if isinstance(exc, ValueError) else "تعذر تحديث الدورة.", "danger")

    previous_cycles = (
        InvFixedAssetCycle.query
        .options(noload(InvFixedAssetCycle.entries), noload(InvFixedAssetCycle.manual_previous_lines))
        .filter(InvFixedAssetCycle.status == "CLOSED")
        .filter(InvFixedAssetCycle.id != cycle.id)
        .order_by(InvFixedAssetCycle.inventory_date.desc(), InvFixedAssetCycle.id.desc())
        .all()
    )
    return render_template(
        "portal/inventory/fixed_assets/cycle_form.html",
        form=form_values,
        previous_cycles=previous_cycles,
        committees=Committee.query.filter_by(is_active=True).order_by(Committee.name_ar.asc()).all(),
        is_edit=True,
        **_choices(include_items=False, include_categories=False),
    )


@portal_bp.route("/inventory/fixed-assets/cycles/<int:cycle_id>")
@login_required
def fixed_asset_cycle_view(cycle_id: int):
    cycle = InvFixedAssetCycle.query.options(
        noload(InvFixedAssetCycle.entries),
        noload(InvFixedAssetCycle.manual_previous_lines),
    ).get_or_404(cycle_id)
    if not (current_user.has_perm(STORE_READ) or _can_scan_cycle(cycle)):
        abort(403)
    status = (request.args.get("status") or "").strip().upper()
    search = (request.args.get("q") or "").strip()
    query = (
        InvFixedAssetEntry.query
        .options(*_cycle_entry_load_options())
        .filter(InvFixedAssetEntry.cycle_id == cycle.id)
        .join(InvFixedAsset)
    )
    if status in INVENTORY_RESULT_LABELS:
        query = query.filter(InvFixedAssetEntry.result_status == status)
    if search:
        like = f"%{search}%"
        query = query.filter(or_(
            InvFixedAsset.asset_tag.ilike(like),
            InvFixedAsset.name.ilike(like),
            InvFixedAsset.serial_number.ilike(like),
        ))
    page = max(_optional_int(request.args.get("page")) or 1, 1)
    pagination = query.order_by(InvFixedAsset.asset_tag.asc()).paginate(
        page=page,
        per_page=60,
        error_out=False,
    )
    entries = pagination.items
    recent_scans = (
        InvFixedAssetScanLog.query
        .options(*_cycle_scan_load_options())
        .filter(InvFixedAssetScanLog.cycle_id == cycle.id)
        .order_by(InvFixedAssetScanLog.scanned_at.desc())
        .limit(15)
        .all()
    )
    return render_template(
        "portal/inventory/fixed_assets/cycle_view.html",
        cycle=cycle,
        entries=entries,
        pagination=pagination,
        stats=cycle_result_counts(cycle.id),
        recent_scans=recent_scans,
        selected={"q": search, "status": status},
        can_manage=current_user.has_perm(STORE_MANAGE),
        can_edit=bool(
            current_user.has_perm(STORE_MANAGE)
            and cycle.finalized_at is None
        ),
        can_scan=_can_scan_cycle(cycle),
        can_export=current_user.has_perm(STORE_EXPORT) or current_user.has_perm(STORE_MANAGE),
    )


@portal_bp.route("/inventory/fixed-assets/cycles/<int:cycle_id>/refresh", methods=["POST"])
@login_required
@_perm_any(STORE_MANAGE)
def fixed_asset_cycle_refresh(cycle_id: int):
    cycle = InvFixedAssetCycle.query.get_or_404(cycle_id)
    if cycle.status != "ACTIVE":
        flash("لا يمكن تحديث نطاق دورة مغلقة.", "warning")
        return redirect(url_for("portal.fixed_asset_cycle_view", cycle_id=cycle.id))
    added = populate_cycle_entries(cycle)
    _audit(
        "INV_FIXED_ASSET_CYCLE_REFRESH",
        f"تحديث نطاق دورة {cycle.code}: إضافة {added} أصل",
        "INV_FIXED_ASSET_CYCLE",
        cycle.id,
    )
    db.session.commit()
    flash(f"تم تحديث النطاق وإضافة {added} أصل جديد.", "success")
    return redirect(url_for("portal.fixed_asset_cycle_view", cycle_id=cycle.id))


@portal_bp.route("/inventory/fixed-assets/cycles/<int:cycle_id>/scan", methods=["GET", "POST"])
@login_required
def fixed_asset_cycle_scan(cycle_id: int):
    cycle = InvFixedAssetCycle.query.get_or_404(cycle_id)
    mobile_mode = request.args.get("mobile") == "1" or bool(session.get("fixed_asset_mobile_mode"))
    if mobile_mode:
        device = _require_approved_mobile_device()
        session["fixed_asset_mobile_device_id"] = device.id
    if not _can_scan_cycle(cycle):
        abort(403)
    if cycle.status != "ACTIVE":
        flash("هذه الدورة مغلقة ولا تقبل عمليات جرد جديدة.", "warning")
        return redirect(url_for("portal.fixed_asset_cycle_view", cycle_id=cycle.id))
    session["fixed_asset_active_cycle_id"] = cycle.id
    if request.args.get("paired") == "1" or request.args.get("mobile") == "1":
        session["fixed_asset_mobile_mode"] = True
        session.modified = True

    if request.method == "POST":
        asset = (
            _find_asset_by_system_qr(request.form.get("asset_code"))
            if mobile_mode
            else find_asset_by_identifier(request.form.get("asset_code"))
        )
        if asset is None:
            flash("هذا QR ليس تابعًا للنظام وتم رفضه.", "danger")
            return redirect(url_for(
                "portal.fixed_asset_cycle_scan",
                cycle_id=cycle.id,
                mobile="1" if mobile_mode else None,
            ))
        try:
            warehouse_id = _optional_int(request.form.get("observed_warehouse_id"))
            room_id = _optional_int(request.form.get("observed_room_id"))
            custodian_user_id = _optional_int(request.form.get("observed_custodian_user_id"))
            _entity_or_none(InvWarehouse, warehouse_id, "المستودع")
            _entity_or_none(InvRoom, room_id, "الغرفة")
            _entity_or_none(User, custodian_user_id, "الموظف المسؤول")
            condition = (request.form.get("observed_condition") or asset.asset_condition or "GOOD").upper()
            lifecycle_status = (request.form.get("observed_lifecycle_status") or asset.lifecycle_status or "ACTIVE").upper()
            if condition not in ASSET_CONDITIONS or lifecycle_status not in ASSET_LIFECYCLE_STATUSES:
                raise ValueError("الحالة المسجلة غير صالحة.")
            entry = record_asset_scan(
                cycle,
                asset,
                scanned_by_id=current_user.id,
                observed_warehouse_id=warehouse_id,
                observed_room_id=room_id,
                observed_custodian_user_id=custodian_user_id,
                observed_condition=condition,
                observed_lifecycle_status=lifecycle_status,
                note=_clean_text(request.form.get("note")),
                source=request.form.get("scan_source") or "MANUAL",
                device_info=request.user_agent.string,
            )
            _audit(
                "INV_FIXED_ASSET_SCAN",
                f"جرد الأصل {asset.asset_tag}: {INVENTORY_RESULT_LABELS.get(entry.result_status, entry.result_status)}",
                "INV_FIXED_ASSET_CYCLE",
                cycle.id,
            )
            db.session.commit()
            flash(
                f"تم تسجيل {asset.asset_tag}: {INVENTORY_RESULT_LABELS.get(entry.result_status, entry.result_status)}.",
                "success",
            )
            return redirect(url_for(
                "portal.fixed_asset_cycle_scan",
                cycle_id=cycle.id,
                last=asset.id,
                mobile="1" if mobile_mode else None,
            ))
        except ValueError as exc:
            db.session.rollback()
            flash(str(exc), "danger")

    recent_scans = (
        InvFixedAssetScanLog.query
        .filter(InvFixedAssetScanLog.cycle_id == cycle.id)
        .order_by(InvFixedAssetScanLog.scanned_at.desc())
        .limit(12)
        .all()
    )
    return render_template(
        "portal/inventory/fixed_assets/scan.html",
        cycle=cycle,
        stats=cycle_result_counts(cycle.id),
        recent_scans=recent_scans,
        initial_code=normalize_asset_identifier(request.args.get("code")),
        mobile_mode=mobile_mode,
        **_choices(include_items=False, include_categories=False),
    )


@portal_bp.route("/inventory/fixed-assets/cycles/<int:cycle_id>/scan/lookup")
@login_required
def fixed_asset_cycle_lookup(cycle_id: int):
    cycle = InvFixedAssetCycle.query.get_or_404(cycle_id)
    mobile_mode = request.args.get("mobile") == "1" or bool(session.get("fixed_asset_mobile_mode"))
    if mobile_mode:
        _require_approved_mobile_device()
    if not _can_scan_cycle(cycle):
        abort(403)
    if cycle.status != "ACTIVE":
        return jsonify({"ok": False, "message": "دورة الجرد مغلقة."}), 409
    asset = (
        _find_asset_by_system_qr(request.args.get("code"))
        if mobile_mode
        else find_asset_by_identifier(request.args.get("code"))
    )
    if asset is None:
        return jsonify({"ok": False, "message": "هذا QR ليس تابعًا للنظام وتم رفضه."}), 404
    entry = InvFixedAssetEntry.query.filter_by(cycle_id=cycle.id, asset_id=asset.id).first()
    defaults = {
        "warehouse_id": (entry.observed_warehouse_id if entry and entry.scanned_at else (entry.expected_warehouse_id if entry else asset.warehouse_id)),
        "room_id": (entry.observed_room_id if entry and entry.scanned_at else (entry.expected_room_id if entry else asset.room_id)),
        "custodian_user_id": (entry.observed_custodian_user_id if entry and entry.scanned_at else (entry.expected_custodian_user_id if entry else asset.custodian_user_id)),
        "condition": (entry.observed_condition if entry and entry.scanned_at else (entry.expected_condition if entry else asset.asset_condition)),
        "lifecycle_status": (entry.observed_lifecycle_status if entry and entry.scanned_at else (entry.expected_lifecycle_status if entry else asset.lifecycle_status)),
        "note": entry.note if entry else "",
    }
    return jsonify({
        "ok": True,
        "asset": {
            "id": asset.id,
            "code": asset.qr_token,
            "asset_tag": asset.asset_tag,
            "name": asset.name,
            "serial_number": asset.serial_number or "",
            "manufacturer": asset.manufacturer or "",
            "model": asset.model or "",
            "location": asset.location_label,
            "image_url": (
                url_for("portal.fixed_asset_image", asset_id=asset.id)
                if asset.image_filename
                else None
            ),
            "view_url": url_for("portal.fixed_asset_view", asset_id=asset.id),
        },
        "entry": {
            "expected": bool(entry and entry.was_expected),
            "already_scanned": bool(entry and entry.scanned_at),
            "scan_count": int(entry.scan_count or 0) if entry else 0,
            "result_status": entry.result_status if entry else "UNEXPECTED",
            "result_label": INVENTORY_RESULT_LABELS.get(entry.result_status, entry.result_status) if entry else "خارج نطاق الجرد",
            "expected_location": entry.expected_location_label if entry else asset.location_label,
            "expected_condition": ASSET_CONDITIONS.get(entry.expected_condition, entry.expected_condition) if entry else ASSET_CONDITIONS.get(asset.asset_condition, asset.asset_condition),
            "previous_result": PREVIOUS_RESULT_LABELS.get(entry.previous_result_status, entry.previous_result_status) if entry else "خارج نطاق الدورة",
        },
        "defaults": defaults,
    })


@portal_bp.route("/inventory/fixed-assets/cycles/<int:cycle_id>/close", methods=["POST"])
@login_required
@_perm_any(STORE_MANAGE)
def fixed_asset_cycle_close(cycle_id: int):
    cycle = InvFixedAssetCycle.query.get_or_404(cycle_id)
    try:
        result = close_inventory_cycle(
            cycle,
            closed_by_id=current_user.id,
            apply_observed_values=bool(request.form.get("apply_observed_values")),
        )
        _audit(
            "INV_FIXED_ASSET_CYCLE_CLOSE",
            f"إغلاق دورة {cycle.code}: مفقود {result['missing']}، تسوية {result['reconciled']}",
            "INV_FIXED_ASSET_CYCLE",
            cycle.id,
        )
        db.session.commit()
        flash(
            f"أُغلقت الدورة. صُنّف {result['missing']} أصل كغير موجود، وحُدّث {result['reconciled']} سجل أصل.",
            "success",
        )
    except ValueError as exc:
        db.session.rollback()
        flash(str(exc), "warning")
    return redirect(url_for("portal.fixed_asset_cycle_view", cycle_id=cycle.id))


@portal_bp.route("/inventory/fixed-assets/cycles/<int:cycle_id>/reopen", methods=["POST"])
@login_required
@_perm_any(STORE_MANAGE)
def fixed_asset_cycle_reopen(cycle_id: int):
    cycle = InvFixedAssetCycle.query.get_or_404(cycle_id)
    try:
        reset_count = reopen_inventory_cycle(cycle)
        _audit(
            "INV_FIXED_ASSET_CYCLE_REOPEN",
            f"إعادة فتح دورة {cycle.code} وإعادة {reset_count} أصل إلى الانتظار",
            "INV_FIXED_ASSET_CYCLE",
            cycle.id,
        )
        db.session.commit()
        flash("تمت إعادة فتح دورة الجرد.", "success")
    except ValueError as exc:
        db.session.rollback()
        flash(str(exc), "warning")
    return redirect(url_for("portal.fixed_asset_cycle_view", cycle_id=cycle.id))


@portal_bp.route("/inventory/fixed-assets/cycles/<int:cycle_id>/export.xlsx")
@login_required
@_perm_any(STORE_EXPORT, STORE_MANAGE)
def fixed_asset_cycle_export(cycle_id: int):
    cycle = InvFixedAssetCycle.query.get_or_404(cycle_id)
    entries = (
        InvFixedAssetEntry.query
        .filter(InvFixedAssetEntry.cycle_id == cycle.id)
        .join(InvFixedAsset)
        .order_by(InvFixedAsset.asset_tag.asc())
        .all()
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "نتائج الجرد"
    sheet.append([
        "رقم الدورة", "اسم الدورة", "اللجنة", "الدورة السابقة", "رقم الأصل", "اسم الأصل",
        "التصنيف", "الرقم التسلسلي", "نتيجة اللجنة السابقة", "الموقع المتوقع",
        "صاحب العهدة المتوقع", "الحالة المتوقعة", "الوضع المتوقع", "نتيجة الجرد الحالي",
        "الموقع الفعلي", "صاحب العهدة الفعلي", "الحالة الفعلية", "الوضع الفعلي",
        "وقت القراءة", "بواسطة", "عدد القراءات", "ملاحظات",
    ])
    for entry in entries:
        asset = entry.asset
        sheet.append([
            cycle.code,
            cycle.name,
            cycle.committee_name or "",
            cycle.previous_cycle.code if cycle.previous_cycle else "",
            asset.asset_tag,
            asset.name,
            asset.category.name if asset.category else "",
            asset.serial_number or "",
            PREVIOUS_RESULT_LABELS.get(entry.previous_result_status, entry.previous_result_status or ""),
            entry.expected_location_label,
            entry.expected_custodian.full_name if entry.expected_custodian else "",
            ASSET_CONDITIONS.get(entry.expected_condition, entry.expected_condition or ""),
            ASSET_LIFECYCLE_STATUSES.get(entry.expected_lifecycle_status, entry.expected_lifecycle_status or ""),
            INVENTORY_RESULT_LABELS.get(entry.result_status, entry.result_status),
            entry.observed_location_label if entry.scanned_at else "",
            entry.observed_custodian.full_name if entry.observed_custodian else "",
            ASSET_CONDITIONS.get(entry.observed_condition, entry.observed_condition or ""),
            ASSET_LIFECYCLE_STATUSES.get(entry.observed_lifecycle_status, entry.observed_lifecycle_status or ""),
            entry.scanned_at,
            entry.scanned_by.full_name if entry.scanned_by else "",
            entry.scan_count,
            entry.note or "",
        ])
    _style_worksheet(sheet, {2: 28, 3: 24, 6: 28, 9: 25, 10: 28, 14: 28, 15: 28, 22: 35})
    from services.asset_custody import KINDS, STATUSES
    approvals = workbook.create_sheet("اعتماد الموظفين والعهد")
    approvals.append(["المستند", "النوع", "الموظف", "الحالة", "الأصل", "قرار الموظف", "ملاحظات الموظف", "وقت المراجعة", "وقت الإصدار"])
    for doc in InvAssetDocument.query.filter_by(cycle_id=cycle.id).order_by(InvAssetDocument.id).all():
        for line in doc.lines:
            approvals.append([doc.id, KINDS[doc.kind], doc.employee.full_name, STATUSES[doc.status],
                              line.snapshot["tag"], {"PENDING": "لم يراجع", "ACCEPT": "موافق", "OBJECT": "اعتراض"}.get(line.decision, line.decision),
                              line.employee_note or "", doc.reviewed_at, doc.issued_at])
    _style_worksheet(approvals, {3: 28, 7: 40})
    return _workbook_response(workbook, f"fixed-asset-inventory-{cycle.code}.xlsx")
