# -*- coding: utf-8 -*-
"""Create the consumable-material catalogue prepared from the supplied image.

The first worksheet deliberately uses the column names understood by the
portal's inventory catalogue importer.  The quantity and room columns are
kept in the same workbook for the subsequent opening-stock/room step; the
current catalogue importer ignores those two columns by design.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


CATEGORY = "قرطاسية / ادوات مكتبية"
ROOM = "غرفة مدير المستودع"


ITEMS = [
    {
        "code": "PNCECS-MGR-001",
        "name": "دفتر ملاحظات A4",
        "subcategory": "ورق ودفاتر ملاحظات",
        "specification": "A4",
        "quantity": 100,
    },
    {
        "code": "PNCECS-MGR-002",
        "name": "مشابك ورق مكتبي",
        "subcategory": "كليبسات-مشابك-مطاط",
        "specification": "50 mm",
        "quantity": 28,
    },
    {
        "code": "PNCECS-MGR-003",
        "name": "ملاقط ورق Binder Clips",
        "subcategory": "كليبسات-مشابك-مطاط",
        "specification": "15/8",
        "quantity": 38,
    },
    {
        "code": "PNCECS-MGR-004",
        "name": "ملاقط ورق Binder Clips",
        "subcategory": "كليبسات-مشابك-مطاط",
        "specification": "19 mm",
        "quantity": 50,
    },
    {
        "code": "PNCECS-MGR-005",
        "name": "كباسات مكبس",
        "subcategory": "دبابيس",
        "specification": "عادية - 26/6",
        "quantity": 28,
    },
    {
        "code": "PNCECS-MGR-006",
        "name": "كباسات مكبس",
        "subcategory": "دبابيس",
        "specification": "23/10",
        "quantity": 10,
    },
    {
        "code": "PNCECS-MGR-007",
        "name": "كباسات مكبس",
        "subcategory": "دبابيس",
        "specification": "23/17",
        "quantity": 10,
    },
    {
        "code": "PNCECS-MGR-008",
        "name": "ورق ملاحظات ملون",
        "subcategory": "ورق ودفاتر ملاحظات",
        "specification": "-",
        "quantity": 60,
    },
    {
        "code": "PNCECS-MGR-009",
        "name": "مقص ورق",
        "subcategory": "مقصات وشفرات",
        "specification": "-",
        "quantity": 22,
    },
    {
        "code": "PNCECS-MGR-010",
        "name": "قلم حبر سائل",
        "subcategory": "اقلام سائل",
        "specification": "0.7",
        "quantity": 100,
    },
    {
        "code": "PNCECS-MGR-011",
        "name": "قلم حبر جاف",
        "subcategory": "اقلام حبر",
        "specification": "-",
        "quantity": 140,
    },
    {
        "code": "PNCECS-MGR-012",
        "name": "قلم رصاص",
        "subcategory": "اقلام رصاص",
        "specification": "-",
        "quantity": 100,
    },
    {
        "code": "PNCECS-MGR-013",
        "name": "قلم ماسح حبر",
        "subcategory": "اقلام تخطيط,تصحيح,ملاحظات",
        "specification": "Correction Pen",
        "quantity": 60,
    },
    {
        "code": "PNCECS-MGR-014",
        "name": "خلاعة دبابيس ورق",
        "subcategory": "خلاعة دبابيس",
        "specification": "-",
        "quantity": 10,
    },
    {
        "code": "PNCECS-MGR-015",
        "name": "ورق ملاحظات ملون لاصق",
        "subcategory": "ورق ودفاتر ملاحظات",
        "specification": "-",
        "quantity": 60,
    },
    {
        "code": "PNCECS-MGR-016",
        "name": "أقلام HighLight ملون",
        "subcategory": "اقلام تخطيط,تصحيح,ملاحظات",
        "specification": "-",
        "quantity": 40,
    },
    {
        "code": "PNCECS-MGR-017",
        "name": "ملفات نايلون شفاف ماعون",
        "subcategory": "ملفات ودوسيات",
        "specification": "-",
        "quantity": 10,
    },
    {
        "code": "PNCECS-MGR-018",
        "name": "حاضنة أقلام زري",
        "subcategory": "اكسسوارات ولوازم طاولة مكتب",
        "specification": "-",
        "quantity": 35,
    },
    {
        "code": "PNCECS-MGR-019",
        "name": "لاصق عريض ورقي",
        "subcategory": "لاصق",
        "specification": "-",
        "quantity": 24,
    },
    {
        "code": "PNCECS-MGR-020",
        "name": "آلة حاسبة",
        "subcategory": "آلة حاسبة",
        "specification": "-",
        "quantity": 5,
    },
]


def _style_sheet(sheet, widths: dict[int, int]) -> None:
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in widths.items():
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 32


def build_workbook(output: Path) -> None:
    workbook = Workbook()
    catalogue = workbook.active
    catalogue.title = "الأصناف"
    catalogue_headers = [
        "رمز الصنف",
        "اسم الصنف",
        "التصنيف الرئيسي",
        "التصنيف الفرعي",
        "مستهلك",
        "الوحدة",
        "المواصفة/المقاس",
        "الكمية",
        "الغرفة",
    ]
    catalogue.append(catalogue_headers)
    for item in ITEMS:
        catalogue.append(
            [
                item["code"],
                item["name"],
                CATEGORY,
                item["subcategory"],
                "YES",
                "PCS",
                item["specification"],
                item["quantity"],
                ROOM,
            ]
        )
    _style_sheet(catalogue, {1: 20, 2: 30, 3: 28, 4: 32, 5: 12, 6: 12, 7: 22, 8: 12, 9: 28})

    balances = workbook.create_sheet("الأرصدة الأولية")
    balances.append(["رمز الصنف", "اسم الصنف", "الكمية", "الوحدة", "الغرفة"])
    for item in ITEMS:
        balances.append([item["code"], item["name"], item["quantity"], "PCS", ROOM])
    _style_sheet(balances, {1: 20, 2: 30, 3: 12, 4: 12, 5: 28})

    instructions = workbook.create_sheet("تعليمات")
    instructions.sheet_view.rightToLeft = True
    instructions.append(["ملاحظة", "التفاصيل"])
    instructions.append(["عدد الأصناف", len(ITEMS)])
    instructions.append(["نوع الأصناف", "مستهلكة - YES"])
    instructions.append(["الغرفة المطلوبة", ROOM])
    instructions.append(["رفع دليل الأصناف", "ارفع الملف من بوابة النظام > المستودع > إدارة الأصناف > استيراد دليل الجرد."])
    instructions.append(["الرصيد", "المستورد الحالي ينشئ تعريف الأصناف فقط؛ استخدم ورقة الأرصدة الأولية لإدخال الكميات في سند جرد للمستودع."])
    instructions.append(["الأكواد", "أكواد داخلية ثابتة PNCECS لأن الصورة المرفقة لا تحتوي على رموز أصناف رسمية."])
    _style_sheet(instructions, {1: 24, 2: 100})
    instructions.auto_filter.ref = "A1:B7"

    workbook.properties.title = "أصناف مستهلكة لطلبات المواد - غرفة مدير المستودع"
    workbook.properties.subject = "Inventory consumables catalogue"
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("مواد_مستهلكة_طلب_المواد_غرفة_مدير_المستودع.xlsx"),
        help="Destination XLSX path.",
    )
    args = parser.parse_args()
    build_workbook(args.output)
    # Keep the CLI message ASCII-safe on Windows consoles using a legacy code page.
    print(f"Created workbook with {len(ITEMS)} items.")


if __name__ == "__main__":
    main()
